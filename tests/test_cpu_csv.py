import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import process_php_csvs as tool


def raw_data():
    return pd.DataFrame({'RelTime': np.arange(600, dtype=float),
        'T_CPU_1': 55., 'T_CPU_3': 56., 'T_CPU_5': 50.,
        'T_WATER_IN': 25., 'T_WATER_OUT': 26., 'VFR_WATER': 2.7,
        'T_EVAP_IN': 25., 'T_EVAP_OUT': 31., 'T_COND_IN': 30., 'T_COND_OUT': 26.,
        'W_PSU': 999., 'I_PSU_1': 9., 'V_PSU_2_SP': 48., 'T_PSU_1': 45.})


def config():
    return dict(board_count=3, board_columns=['T_CPU_1', 'T_CPU_3', 'T_CPU_5'],
                start_s=10., durations_s=[200., 200.], cpu_powers=[[10., 20., 0.], [20., 10., 0.]])


def saved_data():
    data = raw_data()
    boards, _, schedule = tool.build_shift2dc_steps(data, config())
    return tool.save_shift2dc_schedule_columns(data, boards, schedule)


def excel_results():
    """Different CPU sets, legacy sensor names, and unequal/zero CPU loads."""
    results = []
    for ratio in [40, 50, 54, 60, 70]:
        data, setup = raw_data(), config()
        if ratio == 54:
            data = data.rename(columns={'T_CPU_3': 'T_BOARD_10'})
            setup.update(board_count=2, board_columns=['T_CPU_1', 'T_BOARD_10'],
                         durations_s=[200.], cpu_powers=[[12.5, 37.5]])
        boards, steps, schedule = tool.build_shift2dc_steps(data, setup)
        name = f'Shift2DC_LTS_EVAP-demo_COND-demo_R1336mzzE_FR{ratio}_Water_TW25_VFR2p7_SS.csv'
        metadata = tool.parse_file_name(Path(name))
        metadata['source_file'] = name
        results.extend(tool.shift2dc_result(step, metadata, boards, row)
                       for step, row in zip(steps, schedule))
    return results


class CPUCsvTests(unittest.TestCase):
    def test_excel_cpu_pairing_and_fr_colors(self):
        rows = excel_results()
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)/'averages.xlsx'
            tool.write_shift2dc_excel(rows, output)
            workbook = load_workbook(output)
            self.addCleanup(workbook.close)
            sheet = workbook['Test averages']
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(len(headers), len(set(headers)))
            self.assertEqual(sheet.max_row, len(rows) + 1)
            for cpu in [1, 3, 5, 10]:
                temp = f'T_CPU_{cpu} [°C]'
                self.assertEqual(headers[headers.index(temp) + 1], f'W_CPU_{cpu} [W]')
            self.assertLess(headers.index('T_CPU_5 [°C]'), headers.index('T_CPU_10 [°C]'))
            expected = {(r['Source File'], r['Step']): r for r in rows}
            colors = {40: '00B050', 50: 'FFC000', 54: 'FF7300', 60: 'FF0000', 70: '7030A0'}
            for row_index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                actual = dict(zip(headers, values))
                source = expected[(actual['Source File'], actual['Step'])]
                for cpu in [1, 3, 5, 10]:
                    for column in [f'T_CPU_{cpu} [°C]', f'W_CPU_{cpu} [W]']:
                        self.assertEqual(actual[column], source.get(column))
                self.assertEqual(actual['Total heat load [W]'], source['Scheduled total power [W]'])
                fr = actual['FR [%]']
                fill = sheet.cell(row_index, headers.index('FR [%]') + 1).fill
                self.assertEqual(fill.fgColor.rgb[-6:], colors[fr])
                self.assertEqual(tool.filling_ratio_line_color(fr, 0), '#' + colors[fr])

    def test_zero_off_intervals_and_vector_plateaus(self):
        saved = saved_data()
        self.assertFalse(set(tool.SHIFT2DC_LEGACY_SCHEDULE_COLUMNS) & set(saved))
        self.assertEqual(set(tool.shift2dc_cpu_power_columns(saved)), {1, 3, 5})
        self.assertTrue((saved.loc[(saved.RelTime < 10) | (saved.RelTime >= 410), ['W_CPU_1', 'W_CPU_3', 'W_CPU_5']] == 0).all().all())
        self.assertTrue((saved['W_CPU_5'] == 0).all())
        for column in ['W_PSU', 'I_PSU_1', 'V_PSU_2_SP']:
            self.assertNotIn(column, saved)
        self.assertTrue(saved['T_PSU_1'].equals(raw_data()['T_PSU_1']))
        _, _, recovered = tool.build_shift2dc_saved_steps(saved)
        self.assertEqual(len(recovered), 2)
        self.assertEqual([r['Scheduled total power [W]'] for r in recovered], [30., 30.])
        self.assertEqual([r['W_CPU_1 [W]'] for r in recovered], [10., 20.])
        self.assertEqual([r['Average from [s]'] for r in recovered], [110., 310.])

    def test_legacy_migration_and_validation(self):
        saved = saved_data()
        saved.loc[saved.RelTime < 10, ['W_CPU_1', 'W_CPU_3', 'W_CPU_5']] = np.nan
        saved['Heat_Load_Step_Start_s'] = 'obsolete'
        saved['Total_Heat_Load_W'] = 99999
        saved['Scheduled_Power_Per_Board_W'] = 999
        cleaned = tool.clean_shift2dc_power_channels(saved)
        self.assertTrue((cleaned.loc[cleaned.RelTime < 10, 'W_CPU_1'] == 0).all())
        self.assertFalse(set(tool.SHIFT2DC_LEGACY_SCHEDULE_COLUMNS) & set(cleaned))
        self.assertEqual(len(tool.build_shift2dc_saved_steps(cleaned)[2]), 2)
        cleaned.loc[20, 'W_CPU_3'] = np.nan
        with self.assertRaisesRegex(ValueError, 'incomplete'):
            tool.build_shift2dc_saved_steps(tool.clean_shift2dc_power_channels(cleaned))
        cleaned.loc[20, 'W_CPU_3'] = -1
        with self.assertRaisesRegex(ValueError, 'negative'):
            tool.build_shift2dc_saved_steps(cleaned)

    def test_off_gap_and_transient_total(self):
        data = raw_data()
        data['W_CPU_1'] = 0.
        data['W_CPU_3'] = 0.
        data.loc[10:209, ['W_CPU_1', 'W_CPU_3']] = [10., 20.]
        data.loc[220:419, ['W_CPU_1', 'W_CPU_3']] = [10., 20.]
        cleaned = tool.clean_shift2dc_power_channels(data)
        self.assertEqual(len(tool.build_shift2dc_saved_steps(cleaned)[2]), 2)
        metadata = tool.parse_file_name(Path('Shift2DC_LTS_EVAP-demo_COND-demo_R1336mzzE_FR60_Water_TW25_VFR2p7_TR.csv'))
        detail = dict(data=cleaned, metadata=metadata, boards=['T_CPU_1', 'T_CPU_3'], t_cu_columns=['T_CPU_1', 'T_CPU_3'])
        panels = tool.transient_panels(detail, {})
        total = next(series['Total heat load'] for _, _, series in panels if 'Total heat load' in series)
        self.assertEqual(total.iloc[0], 0.)
        self.assertEqual(total.iloc[100], 30.)

    def test_raw_and_postprocessed_roundtrip_without_dialog(self):
        name = 'Shift2DC_LTS_EVAP-demo_COND-demo_R1336mzzE_FR60_Water_TW25_VFR2p7_SS.csv'
        with tempfile.TemporaryDirectory() as tmp, \
             patch.object(tool, 'configure_report', return_value=({}, {})), \
             patch.object(tool, 'write_shift2dc_excel'), \
             patch.object(tool, 'write_shift2dc_pdf'), \
             patch.object(tool, 'show_shift2dc_dialog', side_effect=AssertionError('Unexpected schedule prompt')):
            root = Path(tmp)
            source = root/name
            raw_data().to_csv(source, index=False)
            first = tool.process_shift2dc_files([source], root/'out', configurations={name: config()}, open_pdf_when_done=False)
            raw = pd.read_csv(source)
            pd.testing.assert_frame_equal(raw, pd.read_csv(root/'out'/name))
            self.assertEqual(raw.RelTime.iloc[-1], 510.)
            self.assertEqual(len(raw), 511)
            self.assertTrue((raw.loc[raw.RelTime >= 410, ['W_CPU_1', 'W_CPU_3', 'W_CPU_5']] == 0).all().all())
            pd.testing.assert_frame_equal(raw, first['details'][0]['data'], check_dtype=False)
            self.assertFalse(set(tool.SHIFT2DC_LEGACY_SCHEDULE_COLUMNS) & set(raw))
            second = tool.process_shift2dc_files([source], root/'out', configurations={}, open_pdf_when_done=False)
            pd.testing.assert_frame_equal(pd.DataFrame(first['results']), pd.DataFrame(second['results']))
            pd.testing.assert_frame_equal(raw, pd.read_csv(source))

    def test_existing_csv_cutoff_uses_seconds_and_preserves_averages(self):
        name = 'Shift2DC_LTS_EVAP-demo_COND-demo_R1336mzzE_FR60_Water_TW25_VFR2p7_SS.csv'
        time_grids = [np.arange(0, 601, dt) for dt in [.5, 1., 2.]]
        time_grids.append(np.unique(np.r_[np.arange(0, 410, .8), 410.,
                                          np.arange(411.3, 601, 1.7), 510., 510.1]))
        for time in time_grids:
            with self.subTest(samples=len(time)), tempfile.TemporaryDirectory() as tmp, \
                 patch.object(tool, 'configure_report', return_value=({}, {})), \
                 patch.object(tool, 'write_shift2dc_excel'), \
                 patch.object(tool, 'write_shift2dc_pdf'), \
                 patch.object(tool, 'show_shift2dc_dialog', side_effect=AssertionError('Unexpected schedule prompt')):
                data = raw_data().iloc[np.zeros(len(time), dtype=int)].reset_index(drop=True)
                data['RelTime'] = time
                data['T_CPU_1'] = 35. + time / 100.
                boards, _, schedule = tool.build_shift2dc_steps(data, config())
                saved = tool.save_shift2dc_schedule_columns(data, boards, schedule)
                boards, steps, schedule = tool.build_shift2dc_saved_steps(saved)
                cutoff = schedule[-1]['End [s]'] + 100.
                root = Path(tmp)
                source = root/name
                saved.to_csv(source, index=False)
                first = tool.process_shift2dc_files([source], root/'out', open_pdf_when_done=False)
                actual = pd.read_csv(source)
                expected = saved.loc[saved.RelTime <= cutoff].reset_index(drop=True)
                pd.testing.assert_frame_equal(actual, expected, check_dtype=False)
                pd.testing.assert_frame_equal(actual, pd.read_csv(root/'out'/name))
                for before, after in zip(steps, first['details'][0]['steps']):
                    np.testing.assert_allclose(before['T_CPU_1'], after['T_CPU_1'])
                second = tool.process_shift2dc_files([source], root/'out', open_pdf_when_done=False)
                pd.testing.assert_frame_equal(pd.DataFrame(first['results']), pd.DataFrame(second['results']))
                pd.testing.assert_frame_equal(actual, pd.read_csv(source))

    def test_cooldown_preserves_short_tails_and_later_loads(self):
        saved = saved_data()
        for end in [410., 460., 510.]:
            with self.subTest(end=end):
                shorter = saved.loc[saved.RelTime <= end]
                pd.testing.assert_frame_equal(tool.trim_shift2dc_cooldown(shorter), shorter)
        # A brief final load is excluded from 100 s averages, but is still data.
        saved.loc[520:529, 'W_CPU_5'] = 10.
        extended = pd.concat([saved, saved.iloc[-1:].assign(RelTime=630.),
                              saved.iloc[-1:].assign(RelTime=631.)], ignore_index=True)
        tool.build_shift2dc_saved_steps(extended)
        trimmed = tool.trim_shift2dc_cooldown(extended)
        self.assertEqual(trimmed.RelTime.iloc[-1], 630.)
        self.assertEqual(trimmed.loc[525, 'W_CPU_5'], 10.)
        all_off = saved.copy()
        all_off[['W_CPU_1', 'W_CPU_3', 'W_CPU_5']] = 0.
        pd.testing.assert_frame_equal(tool.trim_shift2dc_cooldown(all_off), all_off)
        powered_at_end = saved.copy()
        powered_at_end.loc[410:, 'W_CPU_5'] = 10.
        pd.testing.assert_frame_equal(tool.trim_shift2dc_cooldown(powered_at_end), powered_at_end)

    def test_transient_recording_keeps_full_duration(self):
        name = 'Shift2DC_LTS_EVAP-demo_COND-demo_R1336mzzE_FR60_Water_TW25_VFR2p7_TR.csv'
        with tempfile.TemporaryDirectory() as tmp, \
             patch.object(tool, 'configure_report', return_value=({}, {})), \
             patch.object(tool, 'write_shift2dc_excel'), \
             patch.object(tool, 'write_shift2dc_pdf'):
            root = Path(tmp)
            source = root/name
            saved = saved_data()
            saved.to_csv(source, index=False)
            tool.process_shift2dc_files([source], root/'out', open_pdf_when_done=False)
            pd.testing.assert_frame_equal(pd.read_csv(source), saved)
            pd.testing.assert_frame_equal(pd.read_csv(root/'out'/name), saved)


if __name__ == '__main__':
    unittest.main()
