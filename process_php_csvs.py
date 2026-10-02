import json
import math
import os
import re
import subprocess
import sys
from datetime import date
from io import BytesIO
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties, findfont
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors as reportlab_colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


# USER SETTINGS
sample_size = 100
minimum_step_size = 100
power_change_tolerance = 1e-6

# REFPROP fluid values are read at runtime. Install NIST REFPROP and either
# the ctREFPROP Python bridge or CoolProp with access to the REFPROP backend.

# Properties used for W_OUT. VFR must be in litres per minute.
water_density = 997.0       # kg/m3
water_cp = 4180.0           # J/(kg K)
air_density = 1.204         # kg/m3
air_cp = 1005.0             # J/(kg K)
cfm_to_m3_s = 0.00047194745 # m3/s per CFM

# Filling-ratio color anchors taken from the supplied reference image.
# Intermediate ratios interpolate in RGB; outside this range use the end color.
filling_ratio_colors = {
    70: "7030A0",  # Purple
    60: "FF0000",  # Red
    50: "FFC000",  # Yellow/orange
    40: "00B050",  # Green
}

# Fixed visual encoding used on performance-characterization pages.
# Temperatures between these 10°C anchors are interpolated. Values outside
# the 10-60°C range use the nearest end color. Within one temperature color,
# lower flow rates are lighter and higher flow rates are darker.
performance_temperature_color_scale = {
    10: "#2166AC",  # Blue
    20: "#4393C3",  # Light blue
    30: "#67B8A5",  # Blue-green
    40: "#F6D55C",  # Yellow
    50: "#F28E2B",  # Orange
    60: "#D62728",  # Red
}
performance_low_flow_lightening = 0.38
performance_high_flow_darkening = 0.18

# Report appearance.
report_green = reportlab_colors.HexColor("#1B5E20")
report_dark = reportlab_colors.HexColor("#222222")
report_grey = reportlab_colors.HexColor("#666666")
report_light_grey = reportlab_colors.HexColor("#D9D9D9")

# Use Matplotlib's bundled DejaVu fonts so the report renders identically
# on Windows and Linux instead of depending on a PDF viewer's font substitute.
report_regular_font = "JJReportSans"
report_bold_font = "JJReportSansBold"
pdfmetrics.registerFont(TTFont(
    report_regular_font,
    findfont(FontProperties(family="DejaVu Sans", weight="normal")),
))
pdfmetrics.registerFont(TTFont(
    report_bold_font,
    findfont(FontProperties(family="DejaVu Sans", weight="bold")),
))

lts_temperature_columns = [
    "T_EVAP_IN",
    "T_EVAP_OUT",
    "T_COND_IN",
    "T_COND_OUT",
]

lts_temperature_output_columns = [
    f"{column} [°C]"
    for column in lts_temperature_columns
]


def get_settings_file():
    """Return the persistent per-user settings-file location."""
    app_data = os.environ.get("APPDATA")

    if app_data:
        settings_folder = Path(app_data) / "JJCooling" / "DataProcessor"
    else:
        settings_folder = Path.home() / ".jjcooling" / "DataProcessor"

    settings_folder.mkdir(parents=True, exist_ok=True)
    return settings_folder / "settings.json"


def load_last_folder():
    """Load the last selected folder, falling back to an existing parent."""
    settings_file = get_settings_file()

    try:
        settings = json.loads(settings_file.read_text(encoding="utf-8"))
        folder = Path(settings["last_raw_data_folder"])
    except (FileNotFoundError, KeyError, TypeError, ValueError, OSError):
        return Path.home()

    # If the exact directory was moved or deleted, start at the closest
    # existing parent so the user can navigate back quickly.
    while not folder.exists() and folder != folder.parent:
        folder = folder.parent

    return folder if folder.exists() else Path.home()


def save_last_folder(folder):
    """Remember the selected 00_RawData folder for the next launch."""
    settings_file = get_settings_file()
    settings = {
        "last_raw_data_folder": str(Path(folder).resolve()),
    }
    settings_file.write_text(
        json.dumps(settings, indent=2),
        encoding="utf-8",
    )


def select_raw_data_folder():
    """Open a Windows dialog and ask the user to select 00_RawData."""
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    selected_folder = filedialog.askdirectory(
        title="Select the 00_RawData folder",
        initialdir=str(load_last_folder()),
        mustexist=True,
    )

    root.destroy()

    if not selected_folder:
        return None

    input_folder = Path(selected_folder)

    if input_folder.name.lower() != "00_rawdata":
        raise ValueError(
            "The selected directory must be named 00_RawData.\n"
            f"Selected directory: {input_folder}"
        )

    save_last_folder(input_folder)

    return input_folder


def refprop_fluid_candidates(fluid_name):
    """Return reasonable REFPROP spellings for a filename fluid token."""
    original = str(fluid_name).strip()
    compact = re.sub(r"[^A-Za-z0-9]", "", original)
    candidates = [original, compact, compact.upper()]

    # These aliases cover common filename and REFPROP naming differences.
    known_aliases = {
        "R1336MZZE": ["R1336MZZE", "R1336mzzE"],
        "R1336MZZZ": ["R1336MZZZ", "R1336mzzZ"],
        "R1233ZDE": ["R1233ZDE", "R1233zdE"],
    }
    candidates.extend(known_aliases.get(compact.upper(), []))

    unique_candidates = []
    for candidate in candidates:
        candidate = candidate.strip()
        if candidate and candidate.upper() not in {
            existing.upper() for existing in unique_candidates
        }:
            unique_candidates.append(candidate)
    return unique_candidates


def refprop_installation_folders():
    """Return configured and common REFPROP installation directories."""
    candidates = []
    for environment_name in ["RPPREFIX", "REFPROP_PATH", "REFPROP_HOME"]:
        value = os.environ.get(environment_name)
        if value:
            candidates.append(Path(value))

    if os.name == "nt":
        for environment_name in ["ProgramFiles", "ProgramFiles(x86)"]:
            program_files = os.environ.get(environment_name)
            if program_files:
                candidates.append(Path(program_files) / "REFPROP")
    else:
        candidates.extend([
            Path("/opt/refprop"),
            Path.home() / "REFPROP",
            Path.home() / "refprop",
        ])

    unique_folders = []
    for folder in candidates:
        folder_text = str(folder)
        if folder.exists() and folder_text not in unique_folders:
            unique_folders.append(folder_text)
    return unique_folders


def refprop_result_error(result):
    """Return a readable ctREFPROP error message."""
    error_text = getattr(result, "herr", "")
    if isinstance(error_text, bytes):
        error_text = error_text.decode("utf-8", errors="replace")
    return str(error_text).strip() or "REFPROP returned an unspecified error"


def calculate_critical_diameter_mm(surface_tension, liquid_density, vapor_density):
    """
    Calculate the PHP capillary critical diameter at a saturation state.

    d_crit = 2 * sqrt(sigma / (g * (rho_liquid - rho_vapor)))
    """
    density_difference = float(liquid_density) - float(vapor_density)
    if surface_tension <= 0 or density_difference <= 0:
        raise ValueError("invalid saturation properties for critical diameter")
    gravity = 9.80665
    diameter_m = 2.0 * math.sqrt(
        float(surface_tension) / (gravity * density_difference)
    )
    return diameter_m * 1000.0


def get_refprop_properties_with_ctrefprop(fluid_name):
    """Read the required pure-fluid properties through ctREFPROP."""
    from ctREFPROP.ctREFPROP import REFPROPFunctionLibrary

    installation_folders = refprop_installation_folders()
    if not installation_folders:
        # An empty prefix still works when the REFPROP DLL is on PATH.
        installation_folders = [""]

    errors = []
    for installation_folder in installation_folders:
        try:
            refprop = REFPROPFunctionLibrary(installation_folder)
            if installation_folder:
                refprop.SETPATHdll(installation_folder)
            units = refprop.GETENUMdll(0, "MASS BASE SI").iEnum
        except Exception as error:
            errors.append(str(error))
            continue

        for candidate in refprop_fluid_candidates(fluid_name):
            try:
                def saturation_state(temperature_c, quality, outputs):
                    result = refprop.REFPROPdll(
                        candidate,
                        "TQ",
                        outputs,
                        units,
                        1,
                        0,
                        float(temperature_c) + 273.15,
                        float(quality),
                        [1.0],
                    )
                    if getattr(result, "ierr", 0) > 0:
                        raise RuntimeError(refprop_result_error(result))
                    requested_value_count = len(outputs.split(";"))
                    values = [
                        float(value)
                        for value in result.Output[:requested_value_count]
                    ]
                    if not all(math.isfinite(value) for value in values):
                        raise RuntimeError("REFPROP returned a non-finite value")
                    return values

                pressure_20, = saturation_state(20.0, 0.0, "P")
                pressure_85, liquid_density, surface_tension = (
                    saturation_state(85.0, 0.0, "P;D;STN")
                )
                vapor_density, = saturation_state(85.0, 1.0, "D")
                critical_diameter = calculate_critical_diameter_mm(
                    surface_tension,
                    liquid_density,
                    vapor_density,
                )

                return {
                    "available": True,
                    "matched_fluid": candidate,
                    "critical_diameter_85_mm": critical_diameter,
                    "saturation_pressure_20_kpa": pressure_20 / 1000.0,  # MASS BASE SI returns Pa.
                    "saturation_pressure_85_kpa": pressure_85 / 1000.0,
                    "source": "REFPROP",
                }
            except Exception as error:
                errors.append(f"{candidate}: {error}")

    raise RuntimeError(errors[-1] if errors else "REFPROP could not be loaded")


def get_refprop_properties_with_coolprop(fluid_name):
    """Use CoolProp's REFPROP backend when ctREFPROP is unavailable."""
    from CoolProp.CoolProp import PropsSI

    errors = []
    for candidate in refprop_fluid_candidates(fluid_name):
        refprop_fluid = f"REFPROP::{candidate}"
        try:
            temperature_20 = 20.0 + 273.15
            temperature_85 = 85.0 + 273.15
            pressure_20 = PropsSI(
                "P", "T", temperature_20, "Q", 0, refprop_fluid
            ) / 1000.0
            pressure_85 = PropsSI(
                "P", "T", temperature_85, "Q", 0, refprop_fluid
            ) / 1000.0
            liquid_density = PropsSI(
                "Dmass", "T", temperature_85, "Q", 0, refprop_fluid
            )
            vapor_density = PropsSI(
                "Dmass", "T", temperature_85, "Q", 1, refprop_fluid
            )
            surface_tension = PropsSI(
                "surface_tension",
                "T",
                temperature_85,
                "Q",
                0,
                refprop_fluid,
            )
            critical_diameter = calculate_critical_diameter_mm(
                surface_tension,
                liquid_density,
                vapor_density,
            )
            return {
                "available": True,
                "matched_fluid": candidate,
                "critical_diameter_85_mm": critical_diameter,
                "saturation_pressure_20_kpa": float(pressure_20),
                "saturation_pressure_85_kpa": float(pressure_85),
                "source": "REFPROP",
            }
        except Exception as error:
            errors.append(f"{candidate}: {error}")

    raise RuntimeError(errors[-1] if errors else "REFPROP could not be loaded")


def get_refprop_fluid_properties(fluid_name):
    """Return report properties, without preventing processing on failure."""
    errors = []
    for property_reader in [
        get_refprop_properties_with_ctrefprop,
        get_refprop_properties_with_coolprop,
    ]:
        try:
            return property_reader(fluid_name)
        except Exception as error:
            message = str(error).replace("\n", " ").strip()
            errors.append(message[:240])

    return {
        "available": False,
        "matched_fluid": None,
        "critical_diameter_85_mm": None,
        "saturation_pressure_20_kpa": None,
        "saturation_pressure_85_kpa": None,
        "source": "REFPROP",
        "error": (
            " | ".join(errors)[:420]
            if errors
            else "REFPROP and its Python interface are unavailable"
        ),
    }


def format_refprop_property_line(
    fluid_name,
    properties,
    include_fluid=True,
    include_error=False,
):
    """Format one fluid's REFPROP values for the UI or report."""
    prefix = f"{fluid_name}: " if include_fluid else ""
    if not properties.get("available", False):
        message = prefix + "REFPROP properties unavailable"
        if include_error and properties.get("error"):
            message += f" - {properties['error']}"
            message += (
                ". Install ctREFPROP in this Python environment and make sure "
                "NIST REFPROP is installed."
            )
        return message
    return (
        prefix
        + f"d_crit(85°C) = {properties['critical_diameter_85_mm']:.2f} mm; "
        + f"P_sat(20°C) = {properties['saturation_pressure_20_kpa']:.2f} kPa; "
        + f"P_sat(85°C) = {properties['saturation_pressure_85_kpa']:.2f} kPa"
    )


# Slots are counted from RIGHT to LEFT in the supplied front-view drawing.
# Empty slots: 1, 6, 7. Board sensor IDs increase across occupied slots.
SHIFT2DC_BOARD_SLOTS = {1: 2, 2: 3, 3: 4, 4: 5, 5: 8, 6: 9, 7: 10, 8: 11}
# Approximate proportions traced from the drawing, not physical millimetres.
# Slot openings listed left to right; board height is 360 drawing units.
SHIFT2DC_SLOT_BOUNDS = [(143, 211), (222, 290), (302, 370), (382, 450),
                       (462, 550), (562, 650), (662, 730), (742, 810),
                       (821, 889), (901, 969), (981, 1031)]


def is_board_test_detail(detail):
    return bool(detail.get("boards")) or detail.get("raw_data_only", False)


def shift2dc_board_sensor_layout(columns):
    """Preserve the pictured gaps and stable board positions for every subset."""
    layout = []
    for column in sorted(columns, key=natural_text_sort_key):
        match = re.fullmatch(r"T_(?:CPU|BOARD)_(\d+)(?:\s*\[.*\])?", column, re.I)
        board_id = int(match.group(1)) if match else None
        if board_id not in SHIFT2DC_BOARD_SLOTS:
            raise ValueError(f"No physical slot is configured for {column}. Update SHIFT2DC_BOARD_SLOTS.")
        slot = SHIFT2DC_BOARD_SLOTS[board_id]
        if slot < 1 or slot > len(SHIFT2DC_SLOT_BOUNDS):
            raise ValueError(f"Invalid slot {slot} configured for {column}.")
        left, right = SHIFT2DC_SLOT_BOUNDS[-slot]
        layout.append((column, left - 143, 0, right - left, 360, 1))
    if len({item[1] for item in layout}) != len(layout):
        raise ValueError("Each board must have its own physical slot.")
    return layout, (-6, 894), (-6, 366)


def is_board_report(results):
    """Identify normalized board results without changing the common metric keys."""
    return (
        results is not None
        and 'Temperature source' in results
        and bool(len(results))
        and results['Temperature source'].eq('board').all()
    )


def report_display_text(text, results=None):
    """Use explicit coolant sensor labels and project-specific temperature names."""
    if results is not None and 'Coolant' in results:
        media = set(results['Coolant'].dropna().astype(str).str.lower())
        coolant = 'WATER' if media == {'water'} else 'AIR' if media == {'air'} else 'COOLANT'
        text = re.sub(r'\bT_(IN|OUT)\b', lambda match: f'T_{coolant}_{match[1]}', text)
    if not is_board_report(results):
        return text
    return (text.replace('T_CU', 'T_CPU')
            .replace('T_BOARD', 'T_CPU')
            .replace('Copper', 'CPU').replace('copper', 'CPU')
            .replace('Applied power, W_IN [W]', 'Total heat load [W]')
            .replace('Applied Power', 'Total heat load')
            .replace('W_PSU', 'W_IN'))


def shift2dc_report_rows(results):
    """Adapt server averages to the established reporting schema in memory.

    Generic T_CU keys belong to the existing calculation/plot interface. Actual
    board sensor names and displayed labels remain T_CPU throughout exports.
    """
    adapted = []
    for original in results:
        row = dict(original)
        for suffix in ['MIN [°C]', 'MAX [°C]', 'AVG [°C]', 'deltaLOW [K]', 'deltaHIGH [K]']:
            row['T_CU_' + suffix] = original['T_CPU_' + suffix]
        for suffix in [' [K]', '_MIN [K]', '_MAX [K]']:
            row['DeltaT_CU' + suffix] = original['ΔT_CPU' + suffix]
        row['W_IN [W]'] = original['Scheduled total power [W]']
        row['W_OUT [W]'] = original['Water heat removal [W]']
        row['Temperature source'] = 'board'
        adapted.append(row)
    return adapted


def is_transient_detail(detail):
    return detail['metadata'].get('test_mode', 'SS') == 'TR'


def prepare_transient_data(data):
    """Keep raw channels; combine two measured setpoints when both are available."""
    data = data.copy()
    if find_optional_column(data, ['W_PSU', 'W_HEATER']) is None:
        channels = [find_optional_column(data, [name]) for name in ['W_PSU_1_SP', 'W_PSU_2_SP']]
        if all(column is not None for column in channels):
            data['W_PSU'] = sum(pd.to_numeric(data[c], errors='coerce') for c in channels)
    return data


def report_campaign_frame(all_results, test_details):
    """Inventory-only TR rows supply cover metadata, never calculated averages."""
    source_rows = all_results.to_dict('records') if isinstance(all_results, pd.DataFrame) else all_results
    rows = [dict(row, **{'Test Mode': 'SS'}) for row in source_rows]
    for detail in test_details:
        if not is_transient_detail(detail):
            continue
        meta = detail['metadata']
        fields = {'Source File': 'source_file', 'Comment': 'comment', 'Report Prefix': 'report_prefix',
                  'Part Name': 'part_name', 'Part Number': 'part_number', 'Evaporator Name': 'evaporator_name',
                  'Condenser Name': 'condenser_name', 'Working Fluid': 'fluid', 'Orientation': 'orientation',
                  'Condition Type': 'condition_type', 'Condition Value': 'condition_value', 'Coolant': 'medium',
                  'Flow Type': 'flow_type', 'Flow Unit': 'flow_unit', 'Nominal Flow Rate': 'nominal_flow_rate',
                  'Nominal T_IN [°C]': 'nominal_temperature'}
        row = {label: meta.get(key) for label, key in fields.items()}
        row['Source File'] = detail['source_file']
        row['Test Mode'] = 'TR'
        row['Temperature source'] = 'board' if is_board_test_detail(detail) else 'copper'
        for name in ['W_IN [W]', 'W_OUT [W]', 'Plateau Duration [s]', 'Sample Count', 'T_IN [°C]',
                     'T_OUT [°C]', 'Flow Rate', 'T_CU_AVG [°C]', 'DeltaT_CU [K]', 'Rth [K/W]', 'Subcooling [K]']:
            row[name] = np.nan
        rows.append(row)
    return pd.DataFrame(rows)


def steady_campaign_rows(frame):
    return frame.loc[frame['Test Mode'].ne('TR')].copy()


def transient_panels(detail, configuration):
    """Raw, synchronized values and instantaneous derived metrics; no averaging."""
    data, meta = detail['data'], detail['metadata']
    nan = pd.Series(np.nan, index=data.index, dtype=float)
    def values(names):
        column = find_optional_column(data, names)
        return nan.copy() if column is None else pd.to_numeric(data[column], errors='coerce')
    panels = []
    def add(title, unit, series):
        series = {name: value for name, value in series.items() if np.isfinite(np.asarray(value, dtype=float)).any()}
        if series:
            panels.append((title, unit, series))
    sensors = detail.get('t_cu_columns', [])
    label = 'T_CPU' if is_board_test_detail(detail) else 'T_CU'
    temps = data[sensors].apply(pd.to_numeric, errors='coerce') if sensors else pd.DataFrame(index=data.index)
    add('Individual temperatures', 'Temperature [°C]', {c.replace('T_BOARD', 'T_CPU'): temps[c] for c in sensors})
    medium = meta['medium'].upper()
    inlet = values([f'T_{medium}_IN', f'T_{medium}_INLET'])
    outlet = values([f'T_{medium}_OUT', f'T_{medium}_OUTLET'])
    avg, low, high = temps.mean(axis=1), temps.min(axis=1), temps.max(axis=1)
    add('Temperature spread', 'Temperature [°C]', {label+'_AVG': avg, label+'_MIN': low, label+'_MAX': high})
    add('Temperature rise above coolant inlet', 'ΔT [K]', {'Δ'+label: avg-inlet, 'Δ'+label+'_MIN': low-inlet, 'Δ'+label+'_MAX': high-inlet})
    add('Coolant temperatures', 'Temperature [°C]', {f'T_{medium}_IN': inlet, f'T_{medium}_OUT': outlet})
    flow = values(['VFR_WATER', 'VFR', 'VFR_WATER_IN'] if medium == 'WATER' else
                  ['CFM', 'CFM_AIR', 'AIR_CFM', 'VFR_AIR', 'VFR', 'VFR_AIR_IN'])
    add('Coolant flow', meta['flow_unit'], {meta['flow_type']: flow})
    cpu_power_columns = shift2dc_cpu_power_columns(data) if is_board_test_detail(detail) else {}
    if cpu_power_columns:
        cpu_powers = data[list(cpu_power_columns.values())].apply(pd.to_numeric, errors='coerce')
        power = cpu_powers.sum(axis=1, min_count=len(cpu_power_columns))
    else:
        power = values(['W_PSU', 'W_HEATER', 'Total_Heat_Load_W'])
    volume = flow * cfm_to_m3_s if meta['flow_type'] == 'CFM' else flow/60000
    removal = ((water_density*water_cp if medium == 'WATER' else air_density*air_cp) * volume * (outlet-inlet))
    add('Heat load and coolant heat removal', 'Heat load [W]', {'Total heat load': power, 'W_OUT': removal})
    add('Instantaneous thermal resistance', 'Rth [K/W]', {'Rth': (avg-inlet)/power.where(power > 0)})
    if meta['part_type'] == 'LTS':
        refrigerant = {name: values([name]) for name in PH_STATE_NAMES}
        add('Refrigerant temperatures', 'Temperature [°C]', refrigerant)
        add('Subcooling (T_COND_IN - T_COND_OUT)', 'Subcooling [K]',
            {'Subcooling': refrigerant['T_COND_IN']-refrigerant['T_COND_OUT']})
        pressures = {c: pressure_values_pa(pd.to_numeric(data[c], errors='coerce'),
                     pressure_column_settings(c, configuration))/1e5 for c in lts_ph_pressure_columns(detail)}
        add('Measured Psat', 'Absolute pressure [bar]', pressures)
    else:
        add('Adiabatic temperature', 'Temperature [°C]', {'T_ADIA': values(['T_ADIA'])})
    psu = [c for c in data if re.fullmatch(r'T_PSU(?:_?\d+)?(?:\s*\[.*\])?', str(c), re.I)]
    add('Electronics PSU temperatures', 'Temperature [°C]', {c: pd.to_numeric(data[c], errors='coerce') for c in psu})
    return panels


def draw_transient_pages(pdf, detail, page_number, configuration):
    panels = transient_panels(detail, configuration)
    x, xlabel = raw_t_cu_x_values(detail['data'])
    for start in range(0, max(1, len(panels)), 4):
        pdf.showPage()
        page_number += 1
        width, height = landscape(A4)
        pdf.setFillColor(report_dark)
        pdf.setFont(report_bold_font, 16)
        pdf.drawString(32, height-32, f"Transient test - {detail['metadata']['part_type']}")
        draw_wrapped_text(pdf, detail['source_file'], 32, height-49, width-64,
                          font_size=8, leading=10, maximum_lines=2)
        pdf.setFont(report_regular_font, 8)
        pdf.drawString(32, height-79, 'Raw time series. No steady-state averaging. Rth is instantaneous; missing power is not inferred.')
        figure, axes = plt.subplots(2, 2, figsize=(11.2, 6.3), squeeze=False)
        group = panels[start:start+4]
        for axis, (title, unit, series) in zip(axes.flat, group):
            for name, y in series.items():
                bound = name.endswith(('_MIN', '_MAX'))
                axis.plot(x, y, linewidth=.8, linestyle='--' if bound else '-',
                          color='#2468A2' if title in ('Temperature spread', 'Temperature rise above coolant inlet') else None,
                          label=name)
            axis.set_title(title, fontsize=10)
            style_report_axis(axis, xlabel, unit)
            axis.tick_params(labelsize=8)
            axis.legend(fontsize=6.5, frameon=False, ncol=2 if len(series)>4 else 1)
        for axis in list(axes.flat)[len(group):]:
            axis.axis('off')
        figure.tight_layout(pad=1.8)
        buffer = BytesIO()
        figure.savefig(buffer, format='png', dpi=160, facecolor='white')
        plt.close(figure)
        buffer.seek(0)
        pdf.drawImage(ImageReader(buffer), 25, 34, width=width-50, height=height-126,
                      preserveAspectRatio=True, anchor='c')
        buffer.close()
        draw_page_footer(pdf, page_number, width)
    return page_number


def configure_report(all_results, part_type, test_details, report_configuration=None):
    """Use one detection and setup path for heater and scheduled server tests."""
    campaign = report_campaign_frame(all_results, test_details)
    frame = steady_campaign_rows(campaign)
    primary = primary_report_results(frame)
    ratios = primary.loc[primary['Condition Type'].eq('FR'), 'Condition Value'].dropna().unique()
    _, repeated = find_repeated_test_groups(frame)
    performance = find_performance_characterization_groups(primary)
    fluids = set(campaign['Working Fluid'].dropna())
    fluid_properties = {fluid: get_refprop_fluid_properties(fluid) for fluid in sorted(fluids)}
    if report_configuration is None:
        report_configuration = show_report_configuration_dialog(
            part_type=part_type,
            fluids=fluids,
            coolants=set(campaign['Coolant'].dropna()),
            test_details=test_details,
            fluid_properties=fluid_properties,
            has_filling_ratio_analysis=len(ratios) > 1,
            repeated_group_count=len(repeated),
            performance_characterization_ratios=[ratio for ratio, _ in performance],
            all_results=all_results,
        )
    return report_configuration, fluid_properties


def make_shift2dc_raw_data_chart_image(test_detail):
    """Board traces use the shared per-test page, with schedule windows, no maps."""
    data = test_detail['data']
    figure, axis = plt.subplots(figsize=(11.2, 6.35))
    time = shift2dc_time(data)
    for column in test_detail['boards']:
        axis.plot(time, pd.to_numeric(data[column], errors='coerce'), linewidth=.9, label=column)
    schedule = test_detail['schedule']
    for row in schedule:
        axis.axvline(row['Start [s]'], color='#666666', linestyle='--', linewidth=.8)
        axis.axvspan(row['Average from [s]'], row['End [s]'], color='#6DAD82', alpha=.22)
        axis.text((row['Start [s]'] + row['End [s]']) / 2, .98,
                  shift2dc_step_power_label(row), ha='center', va='top',
                  transform=axis.get_xaxis_transform(), fontsize=8)
    axis.axvline(schedule[-1]['End [s]'], color='#666666', linestyle='--', linewidth=.8)
    axis.set_title(
        f"Raw T_CPU data | {len(test_detail['boards'])} connected boards | "
        f"Start: {schedule[0]['Start [s]']:g} s\n"
        "Dashed lines: scheduled steps | Green bands: final 100 s averaged",
        fontsize=10, pad=10,
    )
    style_report_axis(axis, 'Elapsed time [s]', 'T_CPU [°C]')
    axis.legend(fontsize=8, frameon=False, ncol=4, loc='lower right')
    figure.tight_layout(pad=1.2)
    buffer = BytesIO()
    figure.savefig(buffer, format='png', dpi=190, facecolor='white', bbox_inches='tight')
    plt.close(figure)
    buffer.seek(0)
    return buffer


def shift2dc_extra_columns(detail, kind):
    """Identify auxiliary sensors without confusing PSU power with temperature."""
    if not is_board_test_detail(detail):
        return []
    pattern = (r'^(?:P_?SAT|P_SATURATION)(?:_?\d+)?$' if kind == 'psat'
               else r'^T_PSU(?:_?\d+)?$')
    return [column for column in detail['data'].columns
            if re.fullmatch(pattern, re.split(r'[\[(]', str(column))[0].strip(), re.I)]


PRESSURE_UNIT_PA = {'Pa': 1.0, 'hPa': 100.0, 'kPa': 1000.0,
                    'MPa': 1e6, 'bar': 1e5, 'mbar': 100.0, 'psi': 6894.757293168}


def pressure_header_settings(column):
    """Use explicit header settings when present; otherwise default to bar absolute."""
    text = str(column).lower()
    match = re.search(r'(?<![a-z])(mpa|kpa|hpa|mbar|bar|pa|psi)(?:\s*\(?([ag])\)?)?(?![a-z])', text)
    unit = next((u for u in PRESSURE_UNIT_PA if match and u.lower() == match[1]), None)
    basis = ('Gauge' if re.search(r'gauge|relative', text) else
             'Absolute' if re.search(r'absolute|\babs\b', text) else None)
    if basis is None and match and match[2]:
        basis = 'Gauge' if match[2] == 'g' else 'Absolute'
    return {'unit': unit or 'bar', 'basis': basis or 'Absolute', 'atmospheric_pressure_pa': 101325.0}


def pressure_column_settings(column, configuration):
    settings = pressure_header_settings(column)
    settings.update(configuration.get('pressure_settings', {}).get(column, {}))
    if settings.get('unit') not in PRESSURE_UNIT_PA or settings.get('basis') not in ('Absolute', 'Gauge'):
        raise ValueError(f'{column}: specify pressure units and Absolute or Gauge in report setup.')
    atmosphere = float(settings.get('atmospheric_pressure_pa', 101325.0))
    if not math.isfinite(atmosphere) or atmosphere <= 0:
        raise ValueError('Atmospheric pressure must be a positive finite value.')
    settings['atmospheric_pressure_pa'] = atmosphere
    return settings


def pressure_values_pa(values, settings):
    pressures = pd.to_numeric(pd.Series(values), errors='coerce').to_numpy(dtype=float, copy=True)
    pressures *= PRESSURE_UNIT_PA[settings['unit']]
    if settings['basis'] == 'Gauge':
        pressures += settings['atmospheric_pressure_pa']
    return pressures


def show_pressure_settings_dialog(parent, columns, existing):
    """One explicit configuration per detected header; preserve differing sensors."""
    dialog = tk.Toplevel(parent)
    dialog.title('Psat pressure settings')
    dialog.transient(parent)
    frame = ttk.Frame(dialog, padding=14)
    frame.pack(fill='both', expand=True)
    ttk.Label(frame, text='Default: bar, absolute pressure. Adjust these settings if needed.').grid(
        row=0, column=0, columnspan=4, sticky='w', pady=(0, 10))
    for index, title in enumerate(['Sensor column', 'Units', 'Reference', 'Atmosphere [kPa]']):
        ttk.Label(frame, text=title).grid(row=1, column=index, padx=7, sticky='w')
    variables = {}
    for row, column in enumerate(columns, 2):
        settings = pressure_header_settings(column)
        settings.update(existing.get(column, {}))
        unit = tk.StringVar(value=settings.get('unit') or 'Select...')
        basis = tk.StringVar(value=settings.get('basis') or 'Select...')
        atmosphere = tk.StringVar(value=f"{settings.get('atmospheric_pressure_pa', 101325)/1000:g}")
        variables[column] = (unit, basis, atmosphere)
        ttk.Label(frame, text=column).grid(row=row, column=0, padx=7, pady=4, sticky='w')
        ttk.Combobox(frame, textvariable=unit, values=list(PRESSURE_UNIT_PA), state='readonly', width=9).grid(row=row, column=1, padx=7)
        ttk.Combobox(frame, textvariable=basis, values=['Absolute', 'Gauge'], state='readonly', width=12).grid(row=row, column=2, padx=7)
        ttk.Entry(frame, textvariable=atmosphere, width=12).grid(row=row, column=3, padx=7)
    ttk.Label(frame, text='Atmospheric pressure is added only to gauge readings; use the local value if known.').grid(
        row=len(columns)+2, column=0, columnspan=4, sticky='w', pady=10)
    answer = {'settings': None}
    def accept():
        try:
            result = {column: {'unit': unit.get(), 'basis': basis.get(),
                              'atmospheric_pressure_pa': float(atmosphere.get().replace(',', '.'))*1000}
                      for column, (unit, basis, atmosphere) in variables.items()}
            for column in columns:
                pressure_column_settings(column, {'pressure_settings': result})
        except (TypeError, ValueError) as error:
            messagebox.showerror('Pressure settings', str(error), parent=dialog)
            return
        answer['settings'] = result
        dialog.destroy()
    buttons = ttk.Frame(frame)
    buttons.grid(row=len(columns)+3, column=0, columnspan=4, sticky='e')
    ttk.Button(buttons, text='Cancel', command=dialog.destroy).pack(side='left', padx=7)
    ttk.Button(buttons, text='Use settings', command=accept).pack(side='left')
    dialog.grab_set()
    parent.wait_window(dialog)
    return answer['settings']


def saturation_temperature_calculator(fluid):
    """Return a P_absolute[Pa] -> saturated-vapor T[°C] function and backend name.

    Q=1 is appropriate for comparing with the evaporator vapor outlet; it is
    the dew temperature for a supported refrigerant blend. No fluid substitution.
    CoolProp: https://coolprop.org/coolprop/HighLevelAPI.html
    REFPROP: https://refprop-docs.readthedocs.io/en/latest/DLL/high_level.html
    """
    aliases = {'R1336MZZE': 'R1336mzz(E)', 'R1336MZZZ': 'R1336mzz(Z)',
               'R1233ZDE': 'R1233zd(E)'}
    compact = re.sub(r'[^A-Za-z0-9]', '', str(fluid)).upper()
    candidates = list(dict.fromkeys([aliases.get(compact, str(fluid))] + refprop_fluid_candidates(fluid)))
    try:
        from CoolProp.CoolProp import PropsSI
        for backend in ['HEOS', 'REFPROP']:
            for candidate in candidates:
                name = backend + '::' + candidate
                try:
                    PropsSI('Tcrit', name)  # Check fluid support independently of measured pressure.
                    def calculate(pressure, name=name):
                        return float(PropsSI('T', 'P', float(pressure), 'Q', 1, name)) - 273.15
                    return calculate, 'CoolProp/' + name
                except Exception:
                    continue
    except ImportError:
        pass
    try:
        from ctREFPROP.ctREFPROP import REFPROPFunctionLibrary
        for folder in refprop_installation_folders() or ['']:
            try:
                library = REFPROPFunctionLibrary(folder)
                if folder:
                    library.SETPATHdll(folder)
                units = library.GETENUMdll(0, 'DEFAULT').iEnum  # P in kPa, T in kelvin.
                for candidate in refprop_fluid_candidates(fluid):
                    result = library.REFPROPdll(candidate, 'TQ', 'P', units, 1, 0, 300.0, 1.0, [1.0])
                    if result.ierr > 0:
                        continue
                    def calculate(pressure, library=library, candidate=candidate, units=units):
                        result = library.REFPROPdll(candidate, 'PQ', 'T', units, 1, 0,
                                                   float(pressure)/1000, 1.0, [1.0])
                        if result.ierr > 0:
                            raise ValueError(refprop_result_error(result))
                        return float(result.Output[0]) - 273.15
                    return calculate, 'REFPROP/' + candidate
            except Exception:
                continue
    except ImportError:
        pass
    raise RuntimeError(f'T_sat unavailable for {fluid}. Install CoolProp, or ctREFPROP with NIST REFPROP and this fluid.')


def saturation_temperatures(pressures_pa, fluid):
    temperatures = np.full(len(pressures_pa), np.nan)
    try:
        calculate, source = saturation_temperature_calculator(fluid)
    except RuntimeError as error:
        return temperatures, str(error)
    cache = {}
    for index, pressure in enumerate(pressures_pa):
        if not np.isfinite(pressure) or pressure <= 0:
            continue
        key = float(pressure)
        if key not in cache:
            try:
                value = calculate(key)
                cache[key] = value if math.isfinite(value) else np.nan
            except Exception:
                cache[key] = np.nan
        temperatures[index] = cache[key]
    invalid = int(np.count_nonzero(~np.isfinite(temperatures)))
    note = source + '; saturated vapor (Q=1)'
    if invalid:
        note += f'; {invalid}/{len(temperatures)} invalid/out-of-range samples omitted'
    return temperatures, note


def shift2dc_auxiliary_averages(detail, values):
    """Use the same final 100 s windows, with actual elapsed-time weighting."""
    time = shift2dc_time(detail['data'])
    values = np.asarray(values, dtype=float)
    means = []
    for interval in detail['schedule']:
        start, end = interval['Average from [s]'], interval['Average to [s]']
        positions = np.flatnonzero((time >= start) & (time < end))
        previous = int(np.searchsorted(time, start, side='right')-1)
        if len(positions) and time[positions[0]] > start and previous >= 0:
            positions = np.r_[previous, positions]
        if not len(positions):
            means.append(np.nan)
            continue
        weights = np.r_[time[positions[1:]], end] - np.maximum(time[positions], start)
        valid = np.isfinite(values[positions]).all() and np.all(weights > 0)
        means.append(float(np.average(values[positions], weights=weights)) if valid else np.nan)
    return np.asarray(means)


def make_shift2dc_extra_chart(detail, kind, configuration, pressure_column=None):
    data = detail['data']
    time = shift2dc_time(data)
    power = np.array([row['Scheduled total power [W]'] for row in detail['schedule']])
    order = np.argsort(power, kind='stable')
    notes = []
    if kind == 'psat':
        figure, axes = plt.subplots(2, 2, figsize=(11.2, 6.1))
        pressure_axis, temperature_axis, average_axis, difference_axis = axes.flat
        settings = pressure_column_settings(pressure_column, configuration)
        pressure = pressure_values_pa(data[pressure_column], settings)
        saturation, source = saturation_temperatures(pressure, detail['metadata']['fluid'])
        notes.append(source)
        notes.append(f"Input: {pressure_column}; {settings['unit']} {settings['basis'].lower()}"
                     + (f"; atmosphere {settings['atmospheric_pressure_pa']/1000:g} kPa" if settings['basis']=='Gauge' else ''))
        pressure_axis.plot(time, pressure/1e5, color='#246080', linewidth=1)
        pressure_axis.set(title='Measured saturation pressure', ylabel='Psat [bar absolute]')
        temperature_axis.plot(time, saturation, label='T_sat (from Psat)', color='#244E78')
        mean_sat = shift2dc_auxiliary_averages(detail, saturation)
        average_axis.plot(power[order], mean_sat[order], 'o-', label='T_sat', color='#244E78')
        evap_column = next((column for column in data.columns
                            if re.split(r'[\[(]', str(column))[0].strip().upper() == 'T_EVAP_OUT'), None)
        if evap_column is not None:
            evap = pd.to_numeric(data[evap_column], errors='coerce').to_numpy(dtype=float)
            temperature_axis.plot(time, evap, label='T_EVAP_OUT', color='#C86828', linewidth=1)
            mean_evap = shift2dc_auxiliary_averages(detail, evap)
            average_axis.plot(power[order], mean_evap[order], 's-', label='T_EVAP_OUT', color='#C86828')
            difference = shift2dc_auxiliary_averages(detail, evap-saturation)
            difference_axis.plot(power[order], difference[order], 'o-', color='#72518A')
        else:
            notes.append('T_EVAP_OUT missing: comparison unavailable.')
        difference_axis.axhline(0, color='#333333', linestyle='--', linewidth=1)
        difference_axis.set(title='Outlet deviation', ylabel='T_EVAP_OUT − T_sat [K]')
        temperature_axis.set(title='Evaporator outlet and saturation temperature', ylabel='Temperature [°C]')
        average_axis.set(title='Plateau averages', ylabel='Temperature [°C]')
        temperature_axis.legend(fontsize=8)
        average_axis.legend(fontsize=8)
        time_axes, heat_axes = [pressure_axis, temperature_axis], [average_axis, difference_axis]
    else:
        figure, axes = plt.subplots(1, 2, figsize=(11.2, 6.1))
        for column in shift2dc_extra_columns(detail, 'psu'):
            values = pd.to_numeric(data[column], errors='coerce').to_numpy(dtype=float)
            line, = axes[0].plot(time, values, label=column, linewidth=1)
            means = shift2dc_auxiliary_averages(detail, values)
            axes[1].plot(power[order], means[order], 'o-', label=column, color=line.get_color())
        axes[0].set(title='Electronics PSU temperatures', ylabel='Temperature [°C]')
        axes[1].set(title='Plateau averages', ylabel='Temperature [°C]')
        for axis in axes:
            axis.legend(fontsize=8)
        time_axes, heat_axes = [axes[0]], [axes[1]]
    for axis in time_axes:
        axis.set_xlabel('Relative time [s]')
        for interval in detail['schedule']:
            axis.axvspan(interval['Average from [s]'], interval['Average to [s]'], color='#58A070', alpha=.12)
    for axis in heat_axes:
        axis.set_xlabel('Total heat load [W]')
    for axis in np.asarray(axes).flat:
        axis.grid(True, alpha=.25)
        axis.tick_params(labelsize=8)
    figure.tight_layout(pad=2)
    buffer = BytesIO()
    figure.savefig(buffer, format='png', dpi=180, facecolor='white')
    plt.close(figure)
    buffer.seek(0)
    return buffer, notes


def draw_shift2dc_extra_page(pdf, detail, page_number, kind, configuration, pressure_column=None):
    page_width, page_height = landscape(A4)
    title = ('Saturation pressure and T_sat' if kind == 'psat' else 'Electronics PSU temperatures')
    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, page_height-32, title)
    pdf.setFillColor(report_grey)
    draw_wrapped_text(pdf, f"CSV: {detail['source_file']}  |  Fluid: {detail['metadata']['fluid']}",
                      32, page_height-50, page_width-64, font_size=7.7, leading=9, maximum_lines=2)
    pdf.setStrokeColor(report_light_grey)
    pdf.line(32, page_height-72, page_width-32, page_height-72)
    chart, notes = make_shift2dc_extra_chart(detail, kind, configuration, pressure_column)
    pdf.drawImage(ImageReader(chart), 26, 65, width=page_width-52, height=page_height-143,
                  preserveAspectRatio=True, anchor='c')
    chart.close()
    if notes:
        draw_wrapped_text(pdf, '\n'.join(notes), 32, 56, page_width-64,
                          font_size=7.1, leading=9, maximum_lines=3)
    draw_page_footer(pdf, page_number, page_width)


PH_STATE_NAMES = ('T_EVAP_IN', 'T_EVAP_OUT', 'T_COND_IN', 'T_COND_OUT')


def lts_ph_pressure_columns(detail):
    """Psat detection shared by conventional LTS and Shift2DC p-h diagrams."""
    return [column for column in detail['data'].columns
            if re.fullmatch(r'(?:P_?SAT|P_SATURATION)(?:_?\d+)?',
                            re.split(r'[\[(]', str(column))[0].strip(), re.I)]


def lts_ph_temperature_column(data, name):
    return next((column for column in data.columns
                 if re.split(r'[\[(]', str(column))[0].strip().upper() == name), None)


def ph_property_backend(fluid):
    """Return SI properties from one consistent backend/reference for all states.

    CoolProp PropsSI and REFPROP MASS BASE SI both use Pa, K and J/kg.
    https://coolprop.org/coolprop/HighLevelAPI.html
    https://refprop-docs.readthedocs.io/en/latest/DLL/high_level.html
    """
    aliases = {'R1336MZZE': 'R1336mzz(E)', 'R1336MZZZ': 'R1336mzz(Z)',
               'R1233ZDE': 'R1233zd(E)'}
    token = re.sub(r'[^A-Za-z0-9]', '', str(fluid)).upper()
    candidates = list(dict.fromkeys([aliases.get(token, str(fluid))] + refprop_fluid_candidates(fluid)))
    try:
        from CoolProp.CoolProp import PropsSI
        for backend in ('HEOS', 'REFPROP'):
            for candidate in candidates:
                name = backend + '::' + candidate
                try:
                    PropsSI('Tcrit', name)
                    def prop(output, pair='', a=0, b=0, name=name):
                        value = (PropsSI(output, pair[0], float(a), pair[1], float(b), name)
                                 if pair else PropsSI(output, name))
                        if not math.isfinite(value):
                            raise ValueError('Non-finite property')
                        return float(value)
                    return prop, 'CoolProp/' + name
                except Exception:
                    continue
    except ImportError:
        pass
    try:
        from ctREFPROP.ctREFPROP import REFPROPFunctionLibrary
        for folder in refprop_installation_folders() or ['']:
            try:
                library = REFPROPFunctionLibrary(folder)
                if folder:
                    library.SETPATHdll(folder)
                enum = library.GETENUMdll(0, 'MASS BASE SI')
                if getattr(enum, 'ierr', 0) > 0:
                    continue
                for candidate in refprop_fluid_candidates(fluid):
                    def prop(output, pair='', a=0, b=0, library=library,
                             candidate=candidate, units=enum.iEnum):
                        output = {'Tcrit': 'TC', 'Tmin': 'TTRP'}.get(output, output)
                        result = library.REFPROPdll(candidate, pair, output, units, 1, 0,
                                                   float(a), float(b), [1.0])
                        value = float(result.Output[0])
                        if result.ierr > 0 or not math.isfinite(value) or value < -1e6:
                            raise ValueError(refprop_result_error(result))
                        return value
                    try:
                        prop('Tcrit')
                        return prop, 'REFPROP/' + candidate
                    except Exception:
                        continue
            except Exception:
                continue
    except ImportError:
        pass
    raise RuntimeError(f'No thermodynamic backend for {fluid}. Install CoolProp, or ctREFPROP with NIST REFPROP.')


def ph_plateau_means(detail, column):
    if column is None:
        return np.full(len(detail['steps']), np.nan)
    if detail.get('schedule'):
        return shift2dc_auxiliary_averages(detail, pd.to_numeric(detail['data'][column], errors='coerce'))
    # Conventional tests already contain the selected final averaging windows.
    return np.array([pd.to_numeric(step[column], errors='coerce').mean() for step in detail['steps']])


def build_lts_ph_data(detail, configuration, pressure_column=None):
    """Reconstruct equilibrium states from plateau averages; preserve uncertainty."""
    prop, backend = ph_property_backend(detail['metadata']['fluid'])
    temps = {name: ph_plateau_means(detail, lts_ph_temperature_column(detail['data'], name))
             for name in PH_STATE_NAMES}
    if pressure_column is not None:
        settings = pressure_column_settings(pressure_column, configuration)
        pressures = pressure_values_pa(ph_plateau_means(detail, pressure_column), settings)
        pressure_note = f"Measured {pressure_column} ({settings['unit']}, {settings['basis'].lower()})"
        if settings['basis'] == 'Gauge':
            pressure_note += f"; atmosphere {settings['atmospheric_pressure_pa']/1000:g} kPa"
    elif configuration.get('ph_allow_estimated_pressure', False):
        pressures = []
        for temp in temps['T_EVAP_OUT']:
            try:
                pressures.append(prop('P', 'TQ', temp+273.15, 1) if np.isfinite(temp) else np.nan)
            except Exception:
                pressures.append(np.nan)
        pressures = np.array(pressures)
        pressure_note = 'Estimated pressure: T_EVAP_OUT assumed saturated vapor (Q=1)'
    else:
        raise ValueError('No Psat column. Enable the saturated-vapor pressure estimate in report setup.')
    shift_model = is_shift2dc_file(Path(detail['source_file']))
    tolerance = float(configuration.get('ph_saturation_tolerance_k', .2))
    if not math.isfinite(tolerance) or tolerance < 0:
        raise ValueError('Saturation tolerance must be finite and nonnegative.')
    records = []
    for index, step in enumerate(detail['steps']):
        measured_pressure = float(pressures[index])
        head = 0.0
        density = np.nan
        if shift_model and pressure_column is not None:
            # Explicit model: a 0.19 m liquid column on either side of the tap.
            # Saturated-liquid density at the measured absolute pressure.
            density = prop('D', 'PQ', measured_pressure, 0)
            if not np.isfinite(density) or density <= 0:
                raise ValueError('Liquid density unavailable for hydrostatic correction.')
            head = density * 9.80665 * 0.19
            if measured_pressure - head <= 0:
                raise ValueError('Hydrostatic correction gives nonpositive condenser pressure.')
        states = []
        for state_index, name in enumerate(PH_STATE_NAMES):
            pressure = measured_pressure + (head if state_index < 2 else -head)
            temperature = float(temps[name][index])
            state = dict(name=name, temperature=temperature, pressure=pressure, enthalpy=np.nan, note='',
                         phase='unknown', saturation_liquid_temperature=np.nan, subcooling=np.nan)
            try:
                if not np.isfinite(pressure) or pressure <= 0 or not np.isfinite(temperature):
                    raise ValueError('missing/invalid temperature or pressure')
                bubble = prop('T', 'PQ', pressure, 0) - 273.15
                dew = prop('T', 'PQ', pressure, 1) - 273.15
                state['saturation_liquid_temperature'] = bubble
                if state_index in (0, 3):
                    state['subcooling'] = bubble - temperature
                if shift_model and state_index in (1, 2):
                    state['enthalpy'] = prop('H', 'PQ', pressure, .5)/1000
                    state['quality'] = .5
                    state['model_temperature'] = prop('T', 'PQ', pressure, .5)-273.15
                    state['note'] = 'assumed Q=0.5'
                    state['phase'] = 'two-phase estimate'
                elif min(bubble, dew)-tolerance <= temperature <= max(bubble, dew)+tolerance:
                    quality = 0 if state_index in (0, 3) else 1
                    boundary = bubble if quality == 0 else dew
                    if not configuration.get('ph_assume_saturated_endpoints', True):
                        raise ValueError('quality unknown at saturation')
                    if abs(temperature-boundary) > tolerance:
                        raise ValueError('within blend glide: quality unknown')
                    state['enthalpy'] = prop('H', 'PQ', pressure, quality)/1000
                    state['note'] = f'assumed Q={quality}'
                    state['phase'] = 'sat. liquid' if quality == 0 else 'sat. vapor'
                else:
                    state['enthalpy'] = prop('H', 'PT', pressure, temperature+273.15)/1000
                    state['note'] = 'from P,T'
                    state['phase'] = 'liquid' if temperature < min(bubble, dew) else 'vapor'
            except Exception as error:
                state['note'] = str(error).split(':')[0][:80]
            states.append(state)
        records.append(dict(power=numeric_average(step, get_power_column(step)), states=states,
                            measured_pressure=measured_pressure, liquid_density=density, hydrostatic_head_pa=head))
    # Saturation dome, always calculated with the same backend as the states.
    critical = prop('Tcrit')
    minimum = max(prop('Tmin')+.05, min([s['temperature']+273.15 for r in records for s in r['states']
                                       if np.isfinite(s['temperature'])] or [critical-100])-60)
    dome = []
    for temperature in np.linspace(minimum, critical-.05, 180):
        try:
            p_liq = prop('P', 'TQ', temperature, 0)/1e5
            p_vap = prop('P', 'TQ', temperature, 1)/1e5
            h_liq = prop('H', 'TQ', temperature, 0)/1000
            h_vap = prop('H', 'TQ', temperature, 1)/1000
            if p_liq > 0 and p_vap > 0:
                dome.append((h_liq, p_liq, h_vap, p_vap))
        except Exception:
            continue
    if len(dome) < 2:
        raise ValueError('Saturation dome unavailable for this refrigerant.')
    return records, np.array(dome), backend, pressure_note


def make_lts_ph_chart(records, dome):
    """Up to four plateau diagrams per page; arrows connect measured state locations."""
    columns = 1 if len(records) == 1 else 2
    rows = math.ceil(len(records)/columns)
    figure, axes = plt.subplots(rows, columns, figsize=(11.2, 6.1), squeeze=False)
    palette = ['#2468A2', '#C34444', '#C97822', '#238A80']
    offsets = [(5, 6), (5, 6), (5, -18), (5, -28)]
    for axis, record in zip(axes.flat, records):
        axis.plot(dome[:,0], dome[:,1], color='#687789', linewidth=1, label='Saturation dome')
        axis.plot(dome[:,2], dome[:,3], color='#687789', linewidth=1)
        axis.set_yscale('linear')
        states = record['states']
        # Shade at common pressures, not common temperatures (also handles glide).
        liquid_order = np.argsort(dome[:,1])
        vapor_order = np.argsort(dome[:,3])
        low_p = max(dome[:,1].min(), dome[:,3].min())
        high_p = min(dome[:,1].max(), dome[:,3].max())
        pressure_grid = np.geomspace(low_p, high_p, 250)
        liquid_h = np.interp(np.log(pressure_grid), np.log(dome[liquid_order,1]), dome[liquid_order,0])
        vapor_h = np.interp(np.log(pressure_grid), np.log(dome[vapor_order,3]), dome[vapor_order,2])
        state_h = [state['enthalpy'] for state in states if np.isfinite(state['enthalpy'])]
        left = min([dome[:,0].min(), dome[:,2].min()] + state_h)
        right = max([dome[:,0].max(), dome[:,2].max()] + state_h)
        margin = max(1, .06*(right-left))
        left, right = left-margin, right+margin
        axis.set_xlim(left, right)
        axis.fill_betweenx(pressure_grid, left, liquid_h, color='#D9EAF7', zorder=0)
        axis.fill_betweenx(pressure_grid, liquid_h, vapor_h, color='#ECE3F5', zorder=0)
        axis.fill_betweenx(pressure_grid, vapor_h, right, color='#FBE3CF', zorder=0)

        for index, state in enumerate(states):
            h, pressure = state['enthalpy'], state['pressure']/1e5
            if np.isfinite(h) and np.isfinite(pressure) and pressure > 0:
                axis.scatter([h], [pressure], color=palette[index], marker=['o','s','^','D'][index], s=28, zorder=4)
                label = str(index+1)
                if index in (0, 3) and np.isfinite(state.get('subcooling', np.nan)):
                    subcooling = state['subcooling']
                    label += (f"\nSC = {subcooling:.1f} K" if subcooling >= 0
                              else f"\nAbove sat. by {-subcooling:.1f} K")
                axis.annotate(label, (h, pressure), textcoords='offset points',
                              xytext=((-5, offsets[index][1]) if h > left+.65*(right-left) else offsets[index]),
                              ha='right' if h > left+.65*(right-left) else 'left',
                              color=palette[index], fontsize=7, fontweight='bold')
                following = states[(index+1)%4]
                h2 = following['enthalpy']
                if np.isfinite(h2) and (not math.isclose(h, h2, abs_tol=1e-7)
                                       or not math.isclose(state['pressure'], following['pressure'])):
                    axis.annotate('', xy=(h2, following['pressure']/1e5), xytext=(h, pressure),
                                  arrowprops=dict(arrowstyle='->', color=palette[index], lw=1.3,
                                                  linestyle='--' if index == 3 else '-'))
        axis.set_title(f"Total heat load: {record['power']:g} W | p evap/cond = "
                       f"{states[0]['pressure']/1e5:.3f}/{states[2]['pressure']/1e5:.3f} bar",
                       fontsize=9, fontweight='bold')
        axis.set_xlabel('Specific enthalpy h [kJ/kg]', fontsize=8)
        axis.set_ylabel('Absolute pressure p [bar]', fontsize=8)
        state_pressures = [st['pressure']/1e5 for st in states
                           if np.isfinite(st['pressure']) and st['pressure'] > 0]
        if state_pressures:
            low, high = min(state_pressures), max(state_pressures)
            span = max(high-low, .005*max(high, 1))
            axis.set_ylim(max(1e-6, low-span), high+1.4*span)
            from matplotlib.ticker import MaxNLocator, ScalarFormatter
            axis.yaxis.set_major_locator(MaxNLocator(nbins=5, steps=[1,2,2.5,5,10]))
            formatter = ScalarFormatter(useOffset=False)
            formatter.set_scientific(False)
            axis.yaxis.set_major_formatter(formatter)
        axis.tick_params(labelsize=7)
        axis.grid(True, which='both', alpha=.2)
        # State values stay visible even when coincident points overlap on the path.
        labels = []
        for index, state in enumerate(states):
            t = f"{state['temperature']:.1f}" if np.isfinite(state['temperature']) else 'N/A'
            h = f"{state['enthalpy']:.1f}" if np.isfinite(state['enthalpy']) else 'N/A'
            labels.append(f"{index+1}: measured {t}°C, h={h} | {state.get('phase', 'unknown')} ({state['note']})")
        axis.text(.02, .98, '\n'.join(labels), transform=axis.transAxes, va='top', fontsize=6.8,
                  bbox=dict(facecolor='white', edgecolor='none', alpha=.85))
    for axis in list(axes.flat)[len(records):]:
        axis.axis('off')
    from matplotlib.patches import Patch
    figure.legend(handles=[Patch(facecolor='#D9EAF7', label='Liquid (subcooled)'),
                           Patch(facecolor='#ECE3F5', label='Liquid + vapor'),
                           Patch(facecolor='#FBE3CF', label='Vapor (superheated)')],
                  loc='lower center', ncol=3, frameon=False, fontsize=8, bbox_to_anchor=(.5, 0))
    figure.tight_layout(pad=1.7, rect=(0, .045, 1, 1))
    buffer = BytesIO()
    figure.savefig(buffer, format='png', dpi=180, facecolor='white')
    plt.close(figure)
    buffer.seek(0)
    return buffer


def draw_lts_ph_page(pdf, detail, page_number, records, dome, backend, pressure_note,
                     configuration, error=None):
    width, height = landscape(A4)
    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, height-32, 'LTS pressure-enthalpy diagram')
    pdf.setFillColor(report_grey)
    draw_wrapped_text(pdf, f"{detail['metadata']['fluid']} | {detail['source_file']}",
                      32, height-49, width-64, font_size=7.5, leading=9, maximum_lines=2)
    pdf.setStrokeColor(report_light_grey)
    pdf.line(32, height-71, width-32, height-71)
    pdf.setFont(report_regular_font, 8)
    pdf.drawString(32, height-85, 'Flow order: 1 T_EVAP_IN → 2 T_EVAP_OUT → 3 T_COND_IN → 4 T_COND_OUT → 1')
    if error:
        draw_wrapped_text(pdf, 'p-h diagram unavailable: '+str(error), 40, height-130,
                          width-80, font_size=11, leading=15, maximum_lines=5)
    else:
        chart = make_lts_ph_chart(records, dome)
        pdf.drawImage(ImageReader(chart), 25, 91, width=width-50, height=height-180,
                      preserveAspectRatio=True, anchor='c')
        chart.close()
        tolerance = configuration.get('ph_saturation_tolerance_k', .2)
        notes = [backend+' | '+pressure_note,
                 'One common pressure per plateau; no local pressure drops reconstructed. Arrows join states, not measured intermediate paths.',
                 (f'Within {tolerance:g} K of saturation: Q=0 at 1/4 and Q=1 at 2/3 assumed; quality is not measured.'
                  if configuration.get('ph_assume_saturated_endpoints', True)
                  else 'Saturated states with unknown quality are omitted; gaps are not connected.'),
                 'SC at 1/4 = T_sat,liquid(P) - T_measured; expected liquid at 1/4 and vapor at 2/3. Actual states follow the data. '
                 'This local subcooling differs from the summary metric T_COND_IN - T_COND_OUT.']
        if is_shift2dc_file(Path(detail['source_file'])):
            notes = [backend+' | '+pressure_note,
                     ('p evap = Psat + rho_liq*g*0.19 m; p cond = Psat - rho_liq*g*0.19 m; '
                      'rho_liq = saturated-liquid density at measured Psat.' if records and
                      np.isfinite(records[0].get('liquid_density', np.nan)) else
                      'No measured Psat: estimated common pressure; elevation correction not applied.'),
                     'Approximation: Q=0.5 at 2/3 (liquid-vapor mixture), regardless of measured temperature; h is modeled, not measured.',
                     'SC at 1/4 = T_sat,liquid(local P) - T_measured. Summary subcooling remains T_COND_IN - T_COND_OUT.',
                     'Arrows join estimated states; frictional pressure losses are not modeled.']
        draw_wrapped_text(pdf, '\n'.join(notes), 32, 75, width-64, font_size=7.1, leading=10, maximum_lines=5)
    draw_page_footer(pdf, page_number, width)


def psat_result_columns(results):
    """Canonical absolute-pressure metrics, preserving distinct Psat sensors."""
    columns = results.columns if hasattr(results, 'columns') else results
    return [column for column in columns if re.fullmatch(r'Psat(?:_\d+)? \[bar abs\]', str(column))]


def tsat_result_columns(results):
    columns = results.columns if hasattr(results, 'columns') else results
    return [c for c in columns if re.fullmatch(r'T_SAT(?:_\d+)? \[°C\]', str(c))]


def superheating_result_columns(results):
    """One superheating metric for each independently measured saturation pressure."""
    columns = results.columns if hasattr(results, 'columns') else results
    return [c for c in columns if re.fullmatch(r'Superheating(?:_\d+)? \[K\]', str(c))]


def lts_saturation_averages(detail, column, settings):
    fluid = detail['metadata']['fluid']
    if detail.get('schedule'):
        pressure = pressure_values_pa(pd.to_numeric(detail['data'][column], errors='coerce'), settings)
        temperatures, note = saturation_temperatures(pressure, fluid)
        return shift2dc_auxiliary_averages(detail, temperatures), note
    means = []
    note = ''
    for step in detail['steps']:
        pressure = pressure_values_pa(pd.to_numeric(step[column], errors='coerce'), settings)
        temperatures, note = saturation_temperatures(pressure, fluid)
        means.append(float(np.mean(temperatures)) if len(temperatures) and np.isfinite(temperatures).all() else np.nan)
    return np.asarray(means), note


def add_lts_psat_results(all_results, test_details, configuration):
    """Attach measured plateau means by source and step order, never by heat load alone."""
    rows = [dict(row) for row in (all_results.to_dict('records') if isinstance(all_results, pd.DataFrame) else all_results)]
    for detail in test_details:
        if detail['metadata'].get('part_type') != 'LTS':
            continue
        columns = lts_ph_pressure_columns(detail)
        if not columns:
            continue
        target_rows = [row for row in rows if row['Source File'] == detail['source_file']]
        if len(target_rows) != len(detail['steps']):
            raise ValueError(f"{detail['source_file']}: pressure averages do not match the result rows.")
        if not target_rows:
            continue
        used = set()
        for column in columns:
            base = re.split(r'[\[(]', str(column))[0].strip()
            number = re.search(r'(\d+)$', base)
            metric = 'Psat' + ('_' + str(int(number[1])) if number else '') + ' [bar abs]'
            if metric in used:
                raise ValueError(f'{detail["source_file"]}: duplicate pressure sensor aliases for {metric}.')
            used.add(metric)
            settings = pressure_column_settings(column, configuration or {})
            pressure = pressure_values_pa(ph_plateau_means(detail, column), settings)/1e5
            temperatures, temperature_note = lts_saturation_averages(detail, column, settings)
            temperature_metric = metric.replace('Psat', 'T_SAT').replace('[bar abs]', '[°C]')
            if not np.isfinite(temperatures).all():
                print(f"{detail['source_file']}: some {temperature_metric} values unavailable: {temperature_note}")
            superheating_metric = metric.replace('Psat', 'Superheating').replace('[bar abs]', '[K]')
            for row, value, temperature in zip(target_rows, pressure, temperatures):
                row[metric] = float(value) if np.isfinite(value) and value > 0 else np.nan
                row[temperature_metric] = float(temperature) if np.isfinite(temperature) else np.nan
                evap_out = pd.to_numeric(row.get('T_EVAP_OUT [°C]', np.nan), errors='coerce')
                row[superheating_metric] = (float(evap_out - temperature)
                    if pd.notna(evap_out) and np.isfinite(evap_out) and np.isfinite(temperature) else np.nan)
    return rows


def psat_comparison_groups(results):
    """Compare only compatible assemblies, refrigerants and filling ratio/charge."""
    group_columns = [c for c in ['Evaporator Name', 'Condenser Name', 'Working Fluid',
                                 'Condition Type', 'Condition Value'] if c in results]
    for _, group in results.groupby(group_columns, sort=True, dropna=False):
        for metric in psat_result_columns(group):
            if pd.to_numeric(group[metric], errors='coerce').notna().any():
                yield group, metric


def make_psat_comparison_chart(results, metric):
    condition_columns = performance_condition_columns()
    conditions = results[condition_columns].drop_duplicates()
    variations = performance_condition_variations(conditions)
    flows = sorted({performance_flow_key(row['Flow Type'], row['Nominal Flow Rate'])
                    for _, row in conditions.iterrows()})
    figure, axis = plt.subplots(figsize=(11.2, 6.1))
    for index, (source, source_data) in enumerate(results.groupby('Source File', sort=True)):
        condition = source_data.iloc[0]
        key = tuple(condition[c] for c in condition_columns)
        label = performance_condition_legend_label(key, variations, index+1)
        if str(condition.get('Comment', '')).strip():
            label += ' | ' + str(condition['Comment']).strip()
        # Keep separate runs separate; add an identifier if their condition labels coincide.
        same_condition = results[condition_key_mask(results, condition_columns, key)]
        if same_condition['Source File'].nunique() > 1:
            label += f' | Run {index+1}'
        flow = performance_flow_key(key[2], key[3])
        color = adjust_performance_color_for_flow(performance_temperature_color(key[4]), flows.index(flow), len(flows))
        grouped, _ = add_heat_load_groups(source_data)
        series = grouped.groupby('_Heat Load Group', sort=True).agg({'W_IN [W]':'mean', metric:'mean'}).sort_values('W_IN [W]')
        plot_performance_series(axis, series, metric, color, performance_orientation_marker(key[5]), label)
    first = results.iloc[0]
    title = performance_common_condition_title(first['Condition Value'], conditions, variations)
    if first['Condition Type'] != 'FR':
        title = title.replace(f"FR {format_number(first['Condition Value'])}%", f"Charge {format_number(first['Condition Value'])}", 1)
    axis.set_title(title, fontsize=10, pad=10)
    style_report_axis(axis, 'Total heat load [W]', metric)
    axis.legend(fontsize=8, frameon=False, loc='best')
    figure.tight_layout(pad=1.8)
    buffer = BytesIO()
    figure.savefig(buffer, format='png', dpi=180, facecolor='white')
    plt.close(figure)
    buffer.seek(0)
    return buffer


def draw_psat_comparison_page(pdf, results, metric, page_number):
    width, height = landscape(A4)
    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, height-32, 'Psat comparison')
    first = results.iloc[0]
    subtitle = (f"EVAP: {first.get('Evaporator Name', '')} | COND: {first.get('Condenser Name', '')} | "
                'Plateau-average pressure, bar absolute')
    pdf.setFillColor(report_grey)
    draw_wrapped_text(pdf, subtitle, 32, height-50, width-64, font_size=8, leading=10, maximum_lines=2)
    pdf.setStrokeColor(report_light_grey)
    pdf.line(32, height-72, width-32, height-72)
    chart = make_psat_comparison_chart(results, metric)
    pdf.drawImage(ImageReader(chart), 26, 50, width=width-52, height=height-126,
                  preserveAspectRatio=True, anchor='c')
    chart.close()
    pdf.setFont(report_regular_font, 7)
    pdf.drawString(32, 39, 'Same averaging windows as the summary tables. Commented repeats are excluded when a reference exists.')
    draw_page_footer(pdf, page_number, width)


def superheating_subcooling_groups(results):
    """Keep assemblies/fluids separate and paginate long condition legends."""
    if results.empty:
        return
    keys = [c for c in ['Evaporator Name', 'Condenser Name', 'Working Fluid', 'Coolant', 'Condition Type'] if c in results]
    groups = (group for _, group in results.groupby(keys, sort=True, dropna=False)) if keys else [results]
    for group in groups:
        metrics = superheating_result_columns(group) or ['Superheating [K]']
        sources = list(group['Source File'].drop_duplicates())
        for metric in metrics:
            for start in range(0, len(sources), 6):
                yield group[group['Source File'].isin(sources[start:start+6])], metric


def make_superheating_subcooling_chart(results, metric, configuration):
    """Compare both temperature differences without combining distinct plateaus."""
    figure, axes = plt.subplots(1, 2, figsize=(11.2, 5.9))
    subcooling_limit, _, superheating_limit = summary_red_thresholds(configuration)
    colors = plt.get_cmap('tab10')
    handles, labels = [], []
    for index, (source, source_data) in enumerate(results.groupby('Source File', sort=True)):
        first = source_data.iloc[0]
        condition = (f"FR {format_number(first['Condition Value'])}%" if first['Condition Type'] == 'FR'
                     else f"Charge {format_number(first['Condition Value'])}")
        label = (f"{condition} | {first['Flow Type']} {format_number(first['Nominal Flow Rate'])}"
                 f" | {'T_AIR' if first['Flow Type'] == 'CFM' else 'TW'} {format_number(first['Nominal T_IN [°C]'])}°C")
        orientation = str(first.get('Orientation', '')).strip()
        if orientation and orientation.lower() not in ('nan', 'n/a', 'none'):
            label += f' | {orientation}'
        comment = str(first.get('Comment', '')).strip()
        if comment and comment.lower() != 'nan':
            label += ' | ' + comment
        if label in labels:
            label += f' | Run {index+1}'
        labels.append(label)
        series = source_data.sort_values('W_IN [W]', kind='stable')
        for axis, column in zip(axes, [metric, 'Subcooling [K]']):
            values = pd.to_numeric(series.get(column, pd.Series(np.nan, index=series.index)), errors='coerce')
            # Do not average identical total loads: heterogeneous CPU distributions may differ.
            line, = axis.plot(series['W_IN [W]'], values, marker='o', markersize=4,
                              linewidth=1.5, color=colors(index % 10), label=label)
            if axis is axes[0]:
                handles.append(line)
    for axis, column, title, limit in zip(axes, [metric, 'Subcooling [K]'],
            ['Superheating', 'Subcooling'], [superheating_limit, subcooling_limit]):
        axis.set_title(title, fontsize=12)
        style_report_axis(axis, 'Total heat load [W]', column)
        axis.axhline(limit, color='#A52424', linestyle='--', linewidth=1,
                     label=f'Red-cell threshold: {limit:g} K')
        axis.text(.02, .98, f'Threshold: {limit:g} K', transform=axis.transAxes,
                  ha='left', va='top', color='#A52424', fontsize=8,
                  bbox={'facecolor': 'white', 'edgecolor': 'none', 'alpha': .85, 'pad': 1.5})
        values = pd.to_numeric(results.get(column, pd.Series(dtype=float)), errors='coerce')
        if not np.isfinite(values).any():
            axis.text(.5, .5, 'Unavailable: measured Psat, refrigerant properties\nand T_EVAP_OUT are required.' if title == 'Superheating'
                      else 'Subcooling data unavailable.', transform=axis.transAxes,
                      ha='center', va='center', fontsize=9)
    figure.legend(handles, labels, loc='lower center', bbox_to_anchor=(.5, .015),
                  ncol=1, fontsize=8, frameon=False)
    figure.tight_layout(rect=(0, .10 + .033*len(labels), 1, 1), pad=1.5)
    buffer = BytesIO()
    figure.savefig(buffer, format='png', dpi=180, facecolor='white')
    plt.close(figure)
    buffer.seek(0)
    return buffer


def draw_superheating_subcooling_page(pdf, results, metric, page_number, configuration):
    width, height = landscape(A4)
    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, height-32, 'Superheating and subcooling')
    first = results.iloc[0]
    subtitle = (f"EVAP: {first.get('Evaporator Name', '')} | COND: {first.get('Condenser Name', '')} | "
                f"Fluid: {first.get('Working Fluid', '')} | {metric}")
    pdf.setFillColor(report_grey)
    draw_wrapped_text(pdf, subtitle, 32, height-50, width-64, font_size=8, leading=10, maximum_lines=2)
    pdf.setStrokeColor(report_light_grey)
    pdf.line(32, height-72, width-32, height-72)
    chart = make_superheating_subcooling_chart(results, metric, configuration)
    pdf.drawImage(ImageReader(chart), 26, 58, width=width-52, height=height-134,
                  preserveAspectRatio=True, anchor='c')
    chart.close()
    pdf.setFont(report_regular_font, 7)
    pdf.drawString(32, 45, 'Superheating = T_EVAP_OUT - T_SAT(Psat, refrigerant). Subcooling = T_COND_IN - T_COND_OUT.')
    pdf.drawString(32, 34, '1 K temperature difference = 1°C. Dashed lines show the selected summary red-cell thresholds.')
    draw_page_footer(pdf, page_number, width)


def summary_value_options(results, part_type, test_details=None):
    """Available numeric results and established defaults for the PDF table."""
    frame = pd.DataFrame(results) if results is not None else pd.DataFrame()
    board = is_board_report(frame) or any(is_board_test_detail(d) for d in test_details or [])
    catalog = ['DeltaT_CU [K]', 'Subcooling [K]', *psat_result_columns(frame), *tsat_result_columns(frame), *superheating_result_columns(frame),
               'W_OUT [W]', 'T_ADIA [°C]', 'T_CU_AVG [°C]', 'T_CU_MIN [°C]', 'T_CU_MAX [°C]',
               'DeltaT_CU_MIN [K]', 'DeltaT_CU_MAX [K]', 'Rth [K/W]',
               'T_EVAP_IN [°C]', 'T_EVAP_OUT [°C]', 'T_COND_IN [°C]', 'T_COND_OUT [°C]',
               'T_IN [°C]', 'T_OUT [°C]', 'Flow Rate', 'Water / scheduled power [%]']
    # Raw sensor channels can also be selected, without duplicating internal aliases.
    catalog += [c for c in frame.columns if re.fullmatch(r'T_(?:CPU|CU|ADIA)(?:_[A-Za-z0-9]+)* \[°C\]', str(c))
                and c not in catalog and (not board or re.fullmatch(r'T_CPU_\d+ \[°C\]', str(c)))]
    available = [c for c in dict.fromkeys(catalog) if c in frame and (pd.to_numeric(frame[c], errors='coerce').notna().any() or c in tsat_result_columns(frame) or c in superheating_result_columns(frame))]
    # Pressure averages are added after the user chooses the pressure settings.
    if part_type == 'LTS':
        for detail in test_details or []:
            for column in lts_ph_pressure_columns(detail):
                base = re.split(r'[\[(]', str(column))[0].strip()
                number = re.search(r'(\d+)$', base)
                metric = 'Psat' + ('_'+str(int(number[1])) if number else '') + ' [bar abs]'
                for result_metric in [metric, metric.replace('Psat', 'T_SAT').replace('[bar abs]', '[°C]'),
                                      metric.replace('Psat', 'Superheating').replace('[bar abs]', '[K]')]:
                    if result_metric not in available:
                        available.append(result_metric)
    defaults = ['DeltaT_CU [K]']
    if part_type == 'LTS':
        defaults += ['Subcooling [K]', *psat_result_columns(available)]
    elif 'T_ADIA [°C]' in available:
        defaults.append('T_ADIA [°C]')
    if board:
        defaults.append('W_OUT [W]')
    defaults = [c for c in defaults if c in available]
    return defaults + [c for c in available if c not in defaults], defaults


def summary_value_label(column, results):
    label = {'W_OUT [W]': 'W_OUT [W]',
             'Water / scheduled power [%]': 'W_OUT/W_IN [%]',
             'Flow Rate': 'Measured flow rate'}.get(column, column)
    return report_display_text(label.replace('DeltaT', 'ΔT'), results)


def row_summary_payload(results, selected_values):
    """Keep one row per source plateau with conditions alongside result columns."""
    frame = results.copy()
    fixed = []
    common = []
    for column, label in [('Evaporator Name','EVAP'), ('Condenser Name','COND'),
                          ('Working Fluid','Fluid'), ('Coolant','Coolant'), ('Orientation','Orientation')]:
        if column not in frame:
            continue
        values = frame[column].fillna('').astype(str)
        unique = values.unique()
        if len(unique) == 1:
            if unique[0] and unique[0].upper() != 'N/A':
                common.append(f'{label}: {unique[0]}')
        else:
            fixed.append((column, label))
    kinds = set(frame['Condition Type'].dropna())
    if len(kinds) > 1:
        fixed.append(('Condition Type', 'FR / charge'))
    fixed.append(('Condition Value', 'FR [%]' if kinds == {'FR'} else 'Charge' if kinds == {'Charge'} else 'FR / charge value'))
    flows = set(frame['Flow Type'].dropna())
    if len(flows) > 1:
        fixed.append(('Flow Type','Flow type'))
    fixed.append(('Nominal Flow Rate', 'VFR\n[l/min]' if flows == {'VFR'} else 'CFM' if flows == {'CFM'} else 'Flow\n[l/min or CFM]'))
    media = set(frame['Coolant'].dropna())
    fixed.append(('Nominal T_IN [°C]', 'T_AIR\n[°C]' if media == {'Air'} else 'TW\n[°C]' if media == {'Water'} else 'T_IN nominal\n[°C]'))
    fixed.append(('W_IN [W]', 'Total heat\nload [W]'))
    if 'Comment' in frame and frame['Comment'].fillna('').astype(str).str.strip().ne('').any():
        fixed.append(('Comment', 'Comment'))
    sort_columns = [c for c in ['Evaporator Name', 'Condenser Name', 'Working Fluid', 'Condition Type',
                    'Condition Value','Flow Type','Nominal Flow Rate','Nominal T_IN [°C]','Orientation',
                    'Source File','W_IN [W]'] if c in frame]
    frame = frame.sort_values(sort_columns, kind='stable', na_position='last')
    # Sequential IDs distinguish identical conditions and identify horizontal continuations.
    frame['_Summary Row'] = np.arange(1, len(frame)+1)
    fixed.insert(0, ('_Summary Row', 'Row'))
    metrics = [(c, summary_value_label(c, results).replace(' [', '\n[')) for c in selected_values]
    return frame, fixed, metrics, ' | '.join(common)


def summary_red_thresholds(configuration):
    subcooling = float(configuration.get('summary_subcooling_limit', 5.0))
    cpu = float(configuration.get('summary_cpu_temperature_limit', 100.0))
    superheating = float(configuration.get('summary_superheating_limit', 1.0))
    if not all(math.isfinite(v) for v in (subcooling, cpu, superheating)):
        raise ValueError('Summary thresholds must be finite numbers.')
    return subcooling, cpu, superheating


def summary_cell_exceeds_limit(column, value, results, configuration):
    if not isinstance(value, (int, float, np.number)) or not np.isfinite(value):
        return False
    subcooling, cpu, superheating = summary_red_thresholds(configuration)
    if column in superheating_result_columns([column]):
        return value > superheating
    if column == 'Subcooling [K]':
        return value > subcooling
    label = summary_value_label(column, results)
    return bool(re.fullmatch(r'T_CPU_(?:AVG|MIN|MAX|\d+) \[°C\]', label)) and value > cpu


def draw_condition_summary_pages(pdf, results, part_type, page_number, configuration=None):
    """Render rows of conditions with selectable value columns and readable pagination."""
    configuration = configuration or {}
    available, defaults = summary_value_options(results, part_type)
    chosen = configuration.get('summary_value_columns')
    selected = defaults if chosen is None else [c for c in dict.fromkeys(chosen) if c in available]
    frame, fixed, metrics, context = row_summary_payload(results, selected)
    subcooling_limit, cpu_limit, superheating_limit = summary_red_thresholds(configuration)
    group_columns = [c for c in ['Part Name', 'Evaporator Name', 'Condenser Name', 'Working Fluid',
        'Coolant', 'Condition Type', 'Condition Value', 'Flow Type', 'Nominal Flow Rate',
        'Nominal T_IN [°C]', 'Orientation'] if c in frame]
    group_keys = frame[group_columns].astype(object).where(frame[group_columns].notna(), '').apply(tuple, axis=1)
    frame['_New Condition Group'] = group_keys.ne(group_keys.shift())
    width, height = landscape(A4)
    # Split wide selections, repeating conditions and row IDs on every page.
    per_group = max(1, 11-len(fixed))
    metric_groups = [metrics[i:i+per_group] for i in range(0, len(metrics), per_group)] or [[]]
    rows_per_page = 27
    chunks = [frame.iloc[i:i+rows_per_page] for i in range(0,len(frame),rows_per_page)] or [frame]
    first_page = True
    for group_index, metric_group in enumerate(metric_groups):
        columns = fixed+metric_group
        weights = [0.55 if c == '_Summary Row' else 1.35 if c in {'Evaporator Name','Condenser Name','Working Fluid','Comment'} else 1.0 for c,_ in columns]
        widths = [(width-64)*weight/sum(weights) for weight in weights]
        for chunk in chunks:
            if not first_page:
                pdf.showPage()
                page_number += 1
            first_page = False
            pdf.setFillColor(report_dark)
            pdf.setFont(report_bold_font,16)
            pdf.drawString(32,height-32,'Test summary')
            pdf.setFillColor(report_grey)
            draw_wrapped_text(pdf, context,32,height-49,width-64,font_size=8,leading=10,maximum_lines=2)
            pdf.setFont(report_regular_font,7.5)
            row_range = f"Rows {int(chunk['_Summary Row'].iloc[0])}-{int(chunk['_Summary Row'].iloc[-1])} of {len(frame)}" if len(chunk) else 'No results'
            pdf.drawString(32,height-76,row_range + (f' | Values {group_index+1}/{len(metric_groups)}' if len(metric_groups)>1 else ''))
            top = height-88
            x = 32
            for index,((column,label),cell_width) in enumerate(zip(columns,widths)):
                fill = reportlab_colors.HexColor('#2F6B3B' if index<len(fixed) else '#365B73')
                draw_summary_cell(pdf,x,top,cell_width,29,label,fill,font_name=report_bold_font,
                                  font_size=7,text_color=reportlab_colors.white)
                x += cell_width
            for row_index,(_,row) in enumerate(chunk.iterrows()):
                y = top-29-row_index*15
                x = 32
                fill = reportlab_colors.HexColor('#F1F4F6') if row_index%2 else reportlab_colors.white
                for (column,_),cell_width in zip(columns,widths):
                    value = row.get(column, np.nan)
                    if pd.isna(value):
                        text = ''
                    elif column == '_Summary Row':
                        text = str(int(value))
                    elif isinstance(value,(int,float,np.number)):
                        text = f'{value:.3f}' if column == 'Rth [K/W]' else f'{value:.1f}'
                    else:
                        text = str(value)
                    alert = summary_cell_exceeds_limit(column, value, results, configuration)
                    cell_fill = reportlab_colors.HexColor('#F3A6A6') if alert else fill
                    draw_summary_cell(pdf,x,y,cell_width,15,text,cell_fill,font_size=7)
                    x += cell_width
                if row_index > 0 and row['_New Condition Group']:
                    pdf.setStrokeColor(reportlab_colors.HexColor('#344D43'))
                    pdf.setLineWidth(1.6)
                    pdf.line(32, y, width-32, y)
                    pdf.setLineWidth(.5)
            pdf.setFillColor(report_grey)
            pdf.setFont(report_regular_font,7)
            pdf.drawString(32,42, f'Red cells: Subcooling > {subcooling_limit:g} K; Superheating > {superheating_limit:g} K; T_CPU > {cpu_limit:g}°C. Blanks: unavailable.')
            draw_page_footer(pdf,page_number,width)
    return page_number


def validate_colormap_temperature_bounds(minimum, maximum):
    """Accept any finite user-selected range, including temperatures above 85°C."""
    minimum, maximum = float(minimum), float(maximum)
    if not (math.isfinite(minimum) and math.isfinite(maximum) and minimum < maximum):
        raise ValueError("Colormap limits must be finite, with minimum below maximum.")
    return minimum, maximum


def colormap_result_temperature_range(test_details):
    """Return the range of the plateau averages actually used in the maps."""
    temperatures = [
        float(value)
        for detail in test_details
        for snapshot in build_plateau_snapshots(detail)
        for value in snapshot["temperatures"].values()
        if value is not None and math.isfinite(float(value))
    ]
    return (min(temperatures), max(temperatures)) if temperatures else None


def default_report_configuration(
    has_filling_ratio_analysis,
    repeated_group_count,
    test_details,
    performance_characterization_ratios=None,
):
    """Return sensible defaults for non-interactive or failed-UI runs."""
    performance_characterization_ratios = (
        performance_characterization_ratios or []
    )
    board_tests = any(is_board_test_detail(detail) for detail in test_details)
    return {
        "include_transient_tests": any(is_transient_detail(d) for d in test_details),
        "include_filling_ratio_analysis": bool(has_filling_ratio_analysis),
        "include_performance_characterization": bool(
            performance_characterization_ratios
        ),
        "include_repeated_tests": repeated_group_count > 0,
        "include_psat_comparison": any(lts_ph_pressure_columns(d) for d in test_details),
        "include_ph_diagram": False,
        "ph_allow_estimated_pressure": False,
        "ph_assume_saturated_endpoints": True,
        "ph_saturation_tolerance_k": 0.2,
        "include_psat": any(shift2dc_extra_columns(d, 'psat') for d in test_details),
        "include_psu_temperatures": any(shift2dc_extra_columns(d, 'psu') for d in test_details),
        "pressure_settings": {},
        "include_colormaps": True,
        "include_raw_data": board_tests,
        "include_test_summary": board_tests or bool(has_filling_ratio_analysis),
        "summary_value_columns": None,
        "summary_subcooling_limit": 5.0,
        "summary_superheating_limit": 1.0,
        "include_superheating_subcooling": any(d["metadata"].get("part_type") == "LTS" and not is_transient_detail(d) for d in test_details),
        "summary_cpu_temperature_limit": 100.0,
        "selected_detail_files": {d["source_file"] for d in test_details if not is_transient_detail(d)},
        "selected_raw_files": {detail["source_file"] for detail in test_details},
        "selected_colormap_files": {
            detail["source_file"] for detail in test_details
        },
        "colormap_min_temperature": 25.0,
        "colormap_max_temperature": 85.0,
    }


def report_test_selection_label(test_detail):
    """Build a compact, unique label for one selectable per-test page."""
    metadata = test_detail["metadata"]
    if metadata["condition_type"] == "FR":
        condition = f"FR {metadata['condition_value']:g}%"
    else:
        condition = f"Charge {metadata['condition_value']:g}"
    comment = metadata.get("comment", "").strip() or "reference"
    label_parts = [
        condition,
        flow_rate_report_label(
            metadata["flow_type"],
            metadata["nominal_flow_rate"],
        ),
        inlet_temperature_report_label(
            metadata["medium"],
            metadata["nominal_temperature"],
            compact=True,
        ),
    ]
    if str(metadata.get("orientation", "N/A")).upper() != "N/A":
        label_parts.append(
            f"Orientation {orientation_report_name(metadata['orientation'])}"
        )
    label_parts.append(comment)
    return " | ".join(label_parts)


def show_report_configuration_dialog(
    part_type,
    fluids,
    coolants,
    test_details,
    fluid_properties,
    has_filling_ratio_analysis,
    repeated_group_count,
    performance_characterization_ratios,
    all_results=None,
):
    """Ask which optional analyses and per-test pages belong in the PDF."""
    defaults = default_report_configuration(
        has_filling_ratio_analysis,
        repeated_group_count,
        test_details,
        performance_characterization_ratios,
    )

    board_tests = any(is_board_test_detail(detail) for detail in test_details)
    transient_count = sum(is_transient_detail(d) for d in test_details)
    steady_details = [d for d in test_details if not is_transient_detail(d)]

    try:
        root = tk.Tk()
    except tk.TclError as error:
        print(f"Report interface unavailable ({error}); using default options.")
        return defaults

    root.title("JJ Cooling - Report setup")
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    window_width = min(980, max(760, screen_width - 80))
    window_height = min(760, max(600, screen_height - 80))
    x_position = max(0, (screen_width - window_width) // 2)
    y_position = max(0, (screen_height - window_height) // 2)
    root.geometry(
        f"{window_width}x{window_height}+{x_position}+{y_position}"
    )
    root.minsize(min(820, window_width), min(650, window_height))
    root.attributes("-topmost", True)

    style = ttk.Style(root)
    try:
        style.theme_use("vista" if os.name == "nt" else "clam")
    except tk.TclError:
        pass
    style.configure("Title.TLabel", font=("Segoe UI", 16, "bold"))
    style.configure("Heading.TLabel", font=("Segoe UI", 10, "bold"))

    setup_canvas = tk.Canvas(root, highlightthickness=0)
    setup_scroll = ttk.Scrollbar(root, orient="vertical", command=setup_canvas.yview)
    setup_scroll.pack(side="right", fill="y")
    setup_canvas.pack(side="left", fill="both", expand=True)
    setup_canvas.configure(yscrollcommand=setup_scroll.set)
    outer = ttk.Frame(setup_canvas, padding=18)
    setup_window = setup_canvas.create_window((0, 0), window=outer, anchor="nw")
    outer.bind("<Configure>", lambda _event: setup_canvas.configure(scrollregion=setup_canvas.bbox("all")))
    setup_canvas.bind("<Configure>", lambda event: setup_canvas.itemconfigure(setup_window, width=event.width))
    outer.columnconfigure(0, weight=1)
    outer.rowconfigure(4, weight=1)

    ttk.Label(
        outer,
        text="Report setup",
        style="Title.TLabel",
    ).grid(row=0, column=0, sticky="w")
    ttk.Label(
        outer,
        text="Review the detected campaign and choose the optional PDF sections.",
    ).grid(row=1, column=0, sticky="w", pady=(2, 12))

    summary_frame = ttk.LabelFrame(outer, text="Detected test campaign", padding=10)
    summary_frame.grid(row=2, column=0, sticky="ew")
    for column in range(4):
        summary_frame.columnconfigure(column, weight=1)

    summary_items = [
        ("Part type", part_type),
        ("Number of CSV tests", str(len(test_details))),
        ("Transient tests (_TR)", str(transient_count)),
        ("Working fluid", ", ".join(sorted(fluids))),
        ("Coolant", ", ".join(sorted(coolants))),
        (
            "Filling-ratio analysis",
            "Detected" if has_filling_ratio_analysis else "Not detected",
        ),
        ("Repeated condition groups", str(repeated_group_count)),
        (
            "Performance characterization",
            (
                "Detected for "
                + ", ".join(
                    f"FR {float(ratio):g}%"
                    for ratio in performance_characterization_ratios
                )
                if performance_characterization_ratios
                else "Not detected"
            ),
        ),
    ]
    for item_index, (label, value) in enumerate(summary_items):
        row = (item_index // 3) * 2
        column = item_index % 3
        ttk.Label(
            summary_frame,
            text=label,
            style="Heading.TLabel",
        ).grid(row=row, column=column, sticky="w", padx=(0, 14))
        ttk.Label(
            summary_frame,
            text=value,
        ).grid(row=row + 1, column=column, sticky="w", padx=(0, 14), pady=(1, 8))

    property_lines = [
        format_refprop_property_line(
            fluid,
            fluid_properties[fluid],
            include_error=True,
        )
        for fluid in sorted(fluids)
    ]
    property_color = (
        "#303030"
        if all(fluid_properties[fluid].get("available") for fluid in fluids)
        else "#A51D1D"
    )
    style.configure("Property.TLabel", foreground=property_color)
    property_start_row = 2 * math.ceil(len(summary_items) / 3)
    ttk.Label(
        summary_frame,
        text="REFPROP working-fluid properties",
        style="Heading.TLabel",
    ).grid(
        row=property_start_row,
        column=0,
        columnspan=3,
        sticky="w",
    )
    ttk.Label(
        summary_frame,
        text="\n".join(property_lines),
        justify="left",
        anchor="w",
        style="Property.TLabel",
        wraplength=max(600, window_width - 105),
    ).grid(
        row=property_start_row + 1,
        column=0,
        columnspan=3,
        sticky="ew",
        pady=(1, 2),
    )

    options_frame = ttk.LabelFrame(outer, text="Report sections", padding=10)
    options_frame.grid(row=3, column=0, sticky="ew", pady=(12, 10))

    filling_ratio_variable = tk.BooleanVar(
        value=defaults["include_filling_ratio_analysis"]
    )
    filling_ratio_checkbox = ttk.Checkbutton(
        options_frame,
        text="Include the filling-ratio analysis pages",
        variable=filling_ratio_variable,
    )
    filling_ratio_checkbox.grid(row=0, column=0, sticky="w", padx=(0, 30))
    if not has_filling_ratio_analysis:
        filling_ratio_checkbox.state(["disabled"])

    repeated_tests_variable = tk.BooleanVar(
        value=defaults["include_repeated_tests"]
    )
    repeated_tests_checkbox = ttk.Checkbutton(
        options_frame,
        text="Include the repeated-tests page",
        variable=repeated_tests_variable,
    )
    repeated_tests_checkbox.grid(row=0, column=1, sticky="w")
    if repeated_group_count == 0:
        repeated_tests_checkbox.state(["disabled"])

    performance_variable = tk.BooleanVar(
        value=defaults["include_performance_characterization"]
    )
    performance_ratio_text = ", ".join(
        f"FR {float(ratio):g}%"
        for ratio in performance_characterization_ratios
    )
    performance_checkbox = ttk.Checkbutton(
        options_frame,
        text=(
            "Include performance-characterization pages"
            + (
                f" ({performance_ratio_text})"
                if performance_ratio_text
                else ""
            )
        ),
        variable=performance_variable,
    )
    performance_checkbox.grid(
        row=1,
        column=0,
        columnspan=2,
        sticky="w",
        pady=(8, 0),
    )
    if not performance_characterization_ratios:
        performance_checkbox.state(["disabled"])

    colormap_variable = tk.BooleanVar(value=bool(steady_details))
    colormap_checkbox = ttk.Checkbutton(
        options_frame,
        text=("Include per-test CPU colormap and raw-data pages" if board_tests
              else "Include per-test T_CU colormap and raw-data pages"),
        variable=colormap_variable,
    )
    colormap_checkbox.grid(
        row=2,
        column=0,
        columnspan=2,
        sticky="w",
        pady=(8, 0),
    )

    summary_variable = tk.BooleanVar(value=defaults["include_test_summary"])
    ttk.Checkbutton(
        options_frame, text="Include test summaries (averages and tables)", variable=summary_variable,
    ).grid(row=3, column=0, columnspan=2, sticky="w", pady=(8, 0))

    pressure_columns = list(dict.fromkeys(
        column for detail in test_details for column in lts_ph_pressure_columns(detail)
    ))
    psat_variable = tk.BooleanVar(value=defaults['include_psat'])
    psu_variable = tk.BooleanVar(value=defaults['include_psu_temperatures'])
    pressure_settings_holder = {'settings': {}}

    def edit_pressure_settings():
        settings = show_pressure_settings_dialog(root, pressure_columns, pressure_settings_holder['settings'])
        if settings is not None:
            pressure_settings_holder['settings'] = settings
        return settings is not None

    if board_tests:
        psat_checkbox = ttk.Checkbutton(options_frame,
            text='Include Psat / T_sat comparison pages' + ('' if pressure_columns else ' (no Psat found)'),
            variable=psat_variable)
        psat_checkbox.grid(row=4, column=0, sticky='w', pady=(8, 0))
        pressure_button = ttk.Button(options_frame, text='Pressure units / reference...', command=edit_pressure_settings)
        pressure_button.grid(row=4, column=1, sticky='w', pady=(8, 0))
        if not pressure_columns:
            psat_checkbox.state(['disabled'])
            pressure_button.state(['disabled'])
        psu_checkbox = ttk.Checkbutton(options_frame,
            text='Include electronics PSU temperature pages'
                 + ('' if defaults['include_psu_temperatures'] else ' (no T_PSU found)'), variable=psu_variable)
        psu_checkbox.grid(row=5, column=0, columnspan=2, sticky='w', pady=(8, 0))
        if not defaults['include_psu_temperatures']:
            psu_checkbox.state(['disabled'])

    psat_comparison_variable = tk.BooleanVar(value=defaults['include_psat_comparison'])
    superheating_variable = tk.BooleanVar(value=defaults['include_superheating_subcooling'])
    if part_type == 'LTS':
        ttk.Checkbutton(options_frame, text='Include superheating and subcooling comparison pages',
                        variable=superheating_variable).grid(row=16, column=0, columnspan=2, sticky='w', pady=(8,0))
    ph_variable = tk.BooleanVar(value=False)
    ph_estimated_variable = tk.BooleanVar(value=False)
    ph_endpoints_variable = tk.BooleanVar(value=True)
    if part_type == 'LTS':
        comparison_checkbox = ttk.Checkbutton(options_frame, text='Include Psat comparison pages'
            + ('' if pressure_columns else ' (no Psat found)'), variable=psat_comparison_variable)
        comparison_checkbox.grid(row=11, column=0, columnspan=2, sticky='w', pady=(8, 0))
        if not pressure_columns:
            comparison_checkbox.state(['disabled'])
        ttk.Checkbutton(options_frame, text='Include p-h diagrams (one path per heat-load plateau)',
                        variable=ph_variable).grid(row=6, column=0, columnspan=2, sticky='w', pady=(8, 0))
        if not board_tests:
            pressure_button = ttk.Button(options_frame, text='Pressure units / reference...', command=edit_pressure_settings)
            pressure_button.grid(row=7, column=1, sticky='w')
            if not pressure_columns:
                pressure_button.state(['disabled'])
        ph_estimate_checkbox = ttk.Checkbutton(options_frame,
            text='If Psat is missing: estimate pressure from saturated T_EVAP_OUT', variable=ph_estimated_variable)
        ph_estimate_checkbox.grid(row=8, column=0, columnspan=2, sticky='w')
        ph_endpoint_checkbox = ttk.Checkbutton(options_frame,
            text='Near saturation: assume liquid at evap-in/cond-out, vapor at evap-out/cond-in',
            variable=ph_endpoints_variable)
        ph_endpoint_checkbox.grid(row=9, column=0, columnspan=2, sticky='w')
        ttk.Label(options_frame, text=('Shift2DC: Q=0.5 at evap-out/cond-in; measured Psat corrected by +/- rho_liq*g*0.19 m. '
                  'The near-saturation checkbox applies to liquid endpoints only.' if any(is_shift2dc_file(Path(d['source_file'])) for d in steady_details) else
                  'One pressure is used for all four states. Near saturation means within 0.2 K; uncheck to omit ambiguous points.'),
                  wraplength=820).grid(row=10, column=0, columnspan=2, sticky='w', pady=(2, 0))
        def update_ph_controls(*_args):
            for widget in [ph_estimate_checkbox, ph_endpoint_checkbox]:
                widget.state(['!disabled'] if ph_variable.get() else ['disabled'])
        ph_variable.trace_add('write', update_ph_controls)
        update_ph_controls()

    transient_variable = tk.BooleanVar(value=bool(transient_count))
    transient_checkbox = ttk.Checkbutton(options_frame,
        text=f'Include transient time-series pages ({transient_count} _TR tests detected)',
        variable=transient_variable)
    transient_checkbox.grid(row=13, column=0, columnspan=2, sticky='w', pady=(8,0))
    if not transient_count:
        transient_checkbox.state(['disabled'])

    summary_frame_data = pd.DataFrame(all_results) if all_results is not None else pd.DataFrame(
        [row for detail in test_details for row in detail.get('results', [])]
    )
    if board_tests and len(summary_frame_data) and 'Temperature source' not in summary_frame_data:
        summary_frame_data = pd.DataFrame(shift2dc_report_rows(summary_frame_data.to_dict('records')))
    summary_options, summary_defaults = summary_value_options(summary_frame_data, part_type, test_details)
    summary_values_frame = ttk.LabelFrame(options_frame, text='Values in the summary table', padding=8)
    summary_values_frame.grid(row=12, column=0, columnspan=2, sticky='ew', pady=(10,0))
    ttk.Label(summary_values_frame, text='Conditions and heat load are always shown. Choose the result columns below.').grid(
        row=0, column=0, columnspan=3, sticky='w', pady=(0,5))
    summary_value_variables = {}
    summary_value_widgets = []
    for index, column in enumerate(summary_options):
        variable = tk.BooleanVar(value=column in summary_defaults)
        summary_value_variables[column] = variable
        widget = ttk.Checkbutton(summary_values_frame, text=summary_value_label(column, summary_frame_data), variable=variable)
        widget.grid(row=1+index//3, column=index%3, sticky='w', padx=(0,15), pady=1)
        summary_value_widgets.append(widget)
    def set_summary_values(columns):
        for column, variable in summary_value_variables.items():
            variable.set(column in columns)
    summary_buttons = ttk.Frame(summary_values_frame)
    summary_buttons.grid(row=2+len(summary_options)//3, column=0, columnspan=3, sticky='w', pady=(5,0))
    for label, columns in [('Defaults', summary_defaults), ('Select all', summary_options), ('Clear all', [])]:
        widget = ttk.Button(summary_buttons, text=label, command=lambda columns=columns: set_summary_values(columns))
        widget.pack(side='left', padx=(0,6))
        summary_value_widgets.append(widget)
    def update_summary_controls(*_args):
        for widget in summary_value_widgets:
            widget.state(['!disabled'] if summary_variable.get() else ['disabled'])
    summary_variable.trace_add('write', update_summary_controls)
    update_summary_controls()

    thresholds_frame = ttk.LabelFrame(options_frame, text='Summary red-cell thresholds', padding=8)
    thresholds_frame.grid(row=14, column=0, columnspan=2, sticky='ew', pady=(8,0))
    subcooling_limit_variable = tk.StringVar(value='5')
    cpu_limit_variable = tk.StringVar(value='100')
    superheating_limit_variable = tk.StringVar(value='1')
    ttk.Label(thresholds_frame, text='Subcooling above [K = °C difference]:').grid(row=0, column=0, sticky='w')
    ttk.Entry(thresholds_frame, textvariable=subcooling_limit_variable, width=7).grid(row=0, column=1, padx=(5,18))
    ttk.Label(thresholds_frame, text='T_CPU above [°C]:').grid(row=0, column=2, sticky='w')
    ttk.Entry(thresholds_frame, textvariable=cpu_limit_variable, width=7).grid(row=0, column=3, padx=5)
    if part_type == 'LTS':
        ttk.Label(thresholds_frame, text='Superheating above [K = °C difference]:').grid(row=1, column=0, sticky='w', pady=(6,0))
        ttk.Entry(thresholds_frame, textvariable=superheating_limit_variable, width=7).grid(row=1, column=1, padx=(5,18), pady=(6,0))

    details_frame = ttk.LabelFrame(options_frame, text='Tests for extra details: Psat / T_sat, PSU temperatures and p-h', padding=8)
    details_frame.grid(row=15, column=0, columnspan=2, sticky='ew', pady=(10,0))
    detail_variables = {d['source_file']: tk.BooleanVar(value=True) for d in steady_details}
    detail_buttons = ttk.Frame(details_frame)
    detail_buttons.pack(fill='x')
    def set_all_details(value):
        for variable in detail_variables.values():
            variable.set(value)
    ttk.Button(detail_buttons, text='Select all', command=lambda: set_all_details(True)).pack(side='left')
    ttk.Button(detail_buttons, text='Clear all', command=lambda: set_all_details(False)).pack(side='left', padx=6)
    ttk.Label(details_frame, text='Only enabled extra sections are included, grouped consecutively for each selected test.').pack(anchor='w', pady=4)
    detail_canvas = tk.Canvas(details_frame, height=130, highlightthickness=0)
    detail_scroll = ttk.Scrollbar(details_frame, orient='vertical', command=detail_canvas.yview)
    detail_scroll.pack(side='right', fill='y')
    detail_canvas.pack(fill='x', expand=True)
    detail_canvas.configure(yscrollcommand=detail_scroll.set)
    detail_inner = ttk.Frame(detail_canvas)
    detail_window = detail_canvas.create_window((0,0), window=detail_inner, anchor='nw')
    detail_inner.columnconfigure(1, weight=1)
    detail_inner.bind('<Configure>', lambda _event: detail_canvas.configure(scrollregion=detail_canvas.bbox('all')))
    detail_canvas.bind('<Configure>', lambda event: detail_canvas.itemconfigure(detail_window, width=event.width))
    for index, detail in enumerate(steady_details):
        ttk.Checkbutton(detail_inner, variable=detail_variables[detail['source_file']]).grid(row=index,column=0,sticky='n')
        ttk.Label(detail_inner, text=detail['source_file'], wraplength=760).grid(row=index,column=1,sticky='w',pady=2)

    tests_frame = ttk.LabelFrame(
        outer,
        text="Tests included in the colormap section",
        padding=8,
    )
    tests_frame.grid(row=4, column=0, sticky="nsew")
    tests_frame.columnconfigure(0, weight=1)
    tests_frame.rowconfigure(1, weight=1)

    controls = ttk.Frame(tests_frame)
    controls.grid(row=0, column=0, sticky="ew", pady=(0, 6))
    limits_frame = ttk.Frame(controls)
    limits_frame.pack(fill="x")
    minimum_temperature_variable = tk.StringVar(value="25")
    maximum_temperature_variable = tk.StringVar(value="85")
    ttk.Label(limits_frame, text="Color scale minimum [°C]:").pack(side="left")
    minimum_temperature_entry = ttk.Entry(
        limits_frame, textvariable=minimum_temperature_variable, width=8,
    )
    minimum_temperature_entry.pack(side="left", padx=(6, 18))
    ttk.Label(limits_frame, text="Maximum [°C]:").pack(side="left")
    maximum_temperature_entry = ttk.Entry(
        limits_frame, textvariable=maximum_temperature_variable, width=8,
    )
    maximum_temperature_entry.pack(side="left", padx=(6, 18))
    temperature_range_variable = tk.StringVar()
    ttk.Label(controls, textvariable=temperature_range_variable).pack(
        anchor="w", pady=(5, 0),
    )
    ttk.Label(
        controls, text="These limits apply to all selected maps; values outside them use the end colors.",
    ).pack(anchor="w")

    list_container = ttk.Frame(tests_frame)
    list_container.grid(row=1, column=0, sticky="nsew")
    list_container.columnconfigure(0, weight=1)
    list_container.rowconfigure(0, weight=1)
    tests_canvas = tk.Canvas(
        list_container,
        highlightthickness=1,
        highlightbackground="#C8C8C8",
        height=230,
    )
    tests_scrollbar = ttk.Scrollbar(
        list_container,
        orient="vertical",
        command=tests_canvas.yview,
    )
    tests_canvas.configure(yscrollcommand=tests_scrollbar.set)
    tests_canvas.grid(row=0, column=0, sticky="nsew")
    tests_scrollbar.grid(row=0, column=1, sticky="ns")

    tests_inner = ttk.Frame(tests_canvas, padding=5)
    tests_window = tests_canvas.create_window(
        (0, 0),
        window=tests_inner,
        anchor="nw",
    )
    tests_inner.bind(
        "<Configure>",
        lambda _event: tests_canvas.configure(
            scrollregion=tests_canvas.bbox("all")
        ),
    )
    tests_canvas.bind(
        "<Configure>",
        lambda event: tests_canvas.itemconfigure(
            tests_window,
            width=event.width,
        ),
    )

    test_variables = {}
    test_checkboxes = []
    for test_index, test_detail in enumerate(steady_details):
        source_file = test_detail["source_file"]
        variable = tk.BooleanVar(value=True)
        test_variables[source_file] = variable
        checkbox = ttk.Checkbutton(
            tests_inner,
            text=report_test_selection_label(test_detail),
            variable=variable,
        )
        checkbox.grid(row=test_index, column=0, sticky="w", pady=1)
        test_checkboxes.append(checkbox)

    test_temperature_ranges = {
        detail["source_file"]: colormap_result_temperature_range([detail])
        for detail in test_details
    }

    def update_temperature_range(*_args):
        selected = [name for name, variable in test_variables.items() if variable.get()]
        ranges = [test_temperature_ranges[name] for name in selected
                  if test_temperature_ranges[name] is not None]
        if ranges:
            temperature_range_variable.set(
                f"Selected plateau averages: minimum {min(r[0] for r in ranges):.1f}°C"
                f" | maximum {max(r[1] for r in ranges):.1f}°C"
            )
        else:
            temperature_range_variable.set(
                "Selected plateau averages: no valid temperatures."
                if selected else "Select tests to see their temperature range."
            )

    for variable in test_variables.values():
        variable.trace_add("write", update_temperature_range)
    update_temperature_range()

    selection_buttons = ttk.Frame(tests_frame)
    selection_buttons.grid(row=2, column=0, sticky="w", pady=(7, 0))

    def set_all_tests(selected):
        for variable in test_variables.values():
            variable.set(selected)

    select_all_button = ttk.Button(
        selection_buttons,
        text="Select all",
        command=lambda: set_all_tests(True),
    )
    select_all_button.pack(side="left")
    clear_all_button = ttk.Button(
        selection_buttons,
        text="Clear all",
        command=lambda: set_all_tests(False),
    )
    clear_all_button.pack(side="left", padx=(7, 0))

    def update_colormap_controls(*_args):
        enabled = colormap_variable.get()
        for widget in [
            minimum_temperature_entry,
            maximum_temperature_entry,
            select_all_button,
            clear_all_button,
            *test_checkboxes,
        ]:
            widget.state(["!disabled"] if enabled else ["disabled"])

    colormap_variable.trace_add("write", update_colormap_controls)
    update_colormap_controls()

    result_holder = {"configuration": None}

    def accept_configuration():
        try:
            subcooling_limit, cpu_limit, superheating_limit = summary_red_thresholds({
                'summary_subcooling_limit': subcooling_limit_variable.get().replace(',', '.'),
                'summary_superheating_limit': superheating_limit_variable.get().replace(',', '.'),
                'summary_cpu_temperature_limit': cpu_limit_variable.get().replace(',', '.'),
            })
        except (ValueError, TypeError):
            messagebox.showerror('Invalid summary threshold', 'Enter finite numbers for all thresholds.', parent=root)
            return
        selected_detail_files = {name for name, variable in detail_variables.items() if variable.get()}
        include_colormaps = colormap_variable.get()
        selected_files = {
            source_file
            for source_file, variable in test_variables.items()
            if variable.get()
        }

        if include_colormaps and not selected_files:
            messagebox.showerror(
                "No tests selected",
                "Select at least one test, or disable the per-test section.",
                parent=root,
            )
            return

        minimum_temperature, maximum_temperature = 25.0, 85.0
        if include_colormaps:
            try:
                minimum_temperature, maximum_temperature = validate_colormap_temperature_bounds(
                    minimum_temperature_variable.get().strip().replace(",", "."),
                    maximum_temperature_variable.get().strip().replace(",", "."),
                )
            except (TypeError, ValueError):
                messagebox.showerror(
                    "Invalid color scale",
                    "Enter finite numeric temperatures with the minimum lower than the maximum.",
                    parent=root,
                )
                return

        if ph_variable.get() and not ph_estimated_variable.get():
            missing_pressure = [d['source_file'] for d in steady_details if d['source_file'] in selected_detail_files and not lts_ph_pressure_columns(d)]
            if missing_pressure:
                messagebox.showerror('p-h pressure required',
                    'No Psat column in ' + missing_pressure[0] +
                    '. Enable the saturated T_EVAP_OUT pressure estimate, or disable p-h diagrams.', parent=root)
                return

        if psat_variable.get() or ph_variable.get() or (part_type == 'LTS' and pressure_columns):
            try:
                for column in pressure_columns:
                    pressure_column_settings(column, {'pressure_settings': pressure_settings_holder['settings']})
            except ValueError:
                if not edit_pressure_settings():
                    return

        result_holder["configuration"] = {
            "include_transient_tests": bool(transient_count) and transient_variable.get(),
            "include_filling_ratio_analysis": (
                has_filling_ratio_analysis and filling_ratio_variable.get()
            ),
            "include_performance_characterization": (
                bool(performance_characterization_ratios)
                and performance_variable.get()
            ),
            "include_repeated_tests": (
                repeated_group_count > 0 and repeated_tests_variable.get()
            ),
            "include_psat_comparison": part_type == 'LTS' and psat_comparison_variable.get(),
            "include_ph_diagram": part_type == 'LTS' and ph_variable.get(),
            "ph_allow_estimated_pressure": ph_estimated_variable.get(),
            "ph_assume_saturated_endpoints": ph_endpoints_variable.get(),
            "ph_saturation_tolerance_k": 0.2,
            "include_psat": psat_variable.get(),
            "include_psu_temperatures": psu_variable.get(),
            "pressure_settings": pressure_settings_holder['settings'],
            "include_colormaps": include_colormaps,
            "include_raw_data": include_colormaps and board_tests,
            "include_test_summary": summary_variable.get(),
            "summary_value_columns": [c for c,v in summary_value_variables.items() if v.get()],
            "summary_subcooling_limit": subcooling_limit,
            "summary_superheating_limit": superheating_limit,
            "include_superheating_subcooling": part_type == "LTS" and superheating_variable.get(),
            "summary_cpu_temperature_limit": cpu_limit,
            "selected_detail_files": selected_detail_files,
            "selected_raw_files": selected_files,
            "selected_colormap_files": selected_files,
            "colormap_min_temperature": minimum_temperature,
            "colormap_max_temperature": maximum_temperature,
        }
        root.destroy()

    def cancel_configuration():
        result_holder["configuration"] = None
        root.destroy()

    action_frame = ttk.Frame(outer)
    action_frame.grid(row=5, column=0, sticky="e", pady=(12, 0))
    ttk.Button(
        action_frame,
        text="Cancel",
        command=cancel_configuration,
    ).pack(side="left", padx=(0, 8))
    ttk.Button(
        action_frame,
        text="Generate Excel and PDF",
        command=accept_configuration,
    ).pack(side="left")

    root.protocol("WM_DELETE_WINDOW", cancel_configuration)
    root.after(250, lambda: root.attributes("-topmost", False))
    root.mainloop()
    return result_holder["configuration"]


def open_pdf_automatically(pdf_file):
    """Open the completed PDF with the operating system's default viewer."""
    pdf_path = Path(pdf_file).resolve()
    try:
        if os.name == "nt":
            os.startfile(str(pdf_path))
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(pdf_path)])
        else:
            subprocess.Popen(["xdg-open", str(pdf_path)])
    except (OSError, subprocess.SubprocessError) as error:
        print(f"The PDF was created but could not be opened automatically: {error}")


def read_csv_automatically(file_path):
    """Read a CSV and automatically detect comma/semicolon/tab separators."""
    try:
        data = pd.read_csv(file_path, sep=None, engine="python", encoding="utf-8-sig")
    except UnicodeDecodeError:
        data = pd.read_csv(file_path, sep=None, engine="python", encoding="cp1252")
    data.columns = data.columns.astype(str).str.strip()
    return data


def parse_file_name(file_path):
    """
    Parse both supported filename structures.

    PHP example:
    e-Durable_PHP-AD-ED-5-1-1_R1336mzzE_OS_Ch5_Water_TW30_VFR1_SS.csv

    LTS example:
    R&D2_LTS_EVAP-spreading-A_COND-AD-ED-4-1-1_R1336mzzE_Ch61_Water_TW30_VFR1_SS.csv
    """
    parts = file_path.stem.split("_")

    type_index = None
    type_match = None
    for index, part in enumerate(parts):
        match = re.match(r"^(PHP|LTS)(?:-|$)", part, re.IGNORECASE)
        if match:
            type_index = index
            type_match = match
            break

    if type_index is None or type_match is None:
        raise ValueError(
            f"{file_path.name}: part type PHP or LTS was not found in the filename"
        )

    part_type = type_match.group(1).upper()

    condition_index = None
    condition_match = None
    for index, part in enumerate(parts):
        match = re.fullmatch(
            r"(FR|Ch)([-+]?\d+(?:\.\d+)?)",
            part,
            re.IGNORECASE,
        )
        if match:
            condition_index = index
            condition_match = match
            break

    if condition_index is None or condition_match is None:
        raise ValueError(
            f"{file_path.name}: filling ratio (FR70) or charge (Ch5) was not found"
        )

    number_match = None
    number_index = None
    for index in range(type_index, condition_index):
        part = parts[index]
        match = re.search(r"(\d+-\d+-\d+)$", part)
        if match:
            number_match = match
            number_index = index

    if (number_match is None or number_index is None) and part_type == "LTS":
        for index in range(type_index + 1, condition_index):
            match = re.fullmatch(r"COND-(.+)", parts[index], re.IGNORECASE)
            if match:
                number_match, number_index = match, index
                break

    if number_match is None or number_index is None:
        raise ValueError(
            f"{file_path.name}: part number such as 4-1-1 or 5-1-1 was not found"
        )

    # An orientation is optional. PHP files normally contain OS, OV, etc.
    # The current R&D2 LTS filenames have no orientation field.
    orientation_index = None
    orientation = "N/A"
    for index in range(type_index + 1, condition_index):
        if re.fullmatch(r"O[A-Za-z0-9+.-]+", parts[index], re.IGNORECASE):
            orientation_index = index
            orientation = parts[index][1:]
            break

    if orientation_index is not None:
        fluid_index = orientation_index - 1
    else:
        fluid_index = condition_index - 1

    if fluid_index <= type_index:
        raise ValueError(
            f"{file_path.name}: working fluid was not found before the FR/charge field"
        )

    working_fluid = parts[fluid_index]

    coolant_index = None
    for index in range(condition_index + 1, len(parts)):
        if parts[index].lower() in {"water", "air"}:
            coolant_index = index
            break

    if coolant_index is None:
        raise ValueError(
            f"{file_path.name}: coolant Water or Air was not found"
        )

    condition_type = (
        "FR" if condition_match.group(1).upper() == "FR" else "Charge"
    )
    medium = parts[coolant_index].capitalize()

    report_prefix = "_".join(parts[:type_index]).strip("_ -")
    if not report_prefix:
        report_prefix = "JJ Cooling"

    # The complete component name begins with PHP/LTS and ends with
    # the part number. This supports both one-token PHP component names
    # and multi-token LTS component names.
    part_name = "_".join(parts[type_index:number_index + 1])

    evaporator_name = None
    condenser_name = None
    if part_type == "LTS":
        component_tokens = parts[type_index + 1:number_index + 1]
        for token in component_tokens:
            evaporator_match = re.fullmatch(
                r"EVAP-(.+)",
                token,
                re.IGNORECASE,
            )
            condenser_match = re.fullmatch(
                r"COND-(.+)",
                token,
                re.IGNORECASE,
            )
            if evaporator_match:
                evaporator_name = evaporator_match.group(1)
            if condenser_match:
                condenser_name = condenser_match.group(1)

        if not evaporator_name or not condenser_name:
            raise ValueError(
                f"{file_path.name}: an LTS filename must contain both "
                "EVAP-<name> and COND-<name> component fields"
            )

    nominal_temperature = None
    nominal_flow_rate = None

    # Search the complete filename rather than only underscore-separated
    # tokens. This accepts TW30, T_WATER30, T_WATER_30, TA30, T_AIR30,
    # T_AIR_30, VFR1, VFR_1, CFM100, and CFM_100.
    if medium == "Water":
        temperature_pattern = (
            r"(?:^|_)(?:T_?WATER_?|TW_?)"
            r"([-+]?\d+(?:\.\d+)?)(?=_|$)"
        )
    else:
        temperature_pattern = (
            r"(?:^|_)(?:T_?AIR_?|TA_?)"
            r"([-+]?\d+(?:\.\d+)?)(?=_|$)"
        )

    temperature_match = re.search(
        temperature_pattern,
        file_path.stem,
        re.IGNORECASE,
    )
    if temperature_match:
        nominal_temperature = float(temperature_match.group(1))

    flow_match = re.search(
        r"(?:^|_)(VFR|CFM)_?([-+]?\d+(?:[.pP]\d+)?)(?=_|$)",
        file_path.stem,
        re.IGNORECASE,
    )
    if flow_match:
        flow_type = flow_match.group(1).upper()
        nominal_flow_rate = float(flow_match.group(2).lower().replace("p", "."))
    elif medium == "Air":
        flow_type = "CFM"
    else:
        flow_type = "VFR"

    flow_unit = "CFM" if flow_type == "CFM" else "l/min"

    # SS marks steady state; TR marks transient. Anything after the marker is a
    # free-form comment that distinguishes repeated tests without changing
    # their physical test conditions. Both _SS_Comment and _SSComment are
    # accepted; underscores inside a comment are preserved for the legend.
    comment = ""
    test_mode = "SS"
    for index in range(coolant_index + 1, len(parts)):
        token = parts[index]
        token_upper = token.upper()
        if token_upper in {"SS", "TR"}:
            test_mode = token_upper
            comment = "_".join(parts[index + 1:]).strip("_ ")
            break
        if token_upper[:2] in {"SS", "TR"} and len(token) > 2:
            test_mode = token_upper[:2]
            comment_tokens = [token[2:], *parts[index + 1:]]
            comment = "_".join(comment_tokens).strip("_ ")
            break

    return {
        "part_type": part_type,
        "part_number": number_match.group(1),
        "part_name": part_name,
        "evaporator_name": evaporator_name,
        "condenser_name": condenser_name,
        "report_prefix": report_prefix,
        "fluid": working_fluid,
        "orientation": orientation,
        "condition_type": condition_type,
        "condition_value": float(condition_match.group(2)),
        "medium": medium,
        "nominal_temperature": nominal_temperature,
        "flow_type": flow_type,
        "flow_unit": flow_unit,
        "flow_type_explicit": flow_match is not None,
        "nominal_flow_rate": nominal_flow_rate,
        "comment": comment,
        "test_mode": test_mode,
    }


def find_column(data, candidates):
    lookup = {column.upper(): column for column in data.columns}
    for candidate in candidates:
        if candidate.upper() in lookup:
            return lookup[candidate.upper()]
    raise ValueError(f"none of these columns were found: {', '.join(candidates)}")


def find_optional_column(data, candidates):
    """Return a matching column, or None when no candidate exists."""
    lookup = {column.upper(): column for column in data.columns}
    for candidate in candidates:
        if candidate.upper() in lookup:
            return lookup[candidate.upper()]
    return None


def find_t_cu_columns(data):
    excluded = {
        "T_CU_MIN", "T_CU_MAX", "T_CU_AVG",
        "T_CU_DELTALOW", "T_CU_DELTAHIGH",
    }
    return [
        column for column in data.columns
        if column.upper().startswith("T_CU") and column.upper() not in excluded
    ]


def t_cu_output_name(column):
    """Add the temperature unit to a T_CU heading exactly once."""
    if "[°C]" in column:
        return column
    return f"{column} [°C]"


def create_cleaned_csv(raw_file, output_file):
    """Create W_PSU when PSU setpoints exist and remove original PSU columns."""
    data = read_csv_automatically(raw_file)

    psu_1 = find_optional_column(data, ["W_PSU_1_SP"])
    psu_2 = find_optional_column(data, ["W_PSU_2_SP"])

    if psu_1 is not None and psu_2 is not None:
        combined_power = (
            pd.to_numeric(data[psu_1], errors="coerce")
            + pd.to_numeric(data[psu_2], errors="coerce")
        )
        if combined_power.isna().any():
            raise ValueError("W_PSU_1_SP or W_PSU_2_SP contains non-numeric values")

        first_psu_position = min(
            i for i, column in enumerate(data.columns) if "PSU" in column.upper()
        )
        psu_columns = [column for column in data.columns if "PSU" in column.upper()]
        data = data.drop(columns=psu_columns)
        data.insert(first_psu_position, "W_PSU", combined_power)

    elif psu_1 is not None or psu_2 is not None:
        raise ValueError("only one of W_PSU_1_SP and W_PSU_2_SP was found")

    elif find_optional_column(data, ["W_PSU", "W_HEATER"]) is None:
        raise ValueError(
            "no usable power column was found; expected both PSU setpoints, "
            "W_PSU, or W_HEATER"
        )

    else:
        # When W_PSU already exists, retain it but remove other PSU channels.
        removable_psu_columns = [
            column for column in data.columns
            if "PSU" in column.upper() and column.upper() != "W_PSU"
        ]
        data = data.drop(columns=removable_psu_columns)

    # Use Excel's standard comma-separated format and include a UTF-8 byte-order
    # mark. This makes double-clicked CSV files open in separate Excel columns
    # instead of placing each semicolon-separated row in column A.
    data.to_csv(
        output_file,
        index=False,
        sep=",",
        encoding="utf-8-sig",
        lineterminator="\n",
    )
    return data


def get_medium_columns(data, medium, flow_type=None):
    if medium == "Water":
        inlet = find_column(data, ["T_WATER_IN", "T_WATER_INLET"])
        outlet = find_column(data, ["T_WATER_OUT", "T_WATER_OUTLET"])
        flow = find_column(data, ["VFR_WATER", "VFR", "VFR_WATER_IN"])
    else:
        inlet = find_column(data, ["T_AIR_IN", "T_AIR_INLET"])
        outlet = find_column(data, ["T_AIR_OUT", "T_AIR_OUTLET"])
        if str(flow_type).upper() == "VFR":
            flow_candidates = [
                "VFR_AIR", "VFR", "VFR_AIR_IN",
                "CFM", "CFM_AIR", "AIR_CFM",
            ]
        else:
            flow_candidates = [
                "CFM", "CFM_AIR", "AIR_CFM",
                "VFR_AIR", "VFR", "VFR_AIR_IN",
            ]
        flow = find_column(data, flow_candidates)
    return inlet, outlet, flow


def get_power_column(data):
    return find_column(data, ["W_PSU", "W_HEATER"])


def complete_nominal_conditions(metadata, data):
    """Provide stable file-level flow/T_IN values if filename values are absent."""
    if metadata.get('test_mode') == 'TR':
        medium = metadata['medium'].upper()
        inlet_column = find_optional_column(data, [f'T_{medium}_IN', f'T_{medium}_INLET'])
        flow_column = find_optional_column(data, ['VFR_WATER', 'VFR', 'VFR_WATER_IN'] if medium == 'WATER'
            else ['CFM', 'CFM_AIR', 'AIR_CFM', 'VFR_AIR', 'VFR', 'VFR_AIR_IN'])
    else:
        inlet_column, _, flow_column = get_medium_columns(data, metadata["medium"], metadata["flow_type"])

    # When the filename does not state VFR or CFM, infer the unit from the CSV
    # heading. An explicit filename token takes priority over a generic LabVIEW
    # channel name such as VFR_AIR.
    if flow_column is not None and not metadata.get("flow_type_explicit", False):
        if "CFM" in flow_column.upper():
            metadata["flow_type"] = "CFM"
            metadata["flow_unit"] = "CFM"
        elif "VFR" in flow_column.upper():
            metadata["flow_type"] = "VFR"
            metadata["flow_unit"] = "l/min"

    if metadata["nominal_temperature"] is None and inlet_column is not None:
        inlet_values = pd.to_numeric(data[inlet_column], errors="coerce").dropna()
        if not inlet_values.empty:
            metadata["nominal_temperature"] = round(float(inlet_values.median()), 1)

    if metadata["nominal_flow_rate"] is None and flow_column is not None:
        flow_values = pd.to_numeric(data[flow_column], errors="coerce").dropna()
        if not flow_values.empty:
            metadata["nominal_flow_rate"] = round(
                float(flow_values.median()),
                3,
            )


def calculate_step_duration_seconds(step):
    """Calculate the complete plateau duration from an available time column."""
    time_column = find_optional_column(
        step,
        [
            "RelTime", "REL_TIME", "RelativeTime", "ElapsedTime",
            "Time", "Seconds", "Time_s", "Time [s]",
        ],
    )

    if time_column is None or len(step) < 2:
        return np.nan

    numeric_time = pd.to_numeric(step[time_column], errors="coerce")
    valid_numeric = numeric_time.dropna()

    if len(valid_numeric) >= 2:
        differences = valid_numeric.diff().dropna()
        positive_differences = differences[differences > 0]
        sample_interval = (
            float(positive_differences.median())
            if not positive_differences.empty
            else 0.0
        )
        duration = float(valid_numeric.iloc[-1] - valid_numeric.iloc[0])
        return max(0.0, duration + sample_interval)

    parsed_time = pd.to_datetime(step[time_column], errors="coerce")
    valid_time = parsed_time.dropna()

    if len(valid_time) >= 2:
        differences = valid_time.diff().dropna().dt.total_seconds()
        positive_differences = differences[differences > 0]
        sample_interval = (
            float(positive_differences.median())
            if not positive_differences.empty
            else 0.0
        )
        duration = (valid_time.iloc[-1] - valid_time.iloc[0]).total_seconds()
        return max(0.0, float(duration) + sample_interval)

    return np.nan


def split_into_power_steps(data, power_column):
    """
    Find valid power plateaus and keep their last sample_size rows.

    Short non-zero plateaus are ignored. They are normally synchronization
    artifacts created when PSU 1 and PSU 2 receive a new setpoint one after
    the other, briefly producing an intermediate combined W_PSU value.
    """
    power = pd.to_numeric(data[power_column], errors="coerce")
    if power.isna().any():
        raise ValueError(f"{power_column} contains non-numeric or empty values")

    changes = power.diff().abs().gt(power_change_tolerance)
    changes.iloc[0] = True
    group_number = changes.cumsum()
    steps = []

    for _, step in data.groupby(group_number, sort=False):
        mean_power = pd.to_numeric(step[power_column], errors="coerce").mean()
        is_non_zero_step = abs(mean_power) > power_change_tolerance
        required_rows = max(minimum_step_size, sample_size)
        is_long_enough = len(step) >= required_rows

        if is_non_zero_step and is_long_enough:
            averaging_window = step.tail(sample_size).copy()
            averaging_window.attrs["plateau_duration_s"] = (
                calculate_step_duration_seconds(step)
            )
            averaging_window.attrs["sample_count"] = len(averaging_window)
            steps.append(averaging_window)
    return steps


def numeric_average(data, column):
    values = pd.to_numeric(data[column], errors="coerce")
    if values.notna().sum() == 0:
        raise ValueError(f"column {column} has no numeric values")
    return float(values.mean())


def calculate_step_result(step, metadata, t_cu_columns):
    power_column = get_power_column(step)
    inlet_column, outlet_column, flow_column = get_medium_columns(
        step,
        metadata["medium"],
        metadata["flow_type"],
    )

    w_in = numeric_average(step, power_column)
    t_in = numeric_average(step, inlet_column)
    t_out = numeric_average(step, outlet_column)
    flow_rate = numeric_average(step, flow_column)
    t_cu_averages = {
        t_cu_output_name(column): numeric_average(step, column)
        for column in t_cu_columns
    }

    values = list(t_cu_averages.values())
    t_cu_min = min(values)
    t_cu_max = max(values)
    t_cu_avg = float(np.mean(values))
    delta_low = t_cu_avg - t_cu_min
    delta_high = t_cu_max - t_cu_avg
    delta_t_cu_min = t_cu_min - t_in
    delta_t_cu_max = t_cu_max - t_in
    delta_t_cu = t_cu_avg - t_in
    rth = delta_t_cu / w_in if w_in else np.nan

    density = water_density if metadata["medium"] == "Water" else air_density
    cp = water_cp if metadata["medium"] == "Water" else air_cp
    if metadata["flow_type"] == "CFM":
        volumetric_flow_m3_s = flow_rate * cfm_to_m3_s
    else:
        volumetric_flow_m3_s = flow_rate / 60000.0
    w_out = density * volumetric_flow_m3_s * cp * (t_out - t_in)

    lts_averages = {}
    subcooling = None
    t_adia = None

    if metadata["part_type"] == "LTS":
        for column in lts_temperature_columns:
            actual_column = find_column(step, [column])
            lts_averages[f"{column} [°C]"] = numeric_average(
                step,
                actual_column,
            )

        subcooling = (
            lts_averages["T_COND_IN [°C]"]
            - lts_averages["T_COND_OUT [°C]"]
        )
    else:
        t_adia_column = find_optional_column(step, ["T_ADIA"])
        if t_adia_column is not None:
            t_adia = numeric_average(step, t_adia_column)

    result = {
        "Source File": metadata["source_file"],
        "Comment": metadata["comment"],
        "Report Prefix": metadata["report_prefix"],
        "Part Name": metadata["part_name"],
        "Part Number": metadata["part_number"],
        "Evaporator Name": metadata["evaporator_name"],
        "Condenser Name": metadata["condenser_name"],
        "Working Fluid": metadata["fluid"],
        "Orientation": metadata["orientation"],
        "Condition Type": metadata["condition_type"],
        "Condition Value": metadata["condition_value"],
        "Coolant": metadata["medium"],
        "W_IN [W]": w_in,
        "W_OUT [W]": w_out,
        "Flow Type": metadata["flow_type"],
        "Flow Unit": metadata["flow_unit"],
        "Flow Rate": flow_rate,
        "Nominal Flow Rate": (
            metadata["nominal_flow_rate"]
            if metadata["nominal_flow_rate"] is not None
            else round(flow_rate, 2)
        ),
        "T_IN [°C]": t_in,
        "Nominal T_IN [°C]": (
            metadata["nominal_temperature"]
            if metadata["nominal_temperature"] is not None
            else round(t_in, 1)
        ),
        "T_OUT [°C]": t_out,
        "Plateau Duration [s]": step.attrs.get("plateau_duration_s", np.nan),
        "Sample Count": step.attrs.get("sample_count", len(step)),
        **t_cu_averages,
        **lts_averages,
        "T_CU_MIN [°C]": t_cu_min,
        "T_CU_MAX [°C]": t_cu_max,
        "T_CU_AVG [°C]": t_cu_avg,
        "T_CU_deltaLOW [K]": delta_low,
        "T_CU_deltaHIGH [K]": delta_high,
        "DeltaT_CU_MIN [K]": delta_t_cu_min,
        "DeltaT_CU_MAX [K]": delta_t_cu_max,
        "DeltaT_CU [K]": delta_t_cu,
        "Rth [K/W]": rth,
    }

    if metadata["part_type"] == "LTS":
        result["Subcooling [K]"] = subcooling
    elif t_adia is not None:
        result["T_ADIA [°C]"] = t_adia

    return result


def style_excel(
    excel_file,
    t_cu_columns,
    condition_header,
    flow_header,
    medium_types,
    condition_values,
    source_files,
    repeated_comment_flags,
    component_headers,
):
    workbook = load_workbook(excel_file)
    sheet = workbook["Test averages"]
    headers = {cell.value: cell.column for cell in sheet[1]}

    colors = {
        "W_IN [W]": "F4B183", "W_OUT [W]": "FFE699",
        flow_header: "9DC3E6", "T_IN [°C]": "0070C0",
        "T_OUT [°C]": "C65911", "Rth [K/W]": "FCE4D6",
    }
    for column in t_cu_columns + [
        "T_CU_MIN [°C]", "T_CU_MAX [°C]", "T_CU_AVG [°C]"
    ]:
        colors[column] = "DDEBF7"
    for column in headers:
        if re.fullmatch(r'W_CPU_\d+ \[W\]', str(column)):
            colors[column] = "F4B183"
    for column in lts_temperature_output_columns:
        colors[column] = "DDEBF7"
    for column in [
        "T_CU_deltaLOW [K]", "T_CU_deltaHIGH [K]", "DeltaT_CU [K]",
        "Subcooling [K]", "T_ADIA [°C]",
    ]:
        colors[column] = "E2F0D9"

    thin = Side(style="thin", color="B7B7B7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor=colors.get(cell.value, "D9EAF7"))
        cell.border = Border(bottom=thin)

    sheet.cell(1, headers[condition_header]).fill = PatternFill("solid", fgColor="7030A0")
    sheet.cell(1, headers[condition_header]).font = Font(bold=True, color="FFFFFF")

    # Share the report's FR gradient while preserving the reference anchors.
    if condition_header == "FR [%]":
        condition_column = headers[condition_header]
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, condition_column)
            color = interpolated_filling_ratio_color(cell.value)
            if color is not None:
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(color="000000")

    for header in ["T_IN [°C]", "T_OUT [°C]"]:
        sheet.cell(1, headers[header]).font = Font(bold=True, color="FFFFFF")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30

    for column_index in range(1, sheet.max_column + 1):
        maximum = max(
            len(str(sheet.cell(row, column_index).value or ""))
            for row in range(1, min(sheet.max_row, 100) + 1)
        )
        column_header = sheet.cell(1, column_index).value
        maximum_width = 38 if column_header in component_headers else 23
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(maximum + 2, 11),
            maximum_width,
        )
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column_index)
            cell.border = Border(bottom=thin)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, headers["W_IN [W]"]).number_format = "0.0"
        sheet.cell(row, headers["Rth [K/W]"]).number_format = "0.0000"

    # Commented repetitions stay in the Excel data for traceability. Their
    # component cells are highlighted so they cannot be mistaken for the
    # uncommented reference test used by the main report analysis.
    repeated_fill = PatternFill("solid", fgColor="FFF2CC")
    for result_index, is_commented_repeat in enumerate(repeated_comment_flags):
        if not is_commented_repeat:
            continue
        excel_row = result_index + 2
        for component_header in component_headers:
            sheet.cell(
                excel_row,
                headers[component_header],
            ).fill = repeated_fill

    # Add a double line when the filling ratio or charge changes.
    # Add a thick line when the source CSV changes.
    # A condition change takes priority when both happen together.
    double_line = Side(style="double", color="000000")
    thick_line = Side(style="medium", color="000000")

    for result_index in range(1, len(source_files)):
        condition_changed = (
            condition_values[result_index]
            != condition_values[result_index - 1]
        )
        source_changed = (
            source_files[result_index]
            != source_files[result_index - 1]
        )

        if condition_changed:
            top_line = double_line
        elif source_changed:
            top_line = thick_line
        else:
            continue

        excel_row = result_index + 2
        for column_index in range(1, sheet.max_column + 1):
            cell = sheet.cell(excel_row, column_index)
            cell.border = Border(
                left=cell.border.left,
                right=cell.border.right,
                top=top_line,
                bottom=cell.border.bottom,
            )

    if medium_types == {"Water"}:
        sheet.cell(1, headers["T_IN [°C]"]).value = "T_WATER_IN [°C]"
        sheet.cell(1, headers["T_OUT [°C]"]).value = "T_WATER_OUT [°C]"
    elif medium_types == {"Air"}:
        sheet.cell(1, headers["T_IN [°C]"]).value = "T_AIR_IN [°C]"
        sheet.cell(1, headers["T_OUT [°C]"]).value = "T_AIR_OUT [°C]"

    workbook.save(excel_file)


def create_master_excel(
    all_results,
    all_t_cu_columns,
    part_type,
    excel_output_file,
):
    if not all_results:
        # Never manufacture plateau averages for transient-only campaigns.
        with pd.ExcelWriter(excel_output_file, engine='openpyxl') as writer:
            pd.DataFrame({'Information': ['No steady-state averages. Transient (_TR) tests are available in the report and cleaned CSVs.']}).to_excel(
                writer, sheet_name='Test averages', index=False)
            writer.sheets['Test averages'].column_dimensions['A'].width = 110
        return
    results = pd.DataFrame(all_results)
    board_tests = is_board_report(results)
    condition_types = set(results["Condition Type"])
    if condition_types == {"FR"}:
        condition_header = "FR [%]"
    elif condition_types == {"Charge"}:
        condition_header = "Charge"
    else:
        condition_header = "FR [%] / Charge"

    results[condition_header] = results["Condition Value"]
    flow_types = set(results["Flow Type"].dropna())
    if flow_types == {"VFR"}:
        flow_header = "VFR [l/min]"
    elif flow_types == {"CFM"}:
        flow_header = "CFM"
    else:
        flow_header = "Flow Rate"
    results[flow_header] = results["Flow Rate"]

    results = results.sort_values(
        [
            "Working Fluid", "Condition Value", "Orientation",
            "Flow Type", "Nominal Flow Rate", "Nominal T_IN [°C]",
            "Source File", "W_IN [W]",
        ],
        ascending=[True, False, True, True, True, True, True, True],
        kind="stable",
    )

    # Identify commented CSVs that repeat an existing physical condition.
    _, repeated_groups = find_repeated_test_groups(results)
    commented_repeat_indices = set()
    for _, repeated_group in repeated_groups:
        comments = (
            repeated_group["Comment"]
            .fillna("")
            .astype(str)
            .str.strip()
        )
        commented_repeat_indices.update(
            repeated_group.index[comments.ne("")]
        )

    repeated_comment_flags = [
        result_index in commented_repeat_indices
        for result_index in results.index
    ]

    condition_values = results["Condition Value"].tolist()
    source_files = results["Source File"].tolist()

    if part_type == "LTS":
        results["EVAP"] = results["Evaporator Name"]
        results["COND"] = results["Condenser Name"]
        component_columns = ["EVAP", "COND"]
        comment_component_column = "COND"
    else:
        results[part_type] = results["Part Number"]
        component_columns = [part_type]
        comment_component_column = part_type

    if commented_repeat_indices:
        repeated_mask = results.index.isin(commented_repeat_indices)
        results.loc[repeated_mask, comment_component_column] = (
            results.loc[repeated_mask, comment_component_column].astype(str)
            + " - "
            + results.loc[repeated_mask, "Comment"].astype(str).str.strip()
        )

    lts_columns = []
    if part_type == "LTS":
        lts_columns = [
            *lts_temperature_output_columns,
            "Subcooling [K]",
            *psat_result_columns(results),
            *tsat_result_columns(results),
            *superheating_result_columns(results),
        ]

    php_optional_columns = []
    if (
        part_type == "PHP"
        and "T_ADIA [°C]" in results.columns
        and results["T_ADIA [°C]"].notna().any()
    ):
        php_optional_columns.append("T_ADIA [°C]")

    # Match physical CPU IDs, not positions: a test may use CPUs 1, 3, 5, etc.
    # Missing powers remain blank when campaigns have different CPU sets.
    temperature_power_columns = []
    for column in all_t_cu_columns:
        temperature_power_columns.append(column)
        match = re.fullmatch(r'T_(?:CPU|BOARD)_(\d+) \[°C\]', str(column), re.I)
        if match:
            temperature_power_columns.append(f'W_CPU_{int(match[1])} [W]')

    final_columns = [
        *component_columns,
        "Working Fluid", "Orientation", condition_header, "Coolant",
        "W_IN [W]", "W_OUT [W]", flow_header, "T_IN [°C]", "T_OUT [°C]",
        *temperature_power_columns,
        *lts_columns,
        *php_optional_columns,
        "T_CU_MIN [°C]", "T_CU_MAX [°C]", "T_CU_AVG [°C]",
        "T_CU_deltaLOW [K]", "T_CU_deltaHIGH [K]", "DeltaT_CU [K]", "Rth [K/W]",
    ]
    if board_tests:
        # One schema for every CSV. Each plateau occupies one row; new CSVs
        # are stacked by the existing sort, never joined horizontally.
        final_columns.extend([
            "Connected boards", "Power per board [W]", *sorted(
                [c for c in results if re.fullmatch(r'W_CPU_\d+ \[W\]', str(c))
                 and c not in final_columns], key=natural_text_sort_key),
            "Water / scheduled power [%]",
            "Source File", "Step", "Start [s]", "End [s]",
            "Average from [s]", "Average to [s]", "Sample Count",
        ])
        # Retain the original export's supporting values without crowding the
        # main table. Excel's column outline can expand this metadata block.
        final_columns.extend([
            "Nominal Flow Rate", "Nominal T_IN [°C]", "Unrecovered scheduled power [W]",
            "DeltaT_CU_MIN [K]", "DeltaT_CU_MAX [K]", "Comment", "Report Prefix",
            "Part Name", "Part Number", "Condition Type", "Flow Type", "Flow Unit",
            "Plateau Duration [s]", "Duration [s]", "Start method", "Power basis",
        ])
    results = results.reindex(columns=final_columns)

    # mode="w" deliberately overwrites an existing master workbook.
    with pd.ExcelWriter(excel_output_file, engine="openpyxl", mode="w") as writer:
        results.to_excel(writer, sheet_name="Test averages", index=False)

    style_excel(
        excel_output_file,
        all_t_cu_columns,
        condition_header,
        flow_header,
        set(results["Coolant"]),
        condition_values,
        source_files,
        repeated_comment_flags,
        component_columns,
    )
    if board_tests:
        workbook = load_workbook(excel_output_file)
        sheet = workbook["Test averages"]
        display_results = pd.DataFrame(all_results)
        for cell in sheet[1]:
            cell.value = report_display_text(str(cell.value), display_results).replace("DeltaT_CPU", "ΔT_CPU")
            if cell.value == "W_IN [W]":
                cell.value = "Total heat load [W]"
            elif cell.value == "W_OUT [W]":
                cell.value = "Water heat removal [W]"
            if cell.value == "Water / scheduled power [%]":
                cell.value = "Water / total heat load [%]"
            elif cell.value == "Unrecovered scheduled power [W]":
                cell.value = "Heat load not recovered in water [W]"
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 42
        # Keep all audit values, grouped at the right of the familiar table.
        source_column = next(c.column for c in sheet[1] if c.value == "Source File")
        sheet.column_dimensions[get_column_letter(source_column)].width = 36
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row_index, source_column).font = Font(size=8)
            sheet.cell(row_index, source_column).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row_index].height = 30
        metadata_column = next(c.column for c in sheet[1] if c.value == "Nominal Flow Rate")
        sheet.column_dimensions.group(
            get_column_letter(metadata_column), get_column_letter(sheet.max_column),
            outline_level=1, hidden=True,
        )
        sheet.column_dimensions[get_column_letter(sheet.max_column + 1)].collapsed = True
        sheet.freeze_panes = "G2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = sheet.cell(1, cell.column).value
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0000" if header == "Rth [K/W]" else (
                        "0" if header in {"Step", "Connected boards", "Sample Count"} else "0.0"
                    )
        workbook.save(excel_output_file)


def format_number(value, decimals=1):
    """Format a report number without unnecessary trailing zeros."""
    if value is None or pd.isna(value):
        return "N/A"
    formatted = f"{float(value):.{decimals}f}"
    return formatted.rstrip("0").rstrip(".")


def format_unique_values(values, formatter=str):
    """Format unique values as a comma-separated report field."""
    unique_values = []
    for value in values:
        if pd.isna(value):
            continue
        if value not in unique_values:
            unique_values.append(value)
    if not unique_values:
        return "N/A"
    return ", ".join(formatter(value) for value in unique_values)


def orientation_report_name(value):
    """Expand the common one-letter orientation codes for the report."""
    orientation_names = {
        "S": "Sideways",
        "V": "Vertical",
        "H": "Horizontal",
        "N/A": "Not specified",
    }
    return orientation_names.get(str(value).upper(), str(value))


def flow_rate_report_label(flow_type, flow_rate):
    """Format either a liquid VFR or an air-flow CFM condition."""
    if str(flow_type).upper() == "CFM":
        return f"CFM {format_number(flow_rate, 2)}"
    return f"VFR {format_number(flow_rate, 2)} l/min"


def inlet_temperature_report_label(coolant, temperature, compact=False):
    """Format the nominal inlet-temperature condition for water or air."""
    if str(coolant).strip().lower() == "air":
        temperature_name = "T_AIR" if compact else "T_AIR_IN"
    else:
        temperature_name = "TW" if compact else "T_WATER_IN"
    return f"{temperature_name} {format_number(temperature)}°C"


def format_power_range(results):
    values = pd.to_numeric(results["W_IN [W]"], errors="coerce").dropna()
    if values.empty:
        return "N/A"
    minimum = float(values.min())
    maximum = float(values.max())
    if math.isclose(minimum, maximum, abs_tol=0.05):
        return f"{format_number(minimum)} W"
    return f"{format_number(minimum)}-{format_number(maximum)} W"


def format_plateau_duration(results):
    values = pd.to_numeric(
        results["Plateau Duration [s]"], errors="coerce"
    ).dropna()
    if values.empty:
        return "Not available"
    minimum = float(values.min())
    maximum = float(values.max())
    if math.isclose(minimum, maximum, abs_tol=0.5):
        return f"{format_number(minimum)} s"
    return f"{format_number(minimum)}-{format_number(maximum)} s"


def coolant_property_description(coolants):
    coolant_set = set(coolants)
    if coolant_set == {"Water"}:
        return (
            f"Constant: density {water_density:g} kg/m3, "
            f"cp {water_cp:g} J/(kg K)"
        )
    if coolant_set == {"Air"}:
        return (
            f"Constant: density {air_density:g} kg/m3, "
            f"cp {air_cp:g} J/(kg K)"
        )
    return "Constant coolant density and cp"


def refprop_property_description(fluid_properties):
    """Format all available working-fluid values for the report cover."""
    if not fluid_properties:
        return "REFPROP properties unavailable"
    include_fluid = len(fluid_properties) > 1
    return " | ".join(
        format_refprop_property_line(
            fluid,
            properties,
            include_fluid=include_fluid,
        )
        for fluid, properties in sorted(fluid_properties.items())
    )


def draw_wrapped_text(
    pdf,
    text,
    x,
    y,
    maximum_width,
    font_name=report_regular_font,
    font_size=9.5,
    leading=13,
    maximum_lines=None,
):
    """Draw simple wrapped text and return the next y coordinate."""
    pdf.setFont(font_name, font_size)
    words = str(text).split()
    lines = []
    current_line = ""

    for word in words:
        # Split long component names at character level if they contain no
        # spaces and would otherwise overlap the next report field.
        if pdf.stringWidth(word, font_name, font_size) > maximum_width:
            if current_line:
                lines.append(current_line)
                current_line = ""

            chunk = ""
            for character in word:
                candidate_chunk = chunk + character
                if (
                    chunk
                    and pdf.stringWidth(
                        candidate_chunk,
                        font_name,
                        font_size,
                    ) > maximum_width
                ):
                    lines.append(chunk)
                    chunk = character
                else:
                    chunk = candidate_chunk
            current_line = chunk
            continue

        candidate = word if not current_line else f"{current_line} {word}"
        if pdf.stringWidth(candidate, font_name, font_size) <= maximum_width:
            current_line = candidate
        else:
            if current_line:
                lines.append(current_line)
            current_line = word

    if current_line:
        lines.append(current_line)

    if maximum_lines is not None and len(lines) > maximum_lines:
        lines = lines[:maximum_lines]
        last_line = lines[-1]
        while (
            pdf.stringWidth(last_line + "...", font_name, font_size)
            > maximum_width
            and last_line
        ):
            last_line = last_line[:-1]
        lines[-1] = last_line.rstrip() + "..."

    for line in lines:
        pdf.drawString(x, y, line)
        y -= leading

    return y


def draw_section_heading(pdf, text, x, y, width):
    pdf.setFillColor(report_green)
    pdf.setFont(report_bold_font, 12)
    pdf.drawString(x, y, text)
    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(x, y - 7, x + width, y - 7)


def draw_report_field(pdf, label, value, x, y, width):
    pdf.setFillColor(report_green)
    pdf.setFont(report_bold_font, 9)
    pdf.drawString(x, y, label)
    pdf.setFillColor(report_dark)
    draw_wrapped_text(
        pdf,
        value,
        x,
        y - 14,
        width,
        font_size=8 if label == 'REFPROP fluid properties' else 9,
        leading=10 if label == 'REFPROP fluid properties' else 11,
        maximum_lines=4 if label == 'REFPROP fluid properties' else 2,
    )


def draw_page_footer(pdf, page_number, page_width):
    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.5)
    pdf.line(32, 24, page_width - 32, 24)
    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 7.5)
    pdf.drawString(32, 12, "JJ Cooling - Automatically generated test report")
    pdf.drawRightString(page_width - 32, 12, f"Page {page_number}")


def draw_cover_page(
    pdf,
    results,
    part_type,
    has_filling_ratio_analysis,
    has_performance_characterization=False,
    fluid_properties=None,
):
    page_width, page_height = landscape(A4)

    prefixes = format_unique_values(results["Report Prefix"])
    report_title = f"{prefixes} - {part_type} test report"

    pdf.setTitle(report_title)
    pdf.setAuthor("JJ Cooling")

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 18)
    pdf.drawString(32, page_height - 34, report_title)

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 9)
    pdf.drawRightString(
        page_width - 32,
        page_height - 30,
        f"Report date: {date.today().isoformat()}",
    )

    analysis_names = []
    if has_filling_ratio_analysis:
        analysis_names.append("filling-ratio analysis")
    if has_performance_characterization:
        analysis_names.append("performance characterization")
    analysis_name = (
        " and ".join(analysis_names)
        if analysis_names
        else "test-average summary"
    )
    transient_count = int(results.loc[results.get('Test Mode', pd.Series('SS', index=results.index)).eq('TR'), 'Source File'].nunique())
    transient_only = transient_count == results['Source File'].nunique()
    if transient_count:
        analysis_name = ('transient time-series' if transient_only else analysis_name + f'; {transient_count} transient test(s)')
    pdf.setFillColor(report_dark)
    pdf.setFont(report_regular_font, 9.5)
    pdf.drawString(
        32,
        page_height - 75,
        f"Files: {results['Source File'].nunique()} CSV file(s)  |  "
        f"Analysis: {analysis_name}",
    )

    pdf.setStrokeColor(report_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, page_height - 93, page_width - 32, page_height - 93)

    main_info_y = page_height - 109
    draw_section_heading(
        pdf,
        "Main test information",
        32,
        main_info_y,
        page_width - 64,
    )

    field_width = (page_width - 82) / 4
    x_positions = [
        32,
        32 + field_width + 6,
        32 + 2 * (field_width + 6),
        32 + 3 * (field_width + 6),
    ]

    fields_first_row = [
        (
            "Working fluid",
            format_unique_values(results["Working Fluid"]),
        ),
        (
            "Coolant",
            format_unique_values(results["Coolant"]),
        ),
        (
            "Part name",
            format_unique_values(results["Part Name"]),
        ),
        (
            "Orientation",
            format_unique_values(
                results["Orientation"],
                orientation_report_name,
            ),
        ),
    ]

    board_tests = is_board_report(results)
    fields_second_row = [
        ("Total heat load" if board_tests else "Applied power", format_power_range(results)),
        ("Plateau duration", format_plateau_duration(results)),
        ("Averaging window" if board_tests else "Sample size",
         "Raw samples (no plateau averages)" if transient_only else ("Final 100 seconds per plateau" if board_tests else f"{sample_size} values per plateau")),
        (
            "REFPROP fluid properties",
            refprop_property_description(fluid_properties),
        ),
    ]

    for x, (label, value) in zip(x_positions, fields_first_row):
        draw_report_field(pdf, label, value, x, page_height - 139, field_width)

    for x, (label, value) in zip(x_positions, fields_second_row):
        draw_report_field(pdf, label, value, x, page_height - 187, field_width)

    left_x = 32
    right_x = page_width / 2 + 15
    section_width = page_width / 2 - 47
    content_heading_y = page_height - 254

    draw_section_heading(
        pdf,
        "Context & Objectives",
        left_x,
        content_heading_y,
        section_width,
    )

    context_parts = []
    if has_filling_ratio_analysis:
        context_parts.append(
            "Evaluate how filling ratio affects the component thermal response. "
            "The report compares average copper temperature against applied "
            "power for every available coolant inlet temperature and volumetric "
            "flow rate."
        )
        if part_type == "LTS":
            context_parts.append("Subcooling is also compared for the LTS.")
    if has_performance_characterization:
        context_parts.append(
            "For each characterized filling ratio, compare ΔT_CU, average "
            "copper temperature, and thermal resistance across the tested "
            "flow rates, inlet temperatures, and orientations."
        )
        if part_type == "LTS":
            context_parts.append(
                "The LTS performance characterization also compares "
                "subcooling."
            )

    if context_parts:
        context_text = " ".join(context_parts)
    else:
        context_text = (
            "Summarize the end-of-plateau test averages calculated from the "
            "selected CSV files. Additional analysis pages will be added as the "
            "report structure is developed."
        )

    if board_tests:
        context_text = (
            "Characterize the server LTS using board temperatures, water-side heat removal "
            "and subcooling. Total heat load is the sum of the entered per-CPU heat loads; "
            "electrical power is not measured. The thermal start "
            "aligns the schedule. Averages use each step's final 100 seconds."
        )
    if transient_count:
        context_text = (('' if transient_only else context_text + ' ') +
                        'Transient (_TR) tests are plotted against time when selected in report setup; '
                        'they are excluded from steady-state averages and comparisons.')
    pdf.setFillColor(report_dark)
    draw_wrapped_text(
        pdf,
        report_display_text(context_text, results),
        left_x,
        content_heading_y - 27,
        section_width,
        font_size=9.5,
        leading=14,
        maximum_lines=8,
    )

    draw_section_heading(
        pdf,
        "Nomenclature",
        right_x,
        content_heading_y,
        section_width,
    )

    nomenclature = [
        ("T_IN", "Coolant inlet temperature [°C]"),
        ("T_OUT", "Coolant outlet temperature [°C]"),
        ("T_CU,i", "Individual copper temperature [°C]"),
        ("T_CU,AVG", "Average copper temperature [°C]"),
        ("ΔT_CU", "T_CU,AVG - T_IN [K]"),
        ("W_IN", "Applied input power [W]"),
        ("W_OUT", "Heat transferred to the coolant [W]"),
        ("R_th", "Thermal resistance [K/W]"),
        ("FR", "Working-fluid filling ratio [%]"),
        ("d_crit", "PHP capillary critical diameter [mm]"),
        ("P_sat", "Working-fluid saturation pressure [kPa]"),
    ]

    flow_types = set(results["Flow Type"].dropna())
    if "VFR" in flow_types:
        nomenclature.insert(
            6,
            ("VFR", "Coolant volumetric flow rate [l/min]"),
        )
    if "CFM" in flow_types:
        nomenclature.insert(
            7 if "VFR" in flow_types else 6,
            ("CFM", "Air volumetric flow rate [ft3/min]"),
        )

    if part_type == "LTS":
        nomenclature.extend([
            ("T_EVAP_IN", "Evaporator inlet temperature [°C]"),
            ("T_EVAP_OUT", "Evaporator outlet temperature [°C]"),
            ("T_COND_IN", "Condenser inlet temperature [°C]"),
            ("T_COND_OUT", "Condenser outlet temperature [°C]"),
            ("Subcooling", "T_COND_IN - T_COND_OUT [K]"),
            ("Superheating", "T_EVAP_OUT - T_SAT [K]"),
        ])
    elif (
        "T_ADIA [°C]" in results.columns
        and results["T_ADIA [°C]"].notna().any()
    ):
        nomenclature.append(
            ("T_ADIA", "Adiabatic-section temperature [°C]")
        )

    nomenclature_y = content_heading_y - 26
    for symbol, description in nomenclature:
        symbol = report_display_text(symbol, results)
        description = report_display_text(description, results)
        if board_tests and symbol == "W_IN":
            description = "Total heat load [W]"
        pdf.setFillColor(report_green)
        pdf.setFont(report_bold_font, 7.8)
        pdf.drawString(right_x, nomenclature_y, symbol)
        pdf.setFillColor(report_dark)
        pdf.setFont(report_regular_font, 7.8)
        pdf.drawString(right_x + 92, nomenclature_y, description)
        nomenclature_y -= 13

    equations_y = 190
    draw_section_heading(
        pdf,
        "Equations",
        left_x,
        equations_y,
        section_width,
    )

    equations = [
        "T_CU,AVG = (T_CU,1 + T_CU,2 + ... + T_CU,n) / n",
        "W_OUT = density x volume flow x cp x (T_OUT - T_IN)",
        "ΔT_CU = T_CU,AVG - T_IN",
        "R_th = ΔT_CU / W_IN",
    ]

    if part_type == "LTS":
        equations.append("Subcooling = T_COND_IN - T_COND_OUT")

    equation_line_y = equations_y - 31
    pdf.setFillColor(report_dark)
    pdf.setFont(report_regular_font, 9.5)
    for equation in equations:
        draw_wrapped_text(pdf, f"-  {report_display_text(equation, results)}",
                          left_x + 4, equation_line_y, section_width - 8,
                          font_size=9.5, leading=12, maximum_lines=2)
        equation_line_y -= 31

    draw_page_footer(pdf, 1, page_width)


def chart_grid_shape(chart_count):
    """Choose a compact grid so every available test condition fits one page."""
    if chart_count <= 1:
        return 1, 1
    if chart_count <= 2:
        return 1, 2
    if chart_count <= 4:
        return 2, 2
    if chart_count <= 6:
        return 2, 3
    columns = min(4, math.ceil(math.sqrt(chart_count * 1.4)))
    rows = math.ceil(chart_count / columns)
    return rows, columns


def chart_font_sizes(chart_count):
    if chart_count <= 1:
        return 11, 10, 9
    if chart_count <= 4:
        return 9, 8, 7.5
    if chart_count <= 6:
        return 8, 7, 6.5
    return 7, 6, 5.5


def interpolated_filling_ratio_color(filling_ratio):
    """Return an RGB hex color shared by Excel and plots, or None if unknown."""
    try:
        ratio = float(filling_ratio)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(ratio) or not filling_ratio_colors:
        return None
    anchors = sorted(filling_ratio_colors)
    if ratio <= anchors[0]:
        return filling_ratio_colors[anchors[0]]
    if ratio >= anchors[-1]:
        return filling_ratio_colors[anchors[-1]]
    for lower, upper in zip(anchors, anchors[1:]):
        if lower <= ratio <= upper:
            weight = (ratio - lower) / (upper - lower)
            low_color, high_color = filling_ratio_colors[lower], filling_ratio_colors[upper]
            return ''.join(
                f'{round(int(low_color[i:i+2], 16) * (1 - weight) + int(high_color[i:i+2], 16) * weight):02X}'
                for i in (0, 2, 4)
            )


def filling_ratio_line_color(filling_ratio, fallback_index):
    color = interpolated_filling_ratio_color(filling_ratio)
    if color is not None:
        return f"#{color}"
    fallback_colors = plt.get_cmap("tab10").colors
    return fallback_colors[fallback_index % len(fallback_colors)]


def make_analysis_chart_image(
    results,
    y_column,
    spread_columns=None,
):
    """Create all filling-ratio plots for one report page."""
    results = results[results["Condition Type"] == "FR"].copy()

    group_columns = [
        "Working Fluid",
        "Coolant",
        "Orientation",
        "Flow Type",
        "Nominal Flow Rate",
        "Nominal T_IN [°C]",
    ]

    chart_groups = list(
        results.groupby(group_columns, dropna=False, sort=True)
    )

    rows, columns = chart_grid_shape(len(chart_groups))
    title_size, label_size, tick_size = chart_font_sizes(len(chart_groups))

    figure, axes = plt.subplots(
        rows,
        columns,
        figsize=(11.2, 6.55),
        squeeze=False,
    )
    axes_list = axes.flatten()

    unique_fluids = results["Working Fluid"].nunique()

    for chart_index, (group_key, group_data) in enumerate(chart_groups):
        axis = axes_list[chart_index]
        (
            fluid,
            coolant,
            orientation,
            flow_type,
            nominal_flow_rate,
            nominal_temperature,
        ) = group_key

        filling_ratios = sorted(
            group_data.loc[
                group_data["Condition Type"] == "FR",
                "Condition Value",
            ].dropna().unique(),
            reverse=True,
        )

        for ratio_index, filling_ratio in enumerate(filling_ratios):
            ratio_data = group_data[
                (group_data["Condition Type"] == "FR")
                & np.isclose(
                    group_data["Condition Value"].astype(float),
                    float(filling_ratio),
                )
            ]

            # Repeated CSVs are compared separately on the final report page.
            # For the main filling-ratio overview, average their matching heat
            # loads so duplicate x values do not create misleading line loops.
            grouped_ratio_data, _ = add_heat_load_groups(ratio_data)
            aggregation_columns = {
                "W_IN [W]": "mean",
                y_column: "mean",
            }
            if spread_columns is not None:
                aggregation_columns.update({
                    spread_column: "mean"
                    for spread_column in spread_columns
                })

            series = (
                grouped_ratio_data
                .groupby("_Heat Load Group", sort=True)
                .agg(aggregation_columns)
                .sort_values("W_IN [W]")
            )

            line_color = filling_ratio_line_color(
                filling_ratio,
                ratio_index,
            )

            axis.plot(
                series["W_IN [W]"],
                series[y_column],
                color=line_color,
                linewidth=1.5,
                marker="o",
                markersize=3.5 if len(chart_groups) <= 4 else 2.8,
                markerfacecolor="black",
                markeredgecolor="black",
                markeredgewidth=0.4,
                label=format_number(filling_ratio),
            )

            if spread_columns is not None:
                for spread_column in spread_columns:
                    axis.plot(
                        series["W_IN [W]"],
                        series[spread_column],
                        color=line_color,
                        linewidth=1.05,
                        linestyle="--",
                        alpha=0.85,
                        label="_nolegend_",
                    )

        title_parts = []
        if unique_fluids > 1:
            title_parts.append(str(fluid))
        title_parts.append(
            inlet_temperature_report_label(
                coolant,
                nominal_temperature,
                compact=True,
            )
        )
        title_parts.append(
            flow_rate_report_label(flow_type, nominal_flow_rate)
        )
        if str(orientation).upper() != "N/A":
            title_parts.append(orientation_report_name(orientation))

        axis.set_title(" | ".join(title_parts), fontsize=title_size, pad=7)
        axis.set_xlabel(report_display_text("Applied power, W_IN [W]", results), fontsize=label_size)
        axis.set_ylabel(report_metric_label(y_column, results), fontsize=label_size)
        axis.tick_params(axis="both", labelsize=tick_size)
        axis.grid(True, color="#D3D3D3", linewidth=0.6)
        axis.set_axisbelow(True)

        if not group_data["W_IN [W]"].empty:
            axis.set_xlim(left=0)

        axis.legend(
            title="FR [%]",
            fontsize=tick_size,
            title_fontsize=tick_size,
            frameon=False,
            loc="best",
        )

        for spine in axis.spines.values():
            spine.set_color("#A8A8A8")
            spine.set_linewidth(0.7)

    for unused_axis in axes_list[len(chart_groups):]:
        unused_axis.axis("off")

    figure.patch.set_facecolor("white")
    figure.tight_layout(pad=1.1, h_pad=1.3, w_pad=1.1)

    image_buffer = BytesIO()
    figure.savefig(
        image_buffer,
        format="png",
        dpi=190,
        facecolor="white",
        bbox_inches="tight",
    )
    plt.close(figure)
    image_buffer.seek(0)
    return image_buffer


def add_heat_load_groups(group_data):
    """
    Assign nearly equal measured powers to the same nominal heat-load group.

    This keeps values such as 199.8 W and 200.1 W on one 200 W curve when
    W_IN comes from a measured heater signal instead of an exact setpoint.
    """
    grouped_data = group_data.copy()
    powers = pd.to_numeric(grouped_data["W_IN [W]"], errors="coerce")

    if powers.isna().any():
        raise ValueError("W_IN [W] contains non-numeric values")

    clusters = []
    assignments = {}

    for row_index, power in powers.sort_values(kind="stable").items():
        power = float(power)
        selected_cluster = None

        for cluster_index, cluster_values in enumerate(clusters):
            cluster_center = float(np.mean(cluster_values))
            tolerance = max(
                2.0,
                0.015 * max(abs(power), abs(cluster_center), 1.0),
            )
            if abs(power - cluster_center) <= tolerance:
                selected_cluster = cluster_index
                break

        if selected_cluster is None:
            selected_cluster = len(clusters)
            clusters.append([])

        clusters[selected_cluster].append(power)
        assignments[row_index] = selected_cluster

    grouped_data["_Heat Load Group"] = pd.Series(assignments)
    cluster_centers = {
        cluster_index: float(np.mean(cluster_values))
        for cluster_index, cluster_values in enumerate(clusters)
    }
    return grouped_data, cluster_centers


def style_report_axis(
    axis,
    x_label,
    y_label,
    label_size=8,
    tick_size=7,
):
    """Apply the common report-chart style to one Matplotlib axis."""
    axis.set_xlabel(x_label, fontsize=label_size)
    axis.set_ylabel(y_label, fontsize=label_size)
    axis.tick_params(axis="both", labelsize=tick_size)
    axis.grid(True, color="#D3D3D3", linewidth=0.6)
    axis.set_axisbelow(True)
    for spine in axis.spines.values():
        spine.set_color("#A8A8A8")
        spine.set_linewidth(0.7)


def report_metric_label(metric_column, results=None):
    """Return the human-readable PDF label for an internal result column."""
    labels = {
        "DeltaT_CU [K]": "ΔT_CU [K]",
    }
    return report_display_text(labels.get(metric_column, metric_column), results)


def plot_metric_against_power(
    axis,
    group_data,
    metric_column,
    title,
    spread_columns=None,
):
    """Plot one metric against W_IN, with one line per filling ratio."""
    ratios = sorted(
        group_data["Condition Value"].dropna().unique(),
        reverse=True,
    )
    grouped_data, _ = add_heat_load_groups(group_data)

    for ratio_index, filling_ratio in enumerate(ratios):
        ratio_data = grouped_data[
            np.isclose(
                grouped_data["Condition Value"].astype(float),
                float(filling_ratio),
            )
        ]

        # Replicate CSVs at the same heat load are averaged into one point.
        aggregation_columns = {
            "W_IN [W]": "mean",
            metric_column: "mean",
        }
        if spread_columns is not None:
            aggregation_columns.update({
                spread_column: "mean"
                for spread_column in spread_columns
            })

        series = (
            ratio_data
            .groupby("_Heat Load Group", sort=True)
            .agg(aggregation_columns)
            .sort_values("W_IN [W]")
        )

        line_color = filling_ratio_line_color(
            filling_ratio,
            ratio_index,
        )

        axis.plot(
            series["W_IN [W]"],
            series[metric_column],
            color=line_color,
            linewidth=1.5,
            marker="o",
            markersize=3.8,
            markerfacecolor="black",
            markeredgecolor="black",
            markeredgewidth=0.4,
            label=f"{format_number(filling_ratio)}%",
        )

        if spread_columns is not None:
            for spread_column in spread_columns:
                axis.plot(
                    series["W_IN [W]"],
                    series[spread_column],
                    color=line_color,
                    linewidth=1.05,
                    linestyle="--",
                    alpha=0.85,
                    label="_nolegend_",
                )

    axis.set_title(report_display_text(title, group_data), fontsize=9.5, pad=7)
    style_report_axis(
        axis,
        report_display_text("Applied power, W_IN [W]", group_data),
        report_metric_label(metric_column, group_data),
    )
    axis.set_xlim(left=0)
    axis.legend(
        title="Filling ratio",
        fontsize=6.8,
        title_fontsize=7,
        frameon=False,
        loc="best",
    )


def plot_metric_against_filling_ratio(
    axis,
    group_data,
    metric_column,
    title,
    spread_columns=None,
):
    """Plot one metric against filling ratio, with one line per heat load."""
    grouped_data, cluster_centers = add_heat_load_groups(group_data)
    heat_load_colors = plt.get_cmap("tab10").colors

    heat_load_groups = sorted(
        cluster_centers,
        key=lambda group_number: cluster_centers[group_number],
    )

    for color_index, heat_load_group in enumerate(heat_load_groups):
        heat_load_data = grouped_data[
            grouped_data["_Heat Load Group"] == heat_load_group
        ]

        # Replicate CSVs at the same filling ratio are averaged into one point.
        series_columns = [metric_column]
        if spread_columns is not None:
            series_columns.extend(spread_columns)

        series = (
            heat_load_data
            .groupby("Condition Value", sort=True)[series_columns]
            .mean()
            .sort_index()
        )

        heat_load = cluster_centers[heat_load_group]
        line_color = heat_load_colors[
            color_index % len(heat_load_colors)
        ]
        axis.plot(
            series.index.astype(float),
            series[metric_column],
            color=line_color,
            linewidth=1.5,
            marker="o",
            markersize=3.8,
            label=f"{format_number(heat_load)} W",
        )

        if spread_columns is not None:
            for spread_column in spread_columns:
                axis.plot(
                    series.index.astype(float),
                    series[spread_column],
                    color=line_color,
                    linewidth=1.05,
                    linestyle="--",
                    alpha=0.85,
                    label="_nolegend_",
                )

    filling_ratios = sorted(
        float(value)
        for value in grouped_data["Condition Value"].dropna().unique()
    )
    axis.set_xticks(filling_ratios)
    axis.set_title(report_display_text(title, group_data), fontsize=9.5, pad=7)
    style_report_axis(
        axis,
        "Filling ratio [%]",
        report_metric_label(metric_column, group_data),
    )
    axis.legend(
        title="Heat load",
        fontsize=6.8,
        title_fontsize=7,
        frameon=False,
        loc="best",
    )


def make_detailed_filling_ratio_chart_image(group_data):
    """Create the four detailed filling-ratio charts for one condition."""
    figure, axes = plt.subplots(
        2,
        2,
        figsize=(11.2, 6.35),
        squeeze=False,
    )

    plot_metric_against_power(
        axes[0, 0],
        group_data,
        "DeltaT_CU [K]",
        "ΔT_CU vs Applied Power",
        spread_columns=(
            "DeltaT_CU_MIN [K]",
            "DeltaT_CU_MAX [K]",
        ),
    )
    plot_metric_against_power(
        axes[0, 1],
        group_data,
        "Rth [K/W]",
        "Thermal Resistance vs Applied Power",
    )
    plot_metric_against_filling_ratio(
        axes[1, 0],
        group_data,
        "DeltaT_CU [K]",
        "ΔT_CU vs Filling Ratio",
        spread_columns=(
            "DeltaT_CU_MIN [K]",
            "DeltaT_CU_MAX [K]",
        ),
    )
    plot_metric_against_filling_ratio(
        axes[1, 1],
        group_data,
        "Rth [K/W]",
        "Thermal Resistance vs Filling Ratio",
    )

    figure.patch.set_facecolor("white")
    figure.tight_layout(pad=1.15, h_pad=2.0, w_pad=1.45)

    image_buffer = BytesIO()
    figure.savefig(
        image_buffer,
        format="png",
        dpi=190,
        facecolor="white",
        bbox_inches="tight",
    )
    plt.close(figure)
    image_buffer.seek(0)
    return image_buffer


def draw_analysis_page(
    pdf,
    results,
    y_column,
    page_title,
    page_number,
    spread_columns=None,
):
    page_width, page_height = landscape(A4)

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, page_height - 34, report_display_text(page_title, results))

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 8.5)
    pdf.drawRightString(
        page_width - 32,
        page_height - 30,
        f"Report date: {date.today().isoformat()}",
    )

    if spread_columns is not None:
        pdf.setFillColor(report_grey)
        pdf.setFont(report_regular_font, 8.5)
        pdf.drawString(
            32,
            page_height - 51,
            report_display_text("Solid lines: T_CU_AVG  |  Dashed bounds: T_CU_MIN / T_CU_MAX", results),
        )
        divider_y = page_height - 63
        chart_height = page_height - 103
    else:
        divider_y = page_height - 48
        chart_height = page_height - 88

    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, divider_y, page_width - 32, divider_y)

    chart_image = make_analysis_chart_image(
        results,
        y_column,
        spread_columns=spread_columns,
    )
    pdf.drawImage(
        ImageReader(chart_image),
        26,
        31,
        width=page_width - 52,
        height=chart_height,
        preserveAspectRatio=True,
        anchor="c",
    )
    chart_image.close()

    draw_page_footer(pdf, page_number, page_width)


def draw_detailed_filling_ratio_page(
    pdf,
    group_key,
    group_data,
    page_number,
):
    """Draw one four-chart page for one VFR and inlet temperature."""
    page_width, page_height = landscape(A4)
    (
        fluid,
        coolant,
        orientation,
        flow_type,
        nominal_flow_rate,
        nominal_temperature,
    ) = group_key

    page_title = (
        "Filling Ratio Detail - "
        f"{flow_rate_report_label(flow_type, nominal_flow_rate)} - "
        f"{inlet_temperature_report_label(coolant, nominal_temperature)}"
    )

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 15)
    pdf.drawString(32, page_height - 32, page_title)

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 8.5)
    pdf.drawRightString(
        page_width - 32,
        page_height - 29,
        f"Report date: {date.today().isoformat()}",
    )

    subtitle_parts = [
        f"Working fluid: {fluid}",
        f"Coolant: {coolant}",
        "Dashed bounds: ΔT_CU_MIN / ΔT_CU_MAX",
    ]
    if str(orientation).upper() != "N/A":
        subtitle_parts.append(
            f"Orientation: {orientation_report_name(orientation)}"
        )

    pdf.setFont(report_regular_font, 8.5)
    pdf.drawString(32, page_height - 49, report_display_text("  |  ".join(subtitle_parts), group_data))

    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, page_height - 61, page_width - 32, page_height - 61)

    chart_image = make_detailed_filling_ratio_chart_image(group_data)
    pdf.drawImage(
        ImageReader(chart_image),
        26,
        31,
        width=page_width - 52,
        height=page_height - 101,
        preserveAspectRatio=True,
        anchor="c",
    )
    chart_image.close()

    draw_page_footer(pdf, page_number, page_width)


def summary_condition_keys(results):
    """Return stable, sorted flow/temperature conditions for summary tables."""
    condition_columns = [
        "Working Fluid",
        "Coolant",
        "Orientation",
        "Flow Type",
        "Nominal Flow Rate",
        "Nominal T_IN [°C]",
    ]
    conditions = results[condition_columns].drop_duplicates().copy()
    conditions = conditions.sort_values(
        condition_columns,
        ascending=True,
        kind="stable",
    )
    keys = [
        tuple(row[column] for column in condition_columns)
        for _, row in conditions.iterrows()
    ]
    return condition_columns, keys


def condition_key_mask(data, condition_columns, condition_key):
    """Build a row mask for one complete report condition key."""
    mask = pd.Series(True, index=data.index)
    for column, value in zip(condition_columns, condition_key):
        if pd.isna(value):
            mask &= data[column].isna()
        elif isinstance(value, (int, float, np.integer, np.floating)):
            mask &= np.isclose(
                pd.to_numeric(data[column], errors="coerce"),
                float(value),
            )
        else:
            mask &= data[column].astype(str) == str(value)
    return mask


def summary_condition_label(
    condition_key,
    show_fluid,
    show_orientation,
):
    """Create the multi-line heading used above one table condition."""
    (
        fluid,
        coolant,
        orientation,
        flow_type,
        nominal_flow_rate,
        nominal_temperature,
    ) = condition_key

    lines = [
        flow_rate_report_label(flow_type, nominal_flow_rate),
        inlet_temperature_report_label(
            coolant,
            nominal_temperature,
            compact=True,
        ),
    ]

    qualifiers = []
    if show_fluid:
        qualifiers.append(str(fluid))
    if show_orientation and str(orientation).upper() != "N/A":
        qualifiers.append(orientation_report_name(orientation))
    if qualifiers:
        lines.append(" / ".join(qualifiers))

    return "\n".join(lines)


def fit_report_cell_font(pdf, lines, font_name, font_size, maximum_width):
    """Reduce a table-cell font until every line fits horizontally."""
    fitted_size = font_size
    while fitted_size > 4.0:
        widest_line = max(
            pdf.stringWidth(line, font_name, fitted_size)
            for line in lines
        )
        if widest_line <= maximum_width:
            break
        fitted_size -= 0.25
    return fitted_size


def draw_summary_cell(
    pdf,
    x,
    top_y,
    width,
    height,
    text,
    fill_color,
    font_name=report_regular_font,
    font_size=7,
    text_color=report_dark,
):
    """Draw one bordered summary-table cell with centered, fitted text."""
    pdf.setFillColor(fill_color)
    pdf.setStrokeColor(reportlab_colors.HexColor("#B8C0C4"))
    pdf.setLineWidth(0.45)
    pdf.rect(x, top_y - height, width, height, fill=1, stroke=1)

    lines = str(text).split("\n") if text != "" else [""]
    # Keep cell padding deliberately small so compact summary tables remain
    # readable without wasting vertical or horizontal space.
    maximum_text_width = max(1.0, width - 2)
    fitted_size = fit_report_cell_font(
        pdf,
        lines,
        font_name,
        font_size,
        maximum_text_width,
    )
    fitted_size = min(
        fitted_size,
        max(3.5, height / (len(lines) * 1.2)),
    )
    leading = fitted_size
    text_block_height = leading * (len(lines) - 1)
    text_y = top_y - height / 2 + text_block_height / 2 - fitted_size * 0.33

    pdf.setFillColor(text_color)
    pdf.setFont(font_name, fitted_size)
    for line in lines:
        pdf.drawCentredString(x + width / 2, text_y, line)
        text_y -= leading


def build_filling_ratio_summary_rows(
    ratio_data,
    condition_columns,
    condition_keys,
    metric_columns,
):
    """Build aligned heat-load rows for one filling-ratio table."""
    grouped_data, cluster_centers = add_heat_load_groups(ratio_data)
    heat_load_groups = sorted(
        cluster_centers,
        key=lambda group_number: cluster_centers[group_number],
    )
    table_rows = []

    for heat_load_group in heat_load_groups:
        row = [cluster_centers[heat_load_group]]
        load_data = grouped_data[
            grouped_data["_Heat Load Group"] == heat_load_group
        ]

        for condition_key in condition_keys:
            condition_data = load_data[
                condition_key_mask(
                    load_data,
                    condition_columns,
                    condition_key,
                )
            ]
            for metric_column in metric_columns:
                if metric_column not in condition_data.columns:
                    row.append(np.nan)
                    continue
                values = pd.to_numeric(
                    condition_data[metric_column],
                    errors="coerce",
                ).dropna()
                row.append(float(values.mean()) if not values.empty else np.nan)

        table_rows.append(row)

    return table_rows


def summary_metric_header(metric_column):
    headers = {
        "DeltaT_CU [K]": "ΔT_CU\n[K]",
        "Subcooling [K]": "Subcooling\n[K]",
        "T_ADIA [°C]": "T_ADIA\n[°C]",
        "W_OUT [W]": "Water heat\n[W]",
        "Water / scheduled power [%]": "Water / W_IN\n[%]",
    }
    if metric_column in psat_result_columns([metric_column]):
        return metric_column.replace(' [bar abs]', '\n[bar abs]')
    return headers.get(metric_column, metric_column)


def summary_metric_value(value, metric_column):
    if value is None or pd.isna(value):
        return ""
    return f"{float(value):.1f}"


def draw_one_filling_ratio_table(
    pdf,
    ratio,
    table_rows,
    condition_keys,
    metric_columns,
    top_y,
    data_row_height,
    show_fluid,
    show_orientation,
    page_width,
    results=None,
):
    """Draw one filling-ratio table inside its allocated page block."""
    left_x = 32
    table_width = page_width - 64
    ratio_label_height = 10
    table_top = top_y - ratio_label_height

    # Fixed compact headers prevent the table from expanding simply because
    # more page space happens to be available.
    header_condition_height = 17.0
    header_metric_height = 13.0

    metric_count = len(metric_columns)
    total_column_count = 1 + len(condition_keys) * metric_count
    power_width = min(68.0, max(47.0, table_width * 0.10))
    metric_width = (
        (table_width - power_width) / max(1, total_column_count - 1)
    )

    ratio_color = filling_ratio_line_color(ratio, 0)
    if isinstance(ratio_color, str):
        ratio_fill = reportlab_colors.HexColor(ratio_color)
    else:
        ratio_fill = reportlab_colors.Color(*ratio_color[:3])
    pdf.setFillColor(ratio_fill)
    pdf.rect(left_x, top_y - 8, 6, 6, fill=1, stroke=0)
    pdf.setFillColor(report_green)
    pdf.setFont(report_bold_font, 7.5)
    pdf.drawString(left_x + 10, top_y - 7, f"FR {format_number(ratio)}%")

    header_green = reportlab_colors.HexColor("#2F6B3B")
    header_light = reportlab_colors.HexColor("#D5DFD8")
    header_grey = reportlab_colors.HexColor("#D9E0E3")
    white = reportlab_colors.white

    draw_summary_cell(
        pdf,
        left_x,
        table_top,
        power_width,
        header_condition_height + header_metric_height,
        report_display_text("W_PSU\n[W]", results),
        header_green,
        font_name=report_bold_font,
        font_size=6.3,
        text_color=white,
    )

    x = left_x + power_width
    for condition_key in condition_keys:
        condition_width = metric_width * metric_count
        draw_summary_cell(
            pdf,
            x,
            table_top,
            condition_width,
            header_condition_height,
            summary_condition_label(
                condition_key,
                show_fluid,
                show_orientation,
            ),
            header_light,
            font_name=report_bold_font,
            font_size=6.2,
        )
        for metric_index, metric_column in enumerate(metric_columns):
            draw_summary_cell(
                pdf,
                x + metric_index * metric_width,
                table_top - header_condition_height,
                metric_width,
                header_metric_height,
                report_display_text(summary_metric_header(metric_column), results),
                header_grey,
                font_name=report_bold_font,
                font_size=5.7,
            )
        x += condition_width

    data_top = table_top - header_condition_height - header_metric_height
    for row_index, row_values in enumerate(table_rows):
        row_fill = (
            reportlab_colors.white
            if row_index % 2 == 0
            else reportlab_colors.HexColor("#F1F4F6")
        )
        row_top = data_top - row_index * data_row_height
        draw_summary_cell(
            pdf,
            left_x,
            row_top,
            power_width,
            data_row_height,
            f"{float(row_values[0]):.1f}",
            row_fill,
            font_name=report_bold_font,
            font_size=6.2,
        )

        x = left_x + power_width
        value_index = 1
        for _ in condition_keys:
            for metric_column in metric_columns:
                draw_summary_cell(
                    pdf,
                    x,
                    row_top,
                    metric_width,
                    data_row_height,
                    summary_metric_value(
                        row_values[value_index],
                        metric_column,
                    ),
                    row_fill,
                    font_size=6.0,
                )
                x += metric_width
                value_index += 1

    return (
        ratio_label_height
        + header_condition_height
        + header_metric_height
        + data_row_height * max(1, len(table_rows))
    )


def draw_filling_ratio_summary_page(
    pdf,
    results,
    part_type,
    page_number,
):
    """Draw all filling-ratio test-summary tables on one landscape page."""
    page_width, page_height = landscape(A4)
    filling_ratio_results = results[
        results["Condition Type"] == "FR"
    ].copy()

    metric_columns = ["DeltaT_CU [K]"]
    if part_type == "LTS":
        metric_columns.append("Subcooling [K]")
        metric_columns.extend(c for c in psat_result_columns(filling_ratio_results)
                              if filling_ratio_results[c].notna().any())
    elif (
        "T_ADIA [°C]" in filling_ratio_results.columns
        and filling_ratio_results["T_ADIA [°C]"].notna().any()
    ):
        metric_columns.append("T_ADIA [°C]")

    if is_board_report(results):
        metric_columns.extend(["W_OUT [W]", "Water / scheduled power [%]"])

    condition_columns, condition_keys = summary_condition_keys(
        filling_ratio_results
    )
    ratios = sorted(
        filling_ratio_results["Condition Value"].dropna().unique(),
        reverse=True,
    )

    table_payloads = []
    for ratio in ratios:
        ratio_data = filling_ratio_results[
            np.isclose(
                filling_ratio_results["Condition Value"].astype(float),
                float(ratio),
            )
        ]
        table_payloads.append((
            ratio,
            build_filling_ratio_summary_rows(
                ratio_data,
                condition_columns,
                condition_keys,
                metric_columns,
            ),
        ))

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, page_height - 34, "Filling Ratio Test Summaries")

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 8.5)
    pdf.drawRightString(
        page_width - 32,
        page_height - 30,
        f"Report date: {date.today().isoformat()}",
    )

    metric_description = "ΔT_CU"
    if part_type == "LTS":
        metric_description += " and Subcooling"
        if psat_result_columns(metric_columns):
            metric_description += ", Psat [bar abs]"
    elif len(metric_columns) > 1:
        metric_description += " and T_ADIA"
    pdf.drawString(
        32,
        page_height - 51,
        report_display_text(f"{metric_description} for every tested flow-rate and inlet-temperature condition.", results),
    )

    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, page_height - 63, page_width - 32, page_height - 63)

    content_top = page_height - 72
    content_bottom = 31
    table_gap = 3
    ratio_count = max(1, len(table_payloads))
    content_height = content_top - content_bottom
    total_data_rows = sum(
        max(1, len(table_rows))
        for _, table_rows in table_payloads
    )
    fixed_height = ratio_count * (10.0 + 17.0 + 13.0)
    available_row_height = (
        content_height
        - fixed_height
        - table_gap * (ratio_count - 1)
    ) / max(1, total_data_rows)
    data_row_height = min(10.0, max(3.0, available_row_height))

    show_fluid = filling_ratio_results["Working Fluid"].nunique() > 1
    show_orientation = filling_ratio_results["Orientation"].nunique() > 1

    block_top = content_top
    for ratio, table_rows in table_payloads:
        used_height = draw_one_filling_ratio_table(
            pdf,
            ratio,
            table_rows,
            condition_keys,
            metric_columns,
            block_top,
            data_row_height,
            show_fluid,
            show_orientation,
            page_width,
            results=results,
        )
        block_top -= used_height + table_gap

    draw_page_footer(pdf, page_number, page_width)


def find_repeated_test_groups(results):
    """Return condition groups represented by two or more source CSV files."""
    condition_columns = [
        "Part Name",
        "Part Number",
        "Working Fluid",
        "Orientation",
        "Condition Type",
        "Condition Value",
        "Coolant",
        "Flow Type",
        "Nominal Flow Rate",
        "Nominal T_IN [°C]",
    ]

    repeated_groups = []
    for group_key, group_data in results.groupby(
        condition_columns,
        dropna=False,
        sort=True,
    ):
        if group_data["Source File"].nunique() >= 2:
            repeated_groups.append((group_key, group_data.copy()))

    return condition_columns, repeated_groups


def primary_report_results(results):
    """
    Exclude commented repeats from the main report when a reference exists.

    The complete data remains available for the repeated-tests page and the
    Excel workbook. A commented file is removed only from its duplicated
    physical condition and only when that condition also has at least one
    uncommented reference CSV.
    """
    primary_results = results.copy()
    _, repeated_groups = find_repeated_test_groups(primary_results)
    excluded_indices = []

    for _, group_data in repeated_groups:
        comments = group_data["Comment"].fillna("").astype(str).str.strip()
        if comments.eq("").any():
            excluded_indices.extend(group_data.index[comments.ne("")])

    if excluded_indices:
        primary_results = primary_results.drop(index=excluded_indices)

    return primary_results


def performance_condition_columns():
    """Columns defining one performance-characterization condition."""
    return [
        "Working Fluid",
        "Coolant",
        "Flow Type",
        "Nominal Flow Rate",
        "Nominal T_IN [°C]",
        "Orientation",
    ]


def find_performance_characterization_groups(results):
    """
    Find filling ratios tested at two or more distinct operating conditions.

    Multiple power plateaus and repeated CSVs at the same physical condition
    do not trigger this analysis. At least one of flow rate, inlet temperature,
    orientation, or coolant must differ.
    """
    if results.empty:
        return []

    filling_ratio_results = results[
        results["Condition Type"] == "FR"
    ].copy()
    groups = []
    ratios = sorted(
        filling_ratio_results["Condition Value"].dropna().unique(),
        reverse=True,
    )
    detection_columns = [
        "Coolant",
        "Flow Type",
        "Nominal Flow Rate",
        "Nominal T_IN [°C]",
        "Orientation",
    ]

    for ratio in ratios:
        ratio_data = filling_ratio_results[
            np.isclose(
                pd.to_numeric(
                    filling_ratio_results["Condition Value"],
                    errors="coerce",
                ),
                float(ratio),
            )
        ].copy()
        unique_conditions = ratio_data[
            detection_columns
        ].drop_duplicates()
        if len(unique_conditions) >= 2:
            groups.append((float(ratio), ratio_data))

    return groups


def performance_condition_variations(condition_table):
    """Return which condition dimensions vary across one FR page."""
    return {
        "fluid": condition_table[["Working Fluid"]].drop_duplicates().shape[0] > 1,
        "coolant_temperature": (
            condition_table[["Coolant", "Nominal T_IN [°C]"]]
            .drop_duplicates()
            .shape[0]
            > 1
        ),
        "flow": (
            condition_table[["Flow Type", "Nominal Flow Rate"]]
            .drop_duplicates()
            .shape[0]
            > 1
        ),
        "orientation": (
            condition_table[["Orientation"]].drop_duplicates().shape[0] > 1
        ),
    }


def performance_common_condition_title(
    filling_ratio,
    condition_table,
    variations,
):
    """Place every constant operating condition in the plot titles."""
    first_condition = condition_table.iloc[0]
    title_parts = [f"FR {format_number(filling_ratio)}%"]

    if not variations["fluid"]:
        title_parts.append(str(first_condition["Working Fluid"]))
    if not variations["flow"]:
        title_parts.append(
            flow_rate_report_label(
                first_condition["Flow Type"],
                first_condition["Nominal Flow Rate"],
            )
        )
    if not variations["coolant_temperature"]:
        title_parts.append(
            inlet_temperature_report_label(
                first_condition["Coolant"],
                first_condition["Nominal T_IN [°C]"],
                compact=True,
            )
        )
    if not variations["orientation"]:
        orientation = first_condition["Orientation"]
        if str(orientation).upper() != "N/A":
            title_parts.append(orientation_report_name(orientation))

    return " | ".join(title_parts)


def performance_condition_legend_label(
    condition_key,
    variations,
    fallback_number,
):
    """Put only varying operating conditions in one legend label."""
    (
        fluid,
        coolant,
        flow_type,
        nominal_flow_rate,
        nominal_temperature,
        orientation,
    ) = condition_key

    label_parts = []
    if variations["fluid"]:
        label_parts.append(str(fluid))
    if variations["flow"]:
        label_parts.append(
            flow_rate_report_label(flow_type, nominal_flow_rate)
        )
    if variations["coolant_temperature"]:
        label_parts.append(
            inlet_temperature_report_label(
                coolant,
                nominal_temperature,
                compact=True,
            )
        )
    if variations["orientation"]:
        label_parts.append(orientation_report_name(orientation))

    return (
        " | ".join(label_parts)
        if label_parts
        else f"Condition {fallback_number}"
    )


def performance_temperature_color(temperature):
    """Return the fixed 10-60°C blue-to-red performance color."""
    anchor_temperatures = sorted(performance_temperature_color_scale)
    anchor_colors = [
        performance_temperature_color_scale[value]
        for value in anchor_temperatures
    ]
    color_map = matplotlib.colors.LinearSegmentedColormap.from_list(
        "performance_temperature",
        list(zip(
            np.linspace(0.0, 1.0, len(anchor_temperatures)),
            anchor_colors,
        )),
    )
    normalizer = matplotlib.colors.Normalize(
        vmin=float(anchor_temperatures[0]),
        vmax=float(anchor_temperatures[-1]),
        clip=True,
    )
    numeric_temperature = pd.to_numeric(
        pd.Series([temperature]),
        errors="coerce",
    ).iloc[0]
    if pd.isna(numeric_temperature):
        numeric_temperature = float(anchor_temperatures[0])
    return color_map(normalizer(float(numeric_temperature)))


def adjust_performance_color_for_flow(base_color, flow_position, flow_count):
    """Make low-flow curves lighter and high-flow curves darker."""
    red, green, blue = matplotlib.colors.to_rgb(base_color)
    if flow_count <= 1:
        return red, green, blue

    normalized_position = flow_position / max(1, flow_count - 1)
    if normalized_position <= 0.5:
        lightening = performance_low_flow_lightening * (
            1.0 - 2.0 * normalized_position
        )
        return tuple(
            component + (1.0 - component) * lightening
            for component in (red, green, blue)
        )

    darkening = performance_high_flow_darkening * (
        2.0 * normalized_position - 1.0
    )
    return tuple(
        component * (1.0 - darkening)
        for component in (red, green, blue)
    )


def performance_flow_key(flow_type, flow_rate):
    """Return a sortable, stable key for a VFR or CFM value."""
    numeric_flow_rate = pd.to_numeric(
        pd.Series([flow_rate]),
        errors="coerce",
    ).iloc[0]
    numeric_flow_rate = (
        float(numeric_flow_rate)
        if not pd.isna(numeric_flow_rate)
        else float("inf")
    )
    return str(flow_type).upper(), numeric_flow_rate


def performance_orientation_marker(orientation):
    """Map component orientation to a consistent plot marker."""
    normalized_orientation = str(orientation).strip().upper()
    orientation_markers = {
        "H": "x",
        "HORIZONTAL": "x",
        "V": "o",
        "VERTICAL": "o",
        "S": "s",
        "SIDE": "s",
        "SIDEWAYS": "s",
        "U": "^",
        "UP": "^",
        "UPWARDS": "^",
        "D": "v",
        "DOWN": "v",
        "DOWNWARDS": "v",
        "N/A": "D",
    }
    return orientation_markers.get(normalized_orientation, "D")


def plot_performance_series(
    axis,
    series,
    value_column,
    line_color,
    marker,
    legend_label,
    spread_columns=None,
):
    """Plot one condition and optional same-color dashed spread bounds."""
    marker_face_color = (
        line_color
        if marker in {"x", "+"}
        else "white"
    )
    axis.plot(
        series["W_IN [W]"],
        series[value_column],
        color=line_color,
        linewidth=1.55,
        marker=marker,
        markersize=4.2,
        markerfacecolor=marker_face_color,
        markeredgecolor=line_color,
        markeredgewidth=1.0,
        label=legend_label,
    )
    for spread_column in spread_columns or []:
        axis.plot(
            series["W_IN [W]"],
            series[spread_column],
            color=line_color,
            linewidth=0.95,
            linestyle="--",
            alpha=0.82,
            label="_nolegend_",
        )


def make_performance_characterization_chart_image(
    filling_ratio,
    ratio_data,
):
    """Create all performance curves for every condition at one FR."""
    condition_columns = performance_condition_columns()
    condition_table = (
        ratio_data[condition_columns]
        .drop_duplicates()
        .sort_values(
            [
                "Flow Type",
                "Nominal Flow Rate",
                "Nominal T_IN [°C]",
                "Orientation",
                "Working Fluid",
                "Coolant",
            ],
            kind="stable",
        )
    )
    condition_keys = [
        tuple(row[column] for column in condition_columns)
        for _, row in condition_table.iterrows()
    ]
    variations = performance_condition_variations(condition_table)
    common_title = performance_common_condition_title(
        filling_ratio,
        condition_table,
        variations,
    )

    has_subcooling = (
        "Subcooling [K]" in ratio_data.columns
        and pd.to_numeric(
            ratio_data["Subcooling [K]"],
            errors="coerce",
        ).notna().any()
    )

    figure, axes = plt.subplots(
        2,
        2,
        figsize=(11.2, 6.25),
        squeeze=False,
    )
    delta_axis = axes[0, 0]
    resistance_axis = axes[0, 1]
    average_temperature_axis = axes[1, 0]
    subcooling_axis = axes[1, 1] if has_subcooling else None

    flow_keys = sorted({
        performance_flow_key(condition_key[2], condition_key[3])
        for condition_key in condition_keys
    })
    flow_positions = {
        flow_key: position
        for position, flow_key in enumerate(flow_keys)
    }

    for condition_index, condition_key in enumerate(condition_keys):
        condition_data = ratio_data[
            condition_key_mask(
                ratio_data,
                condition_columns,
                condition_key,
            )
        ].copy()
        grouped_condition, _ = add_heat_load_groups(condition_data)
        aggregation_columns = {
            "W_IN [W]": "mean",
            "DeltaT_CU [K]": "mean",
            "DeltaT_CU_MIN [K]": "mean",
            "DeltaT_CU_MAX [K]": "mean",
            "Rth [K/W]": "mean",
            "T_CU_AVG [°C]": "mean",
            "T_CU_MIN [°C]": "mean",
            "T_CU_MAX [°C]": "mean",
        }
        if has_subcooling:
            aggregation_columns["Subcooling [K]"] = "mean"
        series = (
            grouped_condition
            .groupby("_Heat Load Group", sort=True)
            .agg(aggregation_columns)
            .sort_values("W_IN [W]")
        )
        flow_key = performance_flow_key(
            condition_key[2],
            condition_key[3],
        )
        base_color = performance_temperature_color(condition_key[4])
        line_color = adjust_performance_color_for_flow(
            base_color,
            flow_positions[flow_key],
            len(flow_keys),
        )
        marker = performance_orientation_marker(condition_key[5])
        legend_label = performance_condition_legend_label(
            condition_key,
            variations,
            condition_index + 1,
        )

        plot_performance_series(
            delta_axis,
            series,
            "DeltaT_CU [K]",
            line_color,
            marker,
            legend_label,
            spread_columns=[
                "DeltaT_CU_MIN [K]",
                "DeltaT_CU_MAX [K]",
            ],
        )
        plot_performance_series(
            resistance_axis,
            series,
            "Rth [K/W]",
            line_color,
            marker,
            legend_label,
        )
        plot_performance_series(
            average_temperature_axis,
            series,
            "T_CU_AVG [°C]",
            line_color,
            marker,
            legend_label,
            spread_columns=[
                "T_CU_MIN [°C]",
                "T_CU_MAX [°C]",
            ],
        )
        if has_subcooling:
            plot_performance_series(
                subcooling_axis,
                series,
                "Subcooling [K]",
                line_color,
                marker,
                legend_label,
            )

    delta_axis.set_title(
        report_display_text(f"ΔT_CU vs Applied Power\n{common_title}", ratio_data),
        fontsize=8.7,
        pad=6,
    )
    resistance_axis.set_title(
        report_display_text(f"Rth vs Applied Power\n{common_title}", ratio_data),
        fontsize=8.7,
        pad=6,
    )
    average_temperature_axis.set_title(
        report_display_text(f"T_CU_AVG vs Applied Power\n{common_title}", ratio_data),
        fontsize=8.7,
        pad=6,
    )
    if has_subcooling:
        subcooling_axis.set_title(
            report_display_text(f"Subcooling vs Applied Power\n{common_title}", ratio_data),
            fontsize=8.7,
            pad=6,
        )
    else:
        axes[1, 1].axis("off")

    style_report_axis(
        delta_axis,
        report_display_text("Applied power, W_IN [W]", ratio_data),
        report_display_text("ΔT_CU [K]", ratio_data),
        label_size=8,
        tick_size=7,
    )
    style_report_axis(
        resistance_axis,
        report_display_text("Applied power, W_IN [W]", ratio_data),
        "Rth [K/W]",
        label_size=8,
        tick_size=7,
    )
    style_report_axis(
        average_temperature_axis,
        report_display_text("Applied power, W_IN [W]", ratio_data),
        report_display_text("T_CU_AVG [°C]", ratio_data),
        label_size=8,
        tick_size=7,
    )
    active_axes = [
        delta_axis,
        resistance_axis,
        average_temperature_axis,
    ]
    if has_subcooling:
        style_report_axis(
            subcooling_axis,
            report_display_text("Applied power, W_IN [W]", ratio_data),
            "Subcooling [K]",
            label_size=8,
            tick_size=7,
        )
        active_axes.append(subcooling_axis)

    legend_columns = 2 if len(condition_keys) > 5 else 1
    for axis in active_axes:
        axis.set_xlim(left=0)
        axis.legend(
            fontsize=5.8,
            frameon=False,
            loc="best",
            ncol=legend_columns,
        )

    figure.patch.set_facecolor("white")
    figure.tight_layout(pad=1.0, h_pad=1.1, w_pad=1.4)

    image_buffer = BytesIO()
    figure.savefig(
        image_buffer,
        format="png",
        dpi=190,
        facecolor="white",
        bbox_inches="tight",
    )
    plt.close(figure)
    image_buffer.seek(0)
    return image_buffer


def draw_performance_characterization_page(
    pdf,
    filling_ratio,
    ratio_data,
    page_number,
    page_title=None,
):
    """Draw one performance page for a single filling ratio."""
    page_width, page_height = landscape(A4)
    page_title = page_title or (
        "Performance Characterization - "
        f"FR {format_number(filling_ratio)}%"
    )

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, page_height - 34, page_title)

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 8.5)
    pdf.drawRightString(
        page_width - 32,
        page_height - 30,
        f"Report date: {date.today().isoformat()}",
    )
    pdf.drawString(
        32,
        page_height - 51,
        "Color: inlet temperature (10°C blue to 60°C red)  |  "
        "Shade: flow rate (lower is lighter)  |  Marker: orientation",
    )
    pdf.drawString(
        32,
        page_height - 62,
        report_display_text("Dashed bounds show ΔT_CU_MIN/MAX and T_CU_MIN/MAX.", ratio_data),
    )

    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, page_height - 72, page_width - 32, page_height - 72)

    chart_image = make_performance_characterization_chart_image(
        filling_ratio,
        ratio_data,
    )
    pdf.drawImage(
        ImageReader(chart_image),
        26,
        31,
        width=page_width - 52,
        height=page_height - 112,
        preserveAspectRatio=True,
        anchor="c",
    )
    chart_image.close()

    draw_page_footer(pdf, page_number, page_width)


def repeated_condition_label(group_key, show_fluid, show_orientation):
    """Build a concise chart title for one repeated physical condition."""
    (
        _part_name,
        _part_number,
        fluid,
        orientation,
        condition_type,
        condition_value,
        coolant,
        flow_type,
        nominal_flow_rate,
        nominal_temperature,
    ) = group_key

    if condition_type == "FR":
        condition_text = f"FR {format_number(condition_value)}%"
    else:
        condition_text = f"Charge {format_number(condition_value)}"

    label_parts = [
        condition_text,
        flow_rate_report_label(flow_type, nominal_flow_rate),
        inlet_temperature_report_label(
            coolant,
            nominal_temperature,
            compact=True,
        ),
    ]
    if show_fluid:
        label_parts.insert(0, str(fluid))
    if show_orientation and str(orientation).upper() != "N/A":
        label_parts.append(orientation_report_name(orientation))

    return " | ".join(label_parts)


def repeated_source_labels(group_data):
    """Map each source CSV to a unique legend label based on its SS comment."""
    source_rows = (
        group_data[["Source File", "Comment"]]
        .drop_duplicates(subset=["Source File"])
        .sort_values("Source File", kind="stable")
    )
    used_labels = {}
    labels = {}

    for _, source_row in source_rows.iterrows():
        source_file = source_row["Source File"]
        comment = source_row["Comment"]
        base_label = (
            str(comment).strip()
            if comment is not None and not pd.isna(comment) and str(comment).strip()
            else "reference"
        )
        used_labels[base_label] = used_labels.get(base_label, 0) + 1
        occurrence = used_labels[base_label]
        labels[source_file] = (
            base_label if occurrence == 1 else f"{base_label} ({occurrence})"
        )

    return labels


def plot_repeated_test_metric(
    axis,
    group_data,
    metric_column,
    condition_label,
    title_size,
    label_size,
    tick_size,
):
    """Plot one line per repeated source file using its SS comment."""
    source_labels = repeated_source_labels(group_data)
    line_colors = plt.get_cmap("tab10").colors

    for source_index, (source_file, source_data) in enumerate(
        group_data.groupby("Source File", sort=True)
    ):
        numeric_metric = pd.to_numeric(
            source_data[metric_column],
            errors="coerce",
        )
        valid_source_data = source_data.loc[numeric_metric.notna()].copy()
        if valid_source_data.empty:
            continue

        grouped_source_data, _ = add_heat_load_groups(valid_source_data)
        aggregation_columns = {
            "W_IN [W]": "mean",
            metric_column: "mean",
        }
        spread_columns = None
        if metric_column == "DeltaT_CU [K]":
            candidate_spread_columns = (
                "DeltaT_CU_MIN [K]",
                "DeltaT_CU_MAX [K]",
            )
            if all(
                column in grouped_source_data.columns
                for column in candidate_spread_columns
            ):
                spread_columns = candidate_spread_columns
                aggregation_columns.update({
                    spread_column: "mean"
                    for spread_column in spread_columns
                })

        series = (
            grouped_source_data
            .groupby("_Heat Load Group", sort=True)
            .agg(aggregation_columns)
            .sort_values("W_IN [W]")
        )
        line_color = line_colors[source_index % len(line_colors)]
        axis.plot(
            series["W_IN [W]"],
            series[metric_column],
            color=line_color,
            linewidth=1.5,
            marker="o",
            markersize=3.5,
            label=source_labels[source_file],
        )

        if spread_columns is not None:
            for spread_column in spread_columns:
                axis.plot(
                    series["W_IN [W]"],
                    series[spread_column],
                    color=line_color,
                    linewidth=1.05,
                    linestyle="--",
                    alpha=0.85,
                    label="_nolegend_",
                )

    axis.set_title(
        f"{condition_label}\n{report_metric_label(metric_column, group_data)}",
        fontsize=title_size,
        pad=6,
    )
    style_report_axis(
        axis,
        report_display_text("Applied power, W_IN [W]", group_data),
        report_metric_label(metric_column, group_data),
        label_size=label_size,
        tick_size=tick_size,
    )
    axis.set_xlim(left=0)
    axis.legend(
        fontsize=tick_size,
        frameon=False,
        loc="best",
    )


def make_repeated_tests_chart_image(results, part_type, repeated_groups):
    """Create every repeated-test comparison chart on one report page."""
    metric_columns = ["DeltaT_CU [K]"]
    if part_type == "LTS" and "Subcooling [K]" in results.columns:
        metric_columns.append("Subcooling [K]")

    chart_specs = [
        (group_key, group_data, metric_column)
        for group_key, group_data in repeated_groups
        for metric_column in metric_columns
    ]
    rows, columns = chart_grid_shape(len(chart_specs))
    title_size, label_size, tick_size = chart_font_sizes(len(chart_specs))
    figure, axes = plt.subplots(
        rows,
        columns,
        figsize=(11.2, 6.35),
        squeeze=False,
    )
    axes_list = axes.flatten()

    show_fluid = results["Working Fluid"].nunique() > 1
    show_orientation = results["Orientation"].nunique() > 1

    for chart_index, (group_key, group_data, metric_column) in enumerate(
        chart_specs
    ):
        plot_repeated_test_metric(
            axes_list[chart_index],
            group_data,
            metric_column,
            repeated_condition_label(
                group_key,
                show_fluid,
                show_orientation,
            ),
            title_size,
            label_size,
            tick_size,
        )

    for unused_axis in axes_list[len(chart_specs):]:
        unused_axis.axis("off")

    figure.patch.set_facecolor("white")
    figure.tight_layout(pad=1.1, h_pad=1.5, w_pad=1.2)

    image_buffer = BytesIO()
    figure.savefig(
        image_buffer,
        format="png",
        dpi=190,
        facecolor="white",
        bbox_inches="tight",
    )
    plt.close(figure)
    image_buffer.seek(0)
    return image_buffer


def draw_repeated_tests_page(
    pdf,
    results,
    part_type,
    repeated_groups,
    page_number,
):
    """Append the repeated-test comparison page to the report."""
    page_width, page_height = landscape(A4)

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, 16)
    pdf.drawString(32, page_height - 34, "Repeated tests")

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 8.5)
    pdf.drawRightString(
        page_width - 32,
        page_height - 30,
        f"Report date: {date.today().isoformat()}",
    )

    subtitle = (
        "ΔT_CU comparison for CSV files tested at identical conditions. "
        "Dashed bounds show ΔT_CU_MIN and ΔT_CU_MAX."
    )
    if part_type == "LTS":
        subtitle += " Subcooling is shown alongside ΔT_CU."
    pdf.drawString(32, page_height - 51, report_display_text(subtitle, results))

    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, page_height - 63, page_width - 32, page_height - 63)

    chart_image = make_repeated_tests_chart_image(
        results,
        part_type,
        repeated_groups,
    )
    pdf.drawImage(
        ImageReader(chart_image),
        26,
        31,
        width=page_width - 52,
        height=page_height - 106,
        preserveAspectRatio=True,
        anchor="c",
    )
    chart_image.close()

    draw_page_footer(pdf, page_number, page_width)


def natural_text_sort_key(value):
    """Sort sensor names naturally so _2 follows _1 instead of _10."""
    return [
        int(token) if token.isdigit() else token.lower()
        for token in re.split(r"(\d+)", str(value))
    ]


def short_t_cu_label(column):
    """Return a compact T_CU sensor label for plateau maps."""
    label = re.sub(r"\s*\[°C\]\s*$", "", str(column), flags=re.IGNORECASE)
    match = re.search(r"T_CU(?:_|$)(.*)$", label, re.IGNORECASE)
    if match and match.group(1):
        return match.group(1)
    return label


def t_cu_sensor_layout(t_cu_columns):
    """
    Return rectangle positions for recognized T_CU layouts.

    The five-sensor L/MID/R arrangement matches the static geometry used by
    the earlier animation. Any unfamiliar naming scheme falls back to a
    complete labelled grid so no T_CU channel is omitted.
    """
    columns = sorted(t_cu_columns, key=natural_text_sort_key)
    recognized_positions = {}

    position_patterns = [
        (r"T_CU_L_1$", (0.00, 0.00, 1.30, 0.95, 1)),
        (r"T_CU_L_2$", (0.00, 1.05, 1.30, 0.95, 1)),
        (r"T_CU_R_1$", (1.70, 0.00, 1.30, 0.95, 1)),
        (r"T_CU_R_2$", (1.70, 1.05, 1.30, 0.95, 1)),
        (r"T_CU_MID$", (1.08, 0.55, 0.84, 0.90, 3)),
    ]

    for column in columns:
        normalized_name = re.sub(
            r"\s*\[°C\]\s*$",
            "",
            str(column),
            flags=re.IGNORECASE,
        ).upper()
        for pattern, position in position_patterns:
            if re.search(pattern, normalized_name, re.IGNORECASE):
                recognized_positions[column] = position
                break

    if len(recognized_positions) == len(columns) and recognized_positions:
        layout = [
            (column, *recognized_positions[column])
            for column in columns
        ]
        layout.sort(key=lambda item: item[-1])
        return layout, (-0.05, 3.05), (-0.05, 2.05)

    simple_positions = {}
    simple_patterns = [
        (r"T_CU_L$", (0.00, 0.00, 0.95, 1.00, 1)),
        (r"T_CU_MID$", (1.05, 0.00, 0.95, 1.00, 1)),
        (r"T_CU_R$", (2.10, 0.00, 0.95, 1.00, 1)),
    ]
    for column in columns:
        normalized_name = re.sub(
            r"\s*\[°C\]\s*$",
            "",
            str(column),
            flags=re.IGNORECASE,
        ).upper()
        for pattern, position in simple_patterns:
            if re.search(pattern, normalized_name, re.IGNORECASE):
                simple_positions[column] = position
                break

    if len(simple_positions) == len(columns) and simple_positions:
        layout = [
            (column, *simple_positions[column])
            for column in columns
        ]
        return layout, (-0.05, 3.10), (-0.05, 1.05)

    column_count = min(3, max(1, math.ceil(math.sqrt(len(columns)))))
    row_count = math.ceil(len(columns) / column_count)
    layout = []
    for sensor_index, column in enumerate(columns):
        grid_column = sensor_index % column_count
        grid_row = row_count - 1 - sensor_index // column_count
        layout.append((
            column,
            float(grid_column),
            float(grid_row),
            0.92,
            0.82,
            1,
        ))

    return (
        layout,
        (-0.05, float(column_count)),
        (-0.05, float(row_count)),
    )


def plateau_grid_shape(plateau_count):
    """Choose a compact grid for plateau snapshots in the left page panel."""
    if plateau_count <= 1:
        return 1, 1
    if plateau_count <= 6:
        return math.ceil(plateau_count / 2), 2
    return math.ceil(plateau_count / 3), 3


def build_plateau_snapshots(test_detail):
    """Calculate one T_CU-only color snapshot per accepted plateau."""
    t_cu_columns = test_detail["t_cu_columns"]
    snapshots = []

    for step in test_detail["steps"]:
        snapshot = {
            "power": numeric_average(step, get_power_column(step)),
            "temperatures": {
                column: numeric_average(step, column)
                for column in t_cu_columns
            },
        }
        snapshots.append(snapshot)

    snapshots.sort(key=lambda item: item["power"])
    return snapshots


def raw_t_cu_x_values(data):
    """Use a numeric relative-time column when possible, else row number."""
    time_column = find_optional_column(
        data,
        ["RelTime", "REL_TIME", "RelativeTime", "ElapsedTime", "Seconds", "Time_s", "Time [s]"],
    )
    if time_column is not None:
        numeric_time = pd.to_numeric(data[time_column], errors="coerce")
        if numeric_time.notna().all() and len(numeric_time) >= 2:
            return numeric_time.to_numpy(dtype=float), "Relative time [s]"

    return np.arange(len(data), dtype=float), "Row index"


def draw_plateau_temperature_map(
    axis,
    snapshot,
    sensor_layout,
    x_limits,
    y_limits,
    color_normalizer,
    color_map,
    font_size,
    board_layout=False,
):
    """Draw one static colored T_CU sensor layout for a power plateau."""
    for column, x, y, width, height, z_order in sensor_layout:
        temperature = snapshot["temperatures"][column]
        face_color = color_map(color_normalizer(temperature))
        red, green, blue, _ = face_color
        luminance = 0.2126 * red + 0.7152 * green + 0.0722 * blue
        text_color = "black" if luminance > 0.58 else "white"

        rectangle = plt.Rectangle(
            (x, y),
            width,
            height,
            facecolor=face_color,
            edgecolor="#333333",
            linewidth=0.65,
            zorder=z_order,
        )
        axis.add_patch(rectangle)
        label = (f"CPU{re.search(r'\d+', column).group()} | {temperature:.1f}°C"
                 if board_layout else f"{short_t_cu_label(column)}\n{temperature:.1f}°C")
        axis.text(
            x + width / 2,
            y + height / 2,
            label,
            rotation=90 if board_layout else 0,
            ha="center",
            va="center",
            fontsize=font_size,
            color=text_color,
            fontweight="bold",
            zorder=z_order + 1,
        )

    axis.set_xlim(*x_limits)
    axis.set_ylim(*y_limits)
    axis.set_aspect("equal")
    axis.axis("off")
    axis.set_title(
        ("Total heat load: " if board_layout else "") + f"{format_number(snapshot['power'])} W",
        fontsize=7.5,
        fontweight="bold",
        pad=2,
    )


def make_test_raw_data_chart_image(
    test_detail,
    colormap_max_temperature=85.0,
    colormap_min_temperature=25.0,
):
    """Create plateau T_CU maps and the complete raw T_CU trace plot."""
    board_test = is_board_test_detail(test_detail)
    if board_test and test_detail.get("raw_data_only"):
        return make_shift2dc_raw_data_chart_image(test_detail)
    data = test_detail["data"]
    t_cu_columns = test_detail["t_cu_columns"]
    snapshots = build_plateau_snapshots(test_detail)

    if not snapshots:
        raise ValueError(
            f"{test_detail['source_file']}: no valid plateau snapshots were found"
        )

    minimum_temperature, maximum_temperature = validate_colormap_temperature_bounds(
        colormap_min_temperature, colormap_max_temperature,
    )

    color_normalizer = matplotlib.colors.Normalize(
        vmin=minimum_temperature,
        vmax=maximum_temperature,
        clip=True,
    )
    color_map = plt.colormaps.get_cmap("turbo")
    if board_test:
        sensor_layout, x_limits, y_limits = shift2dc_board_sensor_layout(t_cu_columns)
        plateau_columns = 1 if len(snapshots) <= 4 else 2
        plateau_rows = math.ceil(len(snapshots) / plateau_columns)
    else:
        sensor_layout, x_limits, y_limits = t_cu_sensor_layout(t_cu_columns)
        plateau_rows, plateau_columns = plateau_grid_shape(len(snapshots))

    figure = plt.figure(figsize=(11.2, 6.35), facecolor="white")
    outer_grid = figure.add_gridspec(
        1,
        2,
        width_ratios=[1.3, 1.0],
        left=0.035,
        right=0.985,
        top=0.87 if board_test else 0.91,
        bottom=0.10,
        wspace=0.20,
    )
    left_grid = outer_grid[0].subgridspec(
        plateau_rows + 1,
        plateau_columns,
        height_ratios=[1.0] * plateau_rows + [0.08],
        hspace=0.42,
        wspace=0.16,
    )

    figure.text(
        0.285,
        0.965,
        ("Plateau-average CPU maps\n" if board_test
         else f"Plateau-average T_CU maps (last {sample_size} values)\n")
        +
        f"Color scale: {minimum_temperature:g}°C to "
        f"{maximum_temperature:g}°C",
        ha="center",
        va="top",
        fontsize=9.2,
        fontweight="bold",
    )

    map_font_size = 6.2 if len(snapshots) <= 6 else 5.0
    for snapshot_index, snapshot in enumerate(snapshots):
        map_row = snapshot_index // plateau_columns
        map_column = snapshot_index % plateau_columns
        map_axis = figure.add_subplot(left_grid[map_row, map_column])
        draw_plateau_temperature_map(
            map_axis,
            snapshot,
            sensor_layout,
            x_limits,
            y_limits,
            color_normalizer,
            color_map,
            map_font_size,
            board_layout=board_test,
        )

    for unused_index in range(len(snapshots), plateau_rows * plateau_columns):
        unused_row = unused_index // plateau_columns
        unused_column = unused_index % plateau_columns
        figure.add_subplot(left_grid[unused_row, unused_column]).axis("off")

    colorbar_axis = figure.add_subplot(left_grid[-1, :])
    color_mappable = matplotlib.cm.ScalarMappable(
        norm=color_normalizer,
        cmap=color_map,
    )
    color_mappable.set_array([])
    colorbar = figure.colorbar(
        color_mappable,
        cax=colorbar_axis,
        orientation="horizontal",
    )
    interior_ticks = [
        float(tick)
        for tick in colorbar.get_ticks()
        if minimum_temperature < float(tick) < maximum_temperature
    ]
    colorbar.set_ticks([
        minimum_temperature,
        *interior_ticks,
        maximum_temperature,
    ])
    colorbar.set_label("T_CPU [°C]" if board_test else "T_CU [°C]", fontsize=7)
    colorbar.ax.tick_params(labelsize=6)

    raw_axis = figure.add_subplot(outer_grid[1])
    x_values, x_label = raw_t_cu_x_values(data)
    line_colors = plt.get_cmap("tab10").colors

    for column_index, column in enumerate(
        sorted(t_cu_columns, key=natural_text_sort_key)
    ):
        temperatures = pd.to_numeric(data[column], errors="coerce")
        raw_axis.plot(
            x_values,
            temperatures,
            linewidth=0.9,
            color=line_colors[column_index % len(line_colors)],
            label=re.sub(r"\s*\[°C\]\s*$", "", str(column)),
        )

    raw_axis.set_title("Raw T_CPU data" if board_test else "Raw T_CU data", fontsize=10, fontweight="bold", pad=8)
    style_report_axis(
        raw_axis,
        x_label,
        "T_CPU [°C]" if board_test else "T_CU [°C]",
        label_size=8,
        tick_size=7,
    )
    if board_test:
        for row in test_detail["schedule"]:
            raw_axis.axvspan(row["Average from [s]"], row["End [s]"], color="#6DAD82", alpha=.15)
    raw_axis.legend(
        fontsize=6.2,
        frameon=False,
        loc="best",
        ncol=2 if len(t_cu_columns) > 6 else 1,
    )

    image_buffer = BytesIO()
    figure.savefig(
        image_buffer,
        format="png",
        dpi=190,
        facecolor="white",
        bbox_inches="tight",
    )
    plt.close(figure)
    image_buffer.seek(0)
    return image_buffer


def raw_test_page_title(metadata):
    """Build the condition title used by one raw-data test page."""
    if metadata["condition_type"] == "FR":
        condition_label = f"FR {format_number(metadata['condition_value'])}%"
    else:
        condition_label = f"Charge {format_number(metadata['condition_value'])}"

    title_parts = [
        condition_label,
        flow_rate_report_label(
            metadata["flow_type"],
            metadata["nominal_flow_rate"],
        ),
        inlet_temperature_report_label(
            metadata["medium"],
            metadata["nominal_temperature"],
            compact=True,
        ),
    ]
    if str(metadata["orientation"]).upper() != "N/A":
        title_parts.append(orientation_report_name(metadata["orientation"]))

    return "Test T_CU detail - " + " | ".join(title_parts)


def draw_test_raw_data_page(
    pdf,
    test_detail,
    page_number,
    colormap_max_temperature=85.0,
    colormap_min_temperature=25.0,
):
    """Draw one final report page for one source CSV test."""
    page_width, page_height = landscape(A4)
    metadata = test_detail["metadata"]
    page_title = raw_test_page_title(metadata)
    if is_board_test_detail(test_detail):
        page_title = page_title.replace("T_CU", "T_CPU")
    title_size = fit_report_cell_font(
        pdf,
        [page_title],
        report_bold_font,
        15,
        page_width - 255,
    )

    pdf.setFillColor(report_dark)
    pdf.setFont(report_bold_font, title_size)
    pdf.drawString(32, page_height - 32, page_title)

    pdf.setFillColor(report_grey)
    pdf.setFont(report_regular_font, 8.5)
    pdf.drawRightString(
        page_width - 32,
        page_height - 29,
        f"Report date: {date.today().isoformat()}",
    )

    subtitle = f"CSV: {test_detail['source_file']}"
    comment = str(metadata.get("comment", "")).strip()
    if comment:
        subtitle += f"  |  Comment: {comment}"

    pdf.setFillColor(report_grey)
    draw_wrapped_text(
        pdf,
        subtitle,
        32,
        page_height - 50,
        page_width - 64,
        font_size=7.7,
        leading=9,
        maximum_lines=2,
    )

    pdf.setStrokeColor(report_light_grey)
    pdf.setLineWidth(0.7)
    pdf.line(32, page_height - 70, page_width - 32, page_height - 70)

    chart_image = make_test_raw_data_chart_image(
        test_detail,
        colormap_max_temperature=colormap_max_temperature,
        colormap_min_temperature=colormap_min_temperature,
    )
    pdf.drawImage(
        ImageReader(chart_image),
        26,
        31,
        width=page_width - 52,
        height=page_height - 112,
        preserveAspectRatio=True,
        anchor="c",
    )
    chart_image.close()

    draw_page_footer(pdf, page_number, page_width)


def create_pdf_report(
    all_results,
    part_type,
    report_output_file,
    test_details=None,
    report_configuration=None,
    fluid_properties=None,
):
    """Create the complete report and overwrite an existing report."""
    if part_type == 'LTS':
        all_results = add_lts_psat_results(all_results, test_details or [], report_configuration or {})
    test_details = test_details or []
    campaign = report_campaign_frame(all_results, test_details)
    results = steady_campaign_rows(campaign)
    analysis_results = primary_report_results(results)

    filling_ratios = analysis_results.loc[
        analysis_results["Condition Type"] == "FR",
        "Condition Value",
    ].dropna().unique()

    has_filling_ratio_analysis = len(filling_ratios) > 1
    _, repeated_groups = find_repeated_test_groups(results)
    performance_groups = find_performance_characterization_groups(
        analysis_results
    )
    performance_ratios = [ratio for ratio, _ in performance_groups]
    if report_configuration is None:
        report_configuration = default_report_configuration(
            has_filling_ratio_analysis,
            len(repeated_groups),
            test_details,
            performance_ratios,
        )

    include_filling_ratio_analysis = (
        has_filling_ratio_analysis
        and report_configuration.get(
            "include_filling_ratio_analysis",
            True,
        )
    )
    include_repeated_tests = (
        bool(repeated_groups)
        and report_configuration.get("include_repeated_tests", True)
    )
    include_performance_characterization = (
        bool(performance_groups)
        and report_configuration.get(
            "include_performance_characterization",
            True,
        )
    )
    include_summary = bool(len(analysis_results)) and report_configuration.get("include_test_summary", is_board_report(results) or bool(include_filling_ratio_analysis))
    include_colormaps = report_configuration.get("include_colormaps", True)
    selected_colormap_files = set(
        report_configuration.get(
            "selected_colormap_files",
            {detail["source_file"] for detail in test_details},
        )
    )
    colormap_min_temperature = float(
        report_configuration.get("colormap_min_temperature", 25.0)
    )
    colormap_max_temperature = float(
        report_configuration.get("colormap_max_temperature", 85.0)
    )

    if report_configuration.get('include_psat', False):
        for detail in test_details:
            for column in shift2dc_extra_columns(detail, 'psat'):
                pressure_column_settings(column, report_configuration)

    if part_type == 'LTS' and report_configuration.get('include_ph_diagram', False):
        for detail in test_details:
            for column in lts_ph_pressure_columns(detail):
                pressure_column_settings(column, report_configuration)

    pdf = canvas.Canvas(
        str(report_output_file),
        pagesize=landscape(A4),
    )

    draw_cover_page(
        pdf,
        campaign,
        part_type,
        include_filling_ratio_analysis,
        has_performance_characterization=(
            include_performance_characterization
        ),
        fluid_properties=fluid_properties,
    )

    page_number = 1
    transient_details = [d for d in test_details if is_transient_detail(d)]
    test_details = [d for d in test_details if not is_transient_detail(d)]

    if include_filling_ratio_analysis:
        pdf.showPage()
        page_number += 1
        draw_analysis_page(
            pdf,
            analysis_results,
            "T_CU_AVG [°C]",
            "Filling Ratio Analysis - Average Copper Temperature",
            page_number,
            spread_columns=(
                "T_CU_MIN [°C]",
                "T_CU_MAX [°C]",
            ),
        )

        detailed_group_columns = [
            "Working Fluid",
            "Coolant",
            "Orientation",
            "Flow Type",
            "Nominal Flow Rate",
            "Nominal T_IN [°C]",
        ]
        filling_ratio_results = analysis_results[
            analysis_results["Condition Type"] == "FR"
        ].copy()
        detailed_groups = list(
            filling_ratio_results.groupby(
                detailed_group_columns,
                dropna=False,
                sort=True,
            )
        )

        for group_key, group_data in detailed_groups:
            pdf.showPage()
            page_number += 1
            draw_detailed_filling_ratio_page(
                pdf,
                group_key,
                group_data,
                page_number,
            )

        if include_summary:
            pdf.showPage()
            page_number += 1
            page_number = draw_condition_summary_pages(
                pdf, analysis_results, part_type, page_number, report_configuration,
            )

    if include_performance_characterization:
        for filling_ratio, ratio_data in performance_groups:
            pdf.showPage()
            page_number += 1
            draw_performance_characterization_page(
                pdf,
                filling_ratio,
                ratio_data,
                page_number,
            )

    if include_repeated_tests:
        pdf.showPage()
        page_number += 1
        draw_repeated_tests_page(
            pdf,
            results,
            part_type,
            repeated_groups,
            page_number,
        )

    if include_summary and not include_filling_ratio_analysis:
        if is_board_report(results) and not performance_groups:
            for ratio, ratio_data in analysis_results.groupby("Condition Value", sort=True):
                pdf.showPage()
                page_number += 1
                draw_performance_characterization_page(
                    pdf, ratio, ratio_data, page_number,
                    page_title=f"Test Results - FR {format_number(ratio)}%",
                )
        pdf.showPage()
        page_number += 1
        page_number = draw_condition_summary_pages(pdf, analysis_results, part_type, page_number, report_configuration)

    if part_type == 'LTS' and report_configuration.get('include_superheating_subcooling', True):
        for thermal_group, metric in superheating_subcooling_groups(analysis_results):
            pdf.showPage()
            page_number += 1
            draw_superheating_subcooling_page(pdf, thermal_group, metric, page_number, report_configuration)

    if part_type == 'LTS' and report_configuration.get('include_psat_comparison', False):
        for pressure_group, metric in psat_comparison_groups(analysis_results):
            pdf.showPage()
            page_number += 1
            draw_psat_comparison_page(pdf, pressure_group, metric, page_number)

    selected_raw_files = set(report_configuration.get(
        "selected_raw_files", {detail["source_file"] for detail in test_details},
    ))
    selected_detail_files = set(report_configuration.get('selected_detail_files',
        {d['source_file'] for d in test_details}))
    for detail in test_details:
        if detail['source_file'] in selected_detail_files:
            if report_configuration.get('include_psat', False):
                for column in shift2dc_extra_columns(detail, 'psat'):
                    pdf.showPage()
                    page_number += 1
                    draw_shift2dc_extra_page(pdf, detail, page_number, 'psat', report_configuration, column)
            if report_configuration.get('include_psu_temperatures', False) and shift2dc_extra_columns(detail, 'psu'):
                pdf.showPage()
                page_number += 1
                draw_shift2dc_extra_page(pdf, detail, page_number, 'psu', report_configuration)

            if part_type == 'LTS' and report_configuration.get('include_ph_diagram', False):
                for pressure_column in lts_ph_pressure_columns(detail) or [None]:
                    try:
                        records, dome, backend, pressure_note = build_lts_ph_data(detail, report_configuration, pressure_column)
                        for start in range(0, len(records), 4):
                            pdf.showPage()
                            page_number += 1
                            draw_lts_ph_page(pdf, detail, page_number, records[start:start+4], dome,
                                             backend, pressure_note, report_configuration)
                    except (ValueError, RuntimeError) as error:
                        pdf.showPage()
                        page_number += 1
                        draw_lts_ph_page(pdf, detail, page_number, [], None, '', '', report_configuration, error=error)

        wants_map = include_colormaps and detail['source_file'] in selected_colormap_files
        wants_raw = (not include_colormaps and detail.get('raw_data_only') and
                     report_configuration.get('include_raw_data', True) and detail['source_file'] in selected_raw_files)
        if wants_map or wants_raw:
            pdf.showPage()
            page_number += 1
            draw_test_raw_data_page(pdf,
                dict(detail, raw_data_only=not wants_map) if is_board_test_detail(detail) else detail,
                page_number, colormap_max_temperature=colormap_max_temperature,
                colormap_min_temperature=colormap_min_temperature)

    if report_configuration.get('include_transient_tests', bool(transient_details)):
        for detail in transient_details:
            page_number = draw_transient_pages(pdf, detail, page_number, report_configuration)

    pdf.save()
    return page_number


def main(
    selected_input_folder=None,
    report_configuration=None,
    open_pdf_when_done=True,
    shift2dc_configurations=None,
):
    if selected_input_folder is None:
        input_folder = select_raw_data_folder()
    else:
        input_folder = Path(selected_input_folder)

    if input_folder is None:
        print("No folder was selected. Processing cancelled.")
        return

    if input_folder.name.lower() != "00_rawdata":
        raise ValueError(
            "The selected directory must be named 00_RawData.\n"
            f"Selected directory: {input_folder}"
        )

    output_folder = input_folder.parent / "01_PostProcessedData"

    if sample_size <= 0:
        raise ValueError("sample_size must be greater than zero")
    if minimum_step_size <= 0:
        raise ValueError("minimum_step_size must be greater than zero")

    output_folder.mkdir(parents=True, exist_ok=True)
    raw_files = sorted(input_folder.glob("*.csv"))
    if not raw_files:
        raise FileNotFoundError(f"No CSV files found in:\n{input_folder}")

    shift_files = [f for f in raw_files if is_shift2dc_file(f)]
    if shift_files:
        if len(shift_files) != len(raw_files):
            raise ValueError("Keep Shift2DC server tests in a separate 00_RawData folder from ordinary PSU tests.")
        return process_shift2dc_files(shift_files, output_folder, shift2dc_configurations, open_pdf_when_done, report_configuration)

    metadata_by_file = {}
    cleaned_by_file = {}
    part_types, fluids, orientations = set(), set(), set()
    filling_ratios, charges = set(), set()
    all_t_cu_columns = []

    for raw_file in raw_files:
        metadata = parse_file_name(raw_file)
        metadata["source_file"] = raw_file.name
        part_types.add(metadata["part_type"])
        fluids.add(metadata["fluid"])
        orientations.add(metadata["orientation"])
        target_set = filling_ratios if metadata["condition_type"] == "FR" else charges
        target_set.add(metadata["condition_value"])

        if metadata.get('test_mode') == 'TR':
            cleaned = prepare_transient_data(read_csv_automatically(raw_file))
            cleaned.to_csv(output_folder / raw_file.name, index=False, encoding='utf-8-sig')
        else:
            cleaned = create_cleaned_csv(raw_file, output_folder / raw_file.name)
        complete_nominal_conditions(metadata, cleaned)
        metadata_by_file[raw_file] = metadata
        cleaned_by_file[raw_file] = cleaned
        for column in find_t_cu_columns(cleaned):
            output_column = t_cu_output_name(column)
            if output_column not in all_t_cu_columns:
                all_t_cu_columns.append(output_column)

    if len(part_types) != 1:
        raise ValueError(
            "A selected 00_RawData folder must contain only one part type. "
            f"Found: {sorted(part_types)}"
        )

    part_type = next(iter(part_types))
    excel_output_file = output_folder / f"{part_type}_Test_Averages.xlsx"
    report_output_file = output_folder / f"{part_type}_Test_Report.pdf"

    if not all_t_cu_columns:
        raise ValueError("No T_CU columns were found")

    print(f"Confirmed part type: {part_type}")
    print(f"T_CU columns: {', '.join(all_t_cu_columns)}")
    print(f"Number of CSV tests: {len(raw_files)}")
    print(
        "Minimum valid power-step size: "
        f"{max(minimum_step_size, sample_size)} measurement points"
    )
    print(f"Fluids found: {', '.join(sorted(fluids))}")
    print(f"Orientations found: {', '.join(sorted(orientations))}")
    if filling_ratios:
        print("Filling ratios covered: " + ", ".join(f"{x:g}%" for x in sorted(filling_ratios, reverse=True)))
    if charges:
        print("Charges covered: " + ", ".join(f"{x:g}" for x in sorted(charges, reverse=True)))

    all_results = []
    test_details = []
    for raw_file in raw_files:
        data = cleaned_by_file[raw_file]
        t_cu_columns = find_t_cu_columns(data)
        steps = ([] if metadata_by_file[raw_file].get('test_mode') == 'TR' else
                 split_into_power_steps(data, get_power_column(data)))
        test_details.append({
            "source_file": raw_file.name,
            "metadata": metadata_by_file[raw_file],
            "data": data,
            "t_cu_columns": t_cu_columns,
            "steps": steps,
        })
        for step in steps:
            result = calculate_step_result(step, metadata_by_file[raw_file], t_cu_columns)
            for column in all_t_cu_columns:
                result.setdefault(column, np.nan)
            all_results.append(result)
        print(f"Processed {raw_file.name}: {len(steps)} power step(s)")

    if not all_results and not any(is_transient_detail(d) for d in test_details):
        raise RuntimeError("No non-zero power-step averages were produced")
    report_configuration, fluid_properties = configure_report(
        all_results, part_type, test_details, report_configuration,
    )
    if report_configuration is None:
        print("Report generation cancelled by the user.")
        return

    if part_type == 'LTS':
        all_results = add_lts_psat_results(all_results, test_details, report_configuration)
    create_master_excel(
        all_results,
        all_t_cu_columns,
        part_type,
        excel_output_file,
    )

    report_page_count = create_pdf_report(
        all_results,
        part_type,
        report_output_file,
        test_details=test_details,
        report_configuration=report_configuration,
        fluid_properties=fluid_properties,
    )

    print(f"\nMaster Excel file created:\n{excel_output_file}")
    print("An existing file with the same name was overwritten.")
    print(
        f"\nPDF report created ({report_page_count} page(s)):\n"
        f"{report_output_file}"
    )
    print("An existing report with the same name was overwritten.")

    if open_pdf_when_done:
        open_pdf_automatically(report_output_file)


"""Condition-based simulation imports. Also embedded in process_php_csvs.py."""
import csv
import json
import math
import re
import os
import shutil
import tempfile
from pathlib import Path
from datetime import datetime
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def correlation_token(value):
    return re.sub(r'[^a-z0-9]+', '', str(value).lower().replace('Δ', 'delta').replace('δ', 'delta').replace('²', '2').replace('³', '3'))


CORRELATION_ALIASES = {
    'power': ['powerw', 'power', 'wpsu', 'wpsuw', 'win', 'winw', 'inputq', 'heatloadw', 'heatload'],
    'fluid': ['fluid', 'workingfluid', 'inputfluid', 'refrigerant'],
    'orientation': ['orientation', 'inputorientation', 'orientationcode'],
    'water_flow': ['mfrlmin', 'vfrlmin', 'vfr', 'mfr', 'inputvfr', 'waterflowlmin'],
    'htc': ['htc', 'inputhtc', 'htcwm2k', 'htcwm²k'],
    'air_flow': ['cfm', 'aircfm', 'inputcfm'],
    'inlet_temperature': ['twater', 'twaterc', 'twaterin', 'twaterinc', 'tair', 'tairc', 'tairin', 'tw', 'tin', 'inputtcool', 'inputtcoolant', 'inputtwater', 'inputtair', 'inputtc', 'coolanttemperaturec', 'coolanttemperature'],
    'filling_ratio': ['fr', 'frpercent', 'fillingratio', 'inputfr', 'fillingratiopercent'],
    'charge': ['charge', 'ch', 'chargeg'],
    'coolant': ['coolant', 'coolingfluid'],
    'component': ['php', 'component', 'part', 'project', 'phpproject', 'phptypeproject', 'testarticle', 'inputproject'],
    'evaporator': ['evap', 'evaporator'],
    'condenser': ['cond', 'condenser'],
    'comment': ['comment', 'testcomment'],
}


def correlation_kind(header):
    token = correlation_token(header)
    return next((key for key, aliases in CORRELATION_ALIASES.items() if token in aliases), None)


def correlation_role(header):
    if correlation_kind(header):
        return 'Condition'
    token = correlation_token(header)
    if 'sim' in token:
        return 'Simulation'
    if 'exp' in token or 'testdata' in token:
        return 'Experimental'
    if any(word in token for word in ('tcu', 'tevap', 'deltat', 'rth', 'subcool', 'tadia')):
        return 'Experimental'
    return 'Condition'  # Unfamiliar conditions must be explicitly reviewed.


def correlation_value(value, kind=None):
    if value is None or str(value).strip() == '':
        return None
    text = str(value).strip()
    if kind == 'filling_ratio' and text.endswith('%'):
        text = text[:-1].strip()
    try:
        number = float(text.replace(',', '.') if ',' in text and '.' not in text else text)
        return number if math.isfinite(number) else None
    except ValueError:
        pass
    if kind in ('fluid', 'component'):
        return correlation_token(text)
    if kind == 'orientation':
        return re.sub(r'\s+', ' ', text).casefold()
    return text.casefold()

def correlation_equal(left, right):
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(left, right, rel_tol=1e-9, abs_tol=1e-6)
    return left == right


def correlation_headers(sheet, header_row=1):
    headers = {str(c.value).strip(): c.column for c in sheet[header_row] if c.value is not None and str(c.value).strip()}
    names = [str(c.value).strip() for c in sheet[header_row] if c.value is not None and str(c.value).strip()]
    if len(names) != len(headers):
        raise ValueError('Duplicate column titles. Give every column a unique title first.')
    if not headers:
        raise ValueError('No column titles found on the selected header row.')
    return headers



def correlation_reference_records(sheet, header_row=1):
    """Read test rows, resolving merged or sparse PHP labels within each block."""
    headers = correlation_headers(sheet, header_row)
    power = next((c for h,c in headers.items() if correlation_kind(h)=='power'), None)
    fluid = next((c for h,c in headers.items() if correlation_kind(h)=='fluid'), None)
    if not power or not fluid:
        raise ValueError('The reference must contain Power and Fluid headers.')
    component = next((h for h in headers if correlation_kind(h)=='component'), None)
    previous = None
    rows = []
    for r in range(header_row+1, sheet.max_row+1):
        if sheet.cell(r,power).value is None and sheet.cell(r,fluid).value is None:
            previous = None
            continue
        record = {h: sheet.cell(r,c).value for h,c in headers.items()}
        if component:
            if record[component] is not None and str(record[component]).strip():
                previous = record[component]
            else:
                record[component] = previous
        for h,c in headers.items():
            cell=sheet.cell(r,c)
            if correlation_kind(h)=='filling_ratio' and '%' in cell.number_format and isinstance(record[h],(int,float)):
                record[h] *= 100
        rows.append((r,record))
    return rows


def correlation_prepare_source(titles, records):
    """Unpack exported inputs_json and expose angle/roll pairs without guessing orientation."""
    titles=list(titles)
    for _,record in records:
        raw=record.get('inputs_json')
        if raw:
            try:
                doc=json.loads(raw)
                cfg=doc.get('config',doc.get('inputs',doc))
                if isinstance(cfg,dict):
                    for k,v in cfg.items():
                        key='input_'+k
                        if not isinstance(v,(dict,list)) and record.get(key) in (None,''):
                            record[key]=v
                            if key not in titles:titles.append(key)
            except (ValueError,TypeError):
                raise ValueError('Invalid inputs_json in the simulation export.')
        for canonical,key in [('wall_conduction','wallConduction'),('film_model','filmModel'),('conduction_mode','conductionMode'),('physics_preset','physicsPreset'),('effective_physics','effectivePhysics'),('transport_source','transportSource')]:
            if record.get(canonical) in (None,'') and record.get('input_'+key) not in (None,''):
                record[canonical]=record['input_'+key]
                if canonical not in titles:titles.append(canonical)
        if record.get('input_angle') not in (None,'') and record.get('input_roll') not in (None,''):
            key='Orientation angles (angle, roll)'
            record[key]=f"{float(record['input_angle']):g}, {float(record['input_roll']):g}"
            if key not in titles:titles.append(key)
    return titles,records


CORRELATION_PHYSICS_FIELDS=('app_version','physics_version','wall_conduction','film_model',
                          'conduction_mode','physics_preset','effective_physics','transport_source')


def correlation_physics_signature(record):
    return tuple((key,str(record.get(key,'')).strip()) for key in CORRELATION_PHYSICS_FIELDS)


def correlation_read_source(path, sheet_name=None, header_row=1):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            lines = stream.readlines()
        while lines and (not lines[0].strip() or lines[0].lstrip().startswith('#')):
            lines.pop(0)
        delimiter = None
        if lines and lines[0].strip().lower().startswith('sep='):
            delimiter = lines.pop(0).strip()[4:]
            if delimiter not in (',', ';', '\t'):
                raise ValueError('Unsupported CSV delimiter in sep= metadata.')
        if not lines:
            raise ValueError('Simulation CSV is empty.')
        try:
            dialect = csv.Sniffer().sniff(''.join(lines)[:16384], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(lines, delimiter=delimiter)) if delimiter else list(csv.reader(lines, dialect))
        if not 1 <= header_row <= len(rows):
            raise ValueError('Simulation header row is outside the file.')
        titles = [str(v).strip() for v in rows[header_row-1]]
        if any(not v for v in titles) or len(set(titles)) != len(titles):
            raise ValueError('Simulation columns need unique, nonempty titles.')
        records = []
        for index, row in enumerate(rows[header_row:], header_row+1):
            if not any(str(v).strip() for v in row):
                continue
            if len(row) != len(titles):
                raise ValueError(f'Simulation row {index} has {len(row)} fields; expected {len(titles)}.')
            records.append((index, dict(zip(titles, row))))
        return correlation_prepare_source(titles, records)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Choose an .xlsx or .csv simulation file.')
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    headers = correlation_headers(sheet, header_row)
    records = [(r, {name: sheet.cell(r, c).value for name, c in headers.items()})
               for r in range(header_row+1, sheet.max_row+1)
               if any(sheet.cell(r, c).value is not None for c in headers.values())]
    for r,record in records:
        for h,c in headers.items():
            if correlation_kind(h)=='filling_ratio' and '%' in sheet.cell(r,c).number_format and isinstance(record[h],(int,float)):
                record[h] *= 100
    workbook.close()
    return correlation_prepare_source(list(headers), records)


def correlation_auto_mapping(sheet, header_row, source_titles):
    """Header-based matching; constants stay explicit/editable in the GUI."""
    headers=correlation_headers(sheet,header_row)
    rows=correlation_reference_records(sheet,header_row)
    htc_available=any(correlation_kind(h)=='htc' for h in source_titles)
    defaults={}
    for title in headers:
        kind=correlation_kind(title);role=correlation_role(title)
        candidates=[h for h in source_titles if correlation_token(h)==correlation_token(title)
                    or kind is not None and correlation_kind(h)==kind]
        preset={'role':role,'column':'','constant':'','note':''}
        if role=='Condition':
            if kind=='water_flow' and htc_available:
                preset.update(role='Ignore',note='HTC is available for matching; VFR is optional.')
            elif len(candidates)==1:
                preset['column']=candidates[0]
            elif not candidates and kind=='orientation' and 'Orientation angles (angle, roll)' in source_titles:
                preset['column']='Orientation angles (angle, roll)'
                preset['note']='Map each angle/roll pair explicitly to the reference orientation code.'
            elif not candidates and rows:
                values=[correlation_value(record[title],kind) for _,record in rows]
                if all(v is not None and correlation_equal(v,values[0]) for v in values):
                    preset.update(column='<constant>',constant=str(rows[0][1][title]),note='Check this constant against the simulation inputs.')
        defaults[title]=preset
    return defaults

def correlation_plan(reference, sheet_name, header_row, records, mappings, metrics, version):
    """Read-only plan. mappings: {reference title: {column|constant, values?}}.

    values is an explicit source-value -> reference-value dictionary (e.g. HTC
    to flow). No positional matching, interpolation, or implicit broadcasting.
    """
    if not version or not str(version).strip():
        raise ValueError('Enter a physics / solver configuration label, or choose an existing result column.')
    if not mappings or not metrics:
        raise ValueError('Select test conditions and at least one simulation result.')
    wb = openpyxl.load_workbook(reference, data_only=True)
    ws = wb[sheet_name]
    headers = correlation_headers(ws, header_row)
    missing = set(mappings) - set(headers)
    if missing:
        raise ValueError(f'Unknown reference conditions: {sorted(missing)}')
    kinds = {correlation_kind(name) for name in mappings}
    if not {'power', 'fluid'} <= kinds:
        raise ValueError('Heat load and working fluid must both be matching conditions. Use recognized titles such as Power [W] and Fluid.')
    required={'component','orientation','filling_ratio','inlet_temperature'}
    absent=[h for h in headers if correlation_kind(h) in required and h not in mappings]
    if absent:
        raise ValueError('Required matching conditions: '+', '.join(absent)+'. Map a source column or enter an explicit constant.')
    if any(correlation_kind(h) in ('htc','water_flow','air_flow') for h in headers) and not kinds.intersection({'htc','water_flow','air_flow'}):
        raise ValueError('Match the cooling condition using HTC or flow rate.')
    signatures={correlation_physics_signature(record) for _,record in records}
    if len(signatures)>1:
        raise ValueError('The selected rows contain different physics / solver configurations. Select one configuration before importing.')
    refs = []
    for r,record in correlation_reference_records(ws,header_row):
        vals=[correlation_value(record[h],correlation_kind(h)) for h in mappings]
        if any(v is None for v in vals):
            raise ValueError(f'Reference row {r} has a missing matching condition. Complete it before importing.')
        refs.append((r,vals))
    wb.close()
    if not refs:
        raise ValueError('No complete reference data rows found.')
    audit, assignments = [], defaultdict(list)
    for source_row, record in records:
        entry = {'source_row': source_row, 'reference_row': None, 'status': '', 'detail': '', 'record': record}
        audit.append(entry)
        error = next((v for k, v in record.items() if correlation_token(k) == 'error' and str(v or '').strip()), None)
        run_status=next((str(v).strip().lower() for k,v in record.items() if correlation_token(k)=='status'), '')
        if run_status and run_status not in ('completed','complete','finished','success','succeeded','done','recorded'):
            error=error or ('Simulation status: '+run_status)
        if not error and any(correlation_token(k)=='times' for k in record) and not (record.get('start_s') not in (None,'') and record.get('end_s') not in (None,'')):
            error='A time-series row is not an averaged correlation result. Export averages or batch summaries.'
        if error:
            entry.update(status='Failed simulation', detail=str(error))
        vals = []
        for name, mapping in mappings.items():
            raw = mapping.get('constant') if 'constant' in mapping else record.get(mapping.get('column'))
            if mapping.get('values') is not None:
                normalized = correlation_value(raw)
                candidates = [v for k, v in mapping['values'].items() if correlation_equal(correlation_value(k), normalized)]
                raw = candidates[0] if len(candidates) == 1 else None
            vals.append(correlation_value(raw, correlation_kind(name)))
        if any(v is None for v in vals):
            if not error:
                entry.update(status='Missing condition', detail='A condition or value mapping is missing.')
            continue
        candidates = [r for r, refvals in refs if all(correlation_equal(a, b) for a, b in zip(vals, refvals))]
        if len(candidates) != 1:
            if not error:
                entry.update(status='Unmatched' if not candidates else 'Ambiguous reference', detail=f'Reference rows: {candidates}')
            continue
        entry['reference_row'] = candidates[0]
        if error:
            continue
        valid = {}
        for metric in metrics:
            v = correlation_value(record.get(metric))
            if isinstance(v, (int, float)):
                valid[metric] = v
        entry['results'] = valid
        if not valid:
            entry.update(status='Missing result', detail='No finite numeric result selected.')
            continue
        assignments[candidates[0]].append(entry)
    for entries in assignments.values():
        for entry in entries:
            if len(entries) > 1:
                entry.update(status='Duplicate simulation', detail='Multiple source rows target this reference row; none imported.')
            else:
                entry.update(status='Matched', detail='' if len(entry['results']) == len(metrics) else 'Some results are missing; existing cells are preserved.')
    return {'reference': str(Path(reference).resolve()), 'sheet': sheet_name, 'header_row': header_row,
            'mappings': mappings, 'metrics': list(metrics), 'version': str(version).strip(),
            'audit': audit, 'counts': dict(Counter(e['status'] for e in audit)),
            'reference_rows': len(refs), 'reference_signature': (Path(reference).stat().st_mtime_ns, Path(reference).stat().st_size)}


def correlation_literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # Imported strings must never execute as formulas.


def correlation_shift_formula(formula, context_sheet, target_sheet, start, count):
    """Update A1 references for an insertion, including absolute/chart references."""
    from openpyxl.formula import Tokenizer
    from openpyxl.utils import column_index_from_string
    def shift_range(value):
        prefix, address = value.rsplit('!', 1) if '!' in value else ('', value)
        sheet = prefix.strip("'").replace("''", "'") if prefix else context_sheet
        if sheet != target_sheet:
            return value
        parts = address.split(':')
        shifted = []
        for part in parts:
            match = re.fullmatch(r'(\$?)([A-Za-z]{1,3})(\$?\d+)?', part)
            if not match:
                return value  # Named ranges and structured references stay intact.
            col = column_index_from_string(match[2])
            shifted.append(match[1]+get_column_letter(col+count if col >= start else col)+(match[3] or ''))
        return (prefix+'!' if prefix else '')+':'.join(shifted)
    has_equals = formula.startswith('=')
    tokens = Tokenizer(formula if has_equals else '='+formula)
    for token in tokens.items:
        if token.type == 'OPERAND' and token.subtype == 'RANGE':
            token.value = shift_range(token.value)
    result = ''.join(token.value for token in tokens.items)
    return ('=' if has_equals else '')+result


def correlation_insert_columns(wb, ws, start, count):
    """openpyxl insert_cols alone does not maintain dependent formulas/charts."""
    from copy import copy
    from openpyxl.descriptors.serialisable import Serialisable
    def formula(value, context):
        return correlation_shift_formula(value, context, ws.title, start, count)
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f' and isinstance(cell.value, str):
                    cell.value = formula(cell.value, sheet.title)
        def update_chart(obj):
            if isinstance(obj, Serialisable):
                for field in obj.__elements__:
                    value = getattr(obj, field, None)
                    if field == 'f' and isinstance(value, str):
                        setattr(obj, field, formula(value, sheet.title))
                    elif isinstance(value, (tuple, list)):
                        for item in value:
                            update_chart(item)
                    else:
                        update_chart(value)
        for chart in sheet._charts:
            update_chart(chart)
        for name in sheet.defined_names.values():
            if name.attr_text:
                name.attr_text = formula(name.attr_text, sheet.title)
    for name in wb.defined_names.values():
        if name.attr_text:
            name.attr_text = formula(name.attr_text, '')
    dimensions = [(key, copy(dim)) for key, dim in ws.column_dimensions.items()]
    merges = [str(value) for value in ws.merged_cells.ranges]
    for value in merges:
        ws.unmerge_cells(value)
    ws.insert_cols(start, count)
    for key in list(ws.column_dimensions):
        del ws.column_dimensions[key]
    from openpyxl.utils import column_index_from_string
    for key, dim in dimensions:
        col = column_index_from_string(key)
        newcol = col+count if col >= start else col
        dim.index = get_column_letter(newcol)
        if dim.min and dim.min >= start:
            dim.min += count
        if dim.max and dim.max >= start:
            dim.max += count
        ws.column_dimensions[dim.index] = dim
    for value in merges:
        ws.merge_cells(formula(value, ws.title))
    for chart in ws._charts:
        anchor = chart.anchor
        if isinstance(anchor, str):
            chart.anchor = formula(anchor, ws.title)
        else:
            for field in ('_from', 'to'):
                marker = getattr(anchor, field, None)
                if marker is not None and marker.col >= start-1:
                    marker.col += count
    if ws.freeze_panes:
        ws.freeze_panes = formula(ws.freeze_panes, ws.title)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = formula(ws.auto_filter.ref, ws.title)
    for table in ws.tables.values():
        # A table spanning the insertion requires extra table-column metadata.
        # Block before saving rather than emit an invalid Excel table.
        from openpyxl.utils.cell import range_boundaries
        lo, _, hi, _ = range_boundaries(table.ref)
        if lo < start <= hi:
            from openpyxl.worksheet.table import TableColumn
            for offset in range(count):
                table.tableColumns.insert(start-lo+offset,TableColumn(id=0,name=f'New result {start+offset}'))
            for index,item in enumerate(table.tableColumns,1):item.id=index
        table.ref = formula(table.ref, ws.title)
        if table.autoFilter and table.autoFilter.ref:
            table.autoFilter.ref = formula(table.autoFilter.ref, ws.title)
    for validation in ws.data_validations.dataValidation:
        validation.sqref = ' '.join(formula(str(r), ws.title) for r in validation.sqref.ranges)
        for field in ('formula1', 'formula2'):
            value = getattr(validation, field)
            if isinstance(value, str):
                setattr(validation, field, formula(value, ws.title))
    # Re-key conditional-format ranges as they are dictionary keys.
    rules = list(ws.conditional_formatting._cf_rules.items())
    ws.conditional_formatting._cf_rules.clear()
    for region, entries in rules:
        region.sqref = ' '.join(formula(str(r), ws.title) for r in region.sqref.ranges)
        for rule in entries:
            if rule.formula:
                rule.formula = [formula(v, ws.title) for v in rule.formula]
        ws.conditional_formatting._cf_rules[region] = entries


def correlation_is_temperature(metric):
    token = correlation_token(metric)
    return not ('delta' in token or 'min' in token or 'max' in token) and (
        'tevap' in token or 'tevaporator' in token or 'tcuavg' in token or 'tcumean' in token)


def correlation_version(value):
    match = re.search(r'(?i)(?:\bv)?(\d+(?:\.\d+)+)(?!\d)', str(value))
    return match.group(1) if match else re.split(r'(?i)\bsim\s+', str(value))[-1].strip().casefold()


def correlation_find_result(headers, metric, version, delta=False):
    # Match the full physics/solver label. A shared version number is insufficient.
    label=re.sub(r'\s+', ' ', str(version).strip()).casefold()
    matches=[]
    for title,col in headers.items():
        token=correlation_token(title)
        if correlation_role(title)!='Simulation' or ('delta' in token)!=delta:continue
        suffix=re.split(r'(?i)\bsim\s+',str(title),maxsplit=1)
        suffix=re.sub(r'\s+', ' ', suffix[-1].strip()).casefold()
        if suffix!=label:continue
        compatible=('tevap' in token or 'tcuavg' in token or 'tcumean' in token) if correlation_is_temperature(metric) else correlation_token(metric) in token
        if compatible:matches.append((title,col))
    if len(matches)>1:raise ValueError(f'Multiple result columns match {metric}, configuration {version}.')
    return matches[0] if matches else None

def correlation_format_and_chart(wb, ws, header_row):
    """Refresh the correlation view for all PHPs and current solver labels."""
    from openpyxl.styles import Border, Side
    from openpyxl.chart import ScatterChart, Series, Reference
    from openpyxl.chart.series import SeriesLabel, XYSeries
    from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumData, NumVal
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils.cell import range_boundaries

    headers=correlation_headers(ws,header_row)
    records=correlation_reference_records(ws,header_row)
    if not records:return
    rows=[r for r,_ in records];last=max(rows);lastcol=max(headers.values())
    kindcols={correlation_kind(h):c for h,c in headers.items() if correlation_kind(h)}
    rawexp=next((c for h,c in headers.items() if correlation_role(h)=='Experimental' and correlation_is_temperature(h)),None)
    exp=next((c for h,c in headers.items() if correlation_role(h)=='Experimental' and 'delta' in correlation_token(h)),None)
    inlet=kindcols.get('inlet_temperature');power=kindcols.get('power')
    palette=['2878A0','8B5DA8','BD6D25','36866B','BB4D70','5266AE']
    rawcols=[(h,c) for h,c in headers.items() if correlation_role(h)=='Simulation' and correlation_is_temperature(h)]
    configs=[re.split(r'(?i)\bsim\s+',h,maxsplit=1)[-1] for h,c in rawcols]
    configcolors={v:palette[i%len(palette)] for i,v in enumerate(configs)}
    raw_to_delta={}
    for (h,c),label in zip(rawcols,configs):
        found=correlation_find_result(headers,h,label,delta=True)
        if found:raw_to_delta[c]=found[1]
    component=kindcols.get('component');projects={};project_blocks=[]
    group_colors=['E3ECF5','E8E5F3','E3F0EB','F6ECDD']
    # Records above retain the merged label before the cells are made writable.
    if component:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_col==merged.max_col==component and merged.min_row>header_row:
                ws.unmerge_cells(str(merged))
    previous=None
    for r,record in records:
        php=next((record[h] for h,c in headers.items() if c==component),'All PHPs') or 'Unspecified PHP'
        projects.setdefault(str(php),[]).append(r)
        if project_blocks and project_blocks[-1]['name']==str(php) and project_blocks[-1]['last']==r-1:
            project_blocks[-1]['last']=r
        else:
            project_blocks.append({'name':str(php),'first':r,'last':r})
        if component and ws.cell(r,component).value in (None,''):
            correlation_literal(ws.cell(r,component),php)
        group=tuple(correlation_value(record[h],correlation_kind(h)) for h in headers
                    if correlation_kind(h) in ('component','fluid','orientation','filling_ratio','water_flow','htc','inlet_temperature'))
        boundary=previous is None or group!=previous;previous=group
        project_start=r==projects[str(php)][0]
        for h,c in headers.items():
            cell=ws.cell(r,c);cell.font=Font(name='Arial',size=10,color='203344')
            cell.fill=PatternFill('solid',fgColor='F3F6F9' if r%2==0 else 'FFFFFF')
            cell.border=Border(top=Side(style='medium' if project_start else 'thin',color='8295A6')) if boundary else Border()
            numeric=isinstance(cell.value,(int,float)) or cell.data_type=='f'
            cell.alignment=Alignment(horizontal='center' if correlation_kind(h)=='orientation' else 'right' if numeric else 'left',vertical='center')
            if numeric:
                cell.number_format='0%' if correlation_kind(h)=='filling_ratio' and '%' in cell.number_format else '0' if correlation_kind(h) in ('power','filling_ratio') else '#,##0' if correlation_kind(h)=='htc' else '0.0#' if correlation_kind(h)=='water_flow' else '0.0'
            if c==component:
                cell.fill=PatternFill('solid',fgColor=group_colors[list(projects).index(str(php))%len(group_colors)])
                cell.font=Font(name='Arial',size=10,color='203344',bold=True)
        ws.row_dimensions[r].height=21
        if inlet:
            for raw,delta in ([(rawexp,exp)] if rawexp and exp else [])+list(raw_to_delta.items()):
                t=f'{get_column_letter(raw)}{r}';tin=f'{get_column_letter(inlet)}{r}'
                ws.cell(r,delta).value=f'=IF(AND(ISNUMBER({t}),ISNUMBER({tin})),{t}-{tin},"")'

    for h,c in headers.items():
        role=correlation_role(h)
        color='344D63' if role=='Condition' else '26705B'
        if role=='Simulation':
            label=re.split(r'(?i)\bsim\s+',h,maxsplit=1)[-1]
            color=configcolors.get(label,'2878A0')
        cell=ws.cell(header_row,c)
        cell.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
        cell.fill=PatternFill('solid',fgColor=color)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        cell.border=Border(right=Side(style='thin',color='FFFFFF'),bottom=Side(style='medium',color='344D63'))
        width={'component':18,'fluid':19,'orientation':12,'filling_ratio':10,'power':13,'water_flow':14,'htc':16,'inlet_temperature':15}.get(correlation_kind(h),22 if 'evo' in h.lower() else 19)
        ws.column_dimensions[get_column_letter(c)].width=width
    ws.row_dimensions[header_row].height=44
    ws.freeze_panes=f'E{header_row+1}';ws.sheet_view.showGridLines=False;ws.sheet_view.topLeftCell='A1'
    for cell in ws[header_row+3]:
        if cell.value=='Project names repeat on each row to keep filtering and matching reliable.':
            cell.value='Combined plot: solver colors and project marker shapes.'
    # Excel tables cannot contain merged cells. Keep the project label outside
    # the native data table while retaining filters on the test/result columns.
    table_first=2 if component==1 else 1
    ws.auto_filter.ref=f'{get_column_letter(table_first)}{header_row}:{get_column_letter(lastcol)}{last}'
    for table in list(ws.tables.values()):
        lo,top,hi,bottom=range_boundaries(table.ref)
        if lo in (1,table_first) and top==header_row:
            from openpyxl.worksheet.table import TableColumn
            if component and component!=1 and lo<=component<=hi:
                del ws.tables[table.name]
                continue
            table.ref=ws.auto_filter.ref
            table.tableColumns=[TableColumn(id=i,name=str(ws.cell(header_row,c).value)) for i,c in enumerate(range(table_first,lastcol+1),1)]
            if table.autoFilter:table.autoFilter.ref=table.ref
    if component:
        for block in project_blocks:
            first,end=block['first'],block['last']
            if first<end:
                ws.merge_cells(start_row=first,start_column=component,end_row=end,end_column=component)
            cell=ws.cell(first,component)
            correlation_literal(cell,block['name'])
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cell.border=Border(top=Side(style='medium',color='8295A6'),bottom=Side(style='medium',color='8295A6'))
    # Add a dynamic missing-result cue without recoloring valid numeric results.
    for h,c in rawcols:
        letter=get_column_letter(c);area=f'{letter}{header_row+1}:{letter}{last}'
        for region in list(ws.conditional_formatting._cf_rules):
            if str(region.sqref)==area:del ws.conditional_formatting._cf_rules[region]
        ws.conditional_formatting.add(area,FormulaRule(formula=[f'ISBLANK({letter}{header_row+1})'],fill=PatternFill('solid',fgColor='FFF4D6')))
    if not (exp and rawexp and inlet and power and rawcols):return

    # Rebuild correlation scatter plots only. Other chart types are preserved.
    ws._charts=[ch for ch in ws._charts if not isinstance(ch,ScatterChart)]
    left=lastcol+2;right=left+13
    for c in range(left,right+12):ws.column_dimensions[get_column_letter(c)].width=10
    ws.column_dimensions[get_column_letter(left-1)].width=3
    ws.column_dimensions[get_column_letter(right-1)].width=3
    # Cell-backed parity lines survive Excel/Calc export and recalculation.
    reference_sheet=wb['ChartData'] if 'ChartData' in wb.sheetnames else wb.create_sheet('ChartData')
    rises=[]
    for r in rows:
        tin=correlation_value(ws.cell(r,inlet).value)
        if isinstance(tin,(int,float)):
            rises.extend(v-tin for c in [rawexp]+[c for _,c in rawcols] if isinstance((v:=correlation_value(ws.cell(r,c).value)),(int,float)))
    reference_limit=max(100,math.ceil(max(rises,default=100)/10)*10)
    for c,title in enumerate(['Reference ΔT [K]','y = x','+30%','−30%'],12):
        reference_sheet.cell(1,c,title)
    for r,value in [(2,0),(3,reference_limit)]:
        reference_sheet.cell(r,12,value)
        for c,factor in [(13,1),(14,1.3),(15,.7)]:
            reference_sheet.cell(r,c,f'=L{r}*{factor}')
    def setup(title,xlabel,ylabel):
        chart=ScatterChart();chart.title=title;chart.x_axis.title=xlabel;chart.y_axis.title=ylabel
        chart.width=21.5;chart.height=15.5;chart.legend.position='b';chart.display_blanks='gap'
        chart.scatterStyle='lineMarker'
        return chart
    def add_ref(chart, limit):
        for c,(name,factor) in enumerate([('y = x',1),('+30%',1.3),('−30%',.7)],13):
            se=Series(Reference(reference_sheet,min_col=c,min_row=2,max_row=3),Reference(reference_sheet,min_col=12,min_row=2,max_row=3),title=name)
            se.marker.symbol='none';se.graphicalProperties.line.solidFill='657789' if factor==1 else 'AAB5BF'
            se.graphicalProperties.line.width=12000
            if factor!=1:se.graphicalProperties.line.prstDash='dash'
            chart.series.append(se)
    def add_series(chart,xcol,ycol,selected,name,color,marker='circle',line=False,dashed=False):
        # Split discontinuous groups instead of including intervening other tests.
        blocks=[]
        for r in sorted(selected):
            if not blocks or r!=blocks[-1][-1]+1:blocks.append([r])
            else:blocks[-1].append(r)
        for index,block in enumerate(blocks):
            se=Series(Reference(ws,min_col=ycol,min_row=block[0],max_row=block[-1]),Reference(ws,min_col=xcol,min_row=block[0],max_row=block[-1]),title=name if index==0 else name+' (continued)')
            se.graphicalProperties.line.noFill=not line
            if line:
                se.graphicalProperties.line.solidFill=color;se.graphicalProperties.line.width=16000
                if dashed:se.graphicalProperties.line.prstDash='dash'
            se.marker.symbol=marker;se.marker.size=5
            se.marker.graphicalProperties.solidFill=color;se.marker.graphicalProperties.line.solidFill=color
            chart.series.append(se)
            if index:
                from openpyxl.chart.legend import LegendEntry
                chart.legend.legendEntry.append(LegendEntry(idx=len(chart.series)-1,delete=True))
    project_symbols=['circle','square','triangle','diamond','x','star','plus']
    project_markers={name:project_symbols[i%len(project_symbols)] for i,name in enumerate(projects)}
    def parity(title,selected,col,row,by_project=False):
        chart=setup(title,'Experimental ΔTevap [K]','Simulated ΔTevap [K]')
        values=[]
        for r in selected:
            tin=correlation_value(ws.cell(r,inlet).value)
            if isinstance(tin,(int,float)):
                values.extend(v-tin for c in [rawexp]+[c for _,c in rawcols] if isinstance((v:=correlation_value(ws.cell(r,c).value)),(int,float)))
        limit=max(20, math.ceil(max(values,default=50)/10)*10)
        for index,((h,raw),label) in enumerate(zip(rawcols,configs)):
            if raw in raw_to_delta:
                available=[r for r in selected if ws.cell(r,raw).data_type=='f' or isinstance(correlation_value(ws.cell(r,raw).value),(int,float))]
                if available and by_project:
                    for name,project_rows in projects.items():
                        selected_rows=[r for r in available if r in project_rows]
                        if selected_rows:
                            add_series(chart,exp,raw_to_delta[raw],selected_rows,f'{label} — {name}',configcolors[label],project_markers[name])
                elif available:
                    add_series(chart,exp,raw_to_delta[raw],available,label,configcolors[label],['circle','diamond','triangle','square'][index%4])
        add_ref(chart,limit)
        chart.x_axis.scaling.min=chart.y_axis.scaling.min=0
        chart.x_axis.scaling.max=chart.y_axis.scaling.max=limit
        ws.add_chart(chart,f'{get_column_letter(col)}{row}')
    parity('All PHPs: available correlations',rows,left,6,by_project=True)
    first_name=next(iter(projects));parity(first_name+': solver comparison',projects[first_name],right,6)
    # Preserve the two established T4H heat-load views, grouped by actual inputs.
    fluidcol=kindcols.get('fluid');flowcol=kindcols.get('water_flow')
    firstrows=projects[first_name]
    fluids=list(dict.fromkeys(str(ws.cell(r,fluidcol).value).strip() for r in firstrows)) if fluidcol else []
    power_slots=0
    for fluid in fluids:
        selected=[r for r in firstrows if str(ws.cell(r,fluidcol).value).strip()==fluid]
        groups={}
        for r in selected:
            key=tuple(ws.cell(r,kindcols[k]).value for k in ['filling_ratio','orientation','water_flow'] if k in kindcols)
            groups.setdefault(key,[]).append(r)
        if len(groups)>4:continue
        chart=setup(f'{first_name} · {fluid}','Heat load [W]','Evaporator temperature [°C]')
        for index,(group,rs) in enumerate(groups.items()):
            condition=' / '.join(str(v) for v in group)
            fixed = all(len({ws.cell(r,kindcols[k]).value for r in selected}) == 1 for k in ('filling_ratio','orientation') if k in kindcols)
            if flowcol and fixed:
                condition=f'{ws.cell(rs[0],flowcol).value:g} l/min'
            else:
                condition=' / '.join(f'{ws.cell(rs[0],kindcols[k]).value}{unit}' for k,unit in [('filling_ratio','%'),('orientation',''),('water_flow',' l/min')] if k in kindcols)
            marker=['circle','square','triangle','diamond'][index%4]
            add_series(chart,power,rawexp,rs,condition+' Exp','26705B',marker)
            for (h,raw),label in zip(rawcols,configs):
                add_series(chart,power,raw,rs,condition+' '+label,configcolors[label],marker,True,'evo' in label.lower())
        col=left if power_slots%2==0 else right;row=32+27*(power_slots//2)
        ws.add_chart(chart,f'{get_column_letter(col)}{row}');power_slots+=1
    nextrow=32+27*max(1,(power_slots+1)//2)
    for index,(name,rs) in enumerate(list(projects.items())[1:]):
        parity(name+': all recorded conditions',rs,left if index%2==0 else right,nextrow+27*(index//2))


def correlation_save(plan, output):
    """Preserve sheets, formulas, charts and styles; backup + atomic replacement."""
    reference = Path(plan['reference'])
    if (reference.stat().st_mtime_ns, reference.stat().st_size) != plan['reference_signature']:
        raise ValueError('The reference changed after preview. Rebuild the preview before saving.')
    if not plan['counts'].get('Matched'):
        raise ValueError('No unambiguous successful rows to import. Nothing was written.')
    output = Path(output).resolve()
    if output.suffix.lower() != '.xlsx':
        raise ValueError('Save the correlation workbook as .xlsx.')
    wb = openpyxl.load_workbook(reference)
    ws = wb[plan['sheet']]
    titles = correlation_headers(ws, plan['header_row'])
    temperature_metrics = [m for m in plan['metrics'] if correlation_is_temperature(m)]
    names = [f'Tevap Sim {plan["version"]}' if m in temperature_metrics else f'{m} | Sim {plan["version"]}' for m in plan['metrics']]
    if len(set(names)) != len(names):
        names = [f'{m} | Sim {plan["version"]}' for m in plan['metrics']]
    delta_names = [f'Δ{name}' for m, name in zip(plan['metrics'], names) if m in temperature_metrics]
    raw_names = {}
    delta_by_metric = {}
    for metric, proposed in zip(plan['metrics'], names):
        titles = correlation_headers(ws, plan['header_row'])
        found = correlation_find_result(titles, metric, plan['version'])
        if found:
            raw_names[metric] = found[0]
        else:
            columns = [c for h,c in titles.items() if correlation_role(h)=='Simulation' and 'delta' not in correlation_token(h)]
            col = max(columns,default=max(titles.values()))+1
            correlation_insert_columns(wb,ws,col,1)
            correlation_literal(ws.cell(plan['header_row'],col),proposed)
            raw_names[metric] = proposed
    for metric in temperature_metrics:
        titles=correlation_headers(ws,plan['header_row'])
        found=correlation_find_result(titles,metric,plan['version'],delta=True)
        if found:
            delta_by_metric[metric]=found[0]
        else:
            columns=[c for h,c in titles.items() if 'delta' in correlation_token(h)]
            col=max(columns,default=max(titles.values()))+1
            correlation_insert_columns(wb,ws,col,1)
            title='Δ'+raw_names[metric]
            correlation_literal(ws.cell(plan['header_row'],col),title)
            delta_by_metric[metric]=title
    titles=correlation_headers(ws,plan['header_row'])
    raw_cols={m:titles[h] for m,h in raw_names.items()}
    delta_cols={m:titles[h] for m,h in delta_by_metric.items()}
    names=list(raw_names.values());delta_names=list(delta_by_metric.values())
    first=min(raw_cols.values());delta_first=min(delta_cols.values()) if delta_cols else first
    inlet_columns=[c for h,c in titles.items() if correlation_kind(h)=='inlet_temperature']
    if temperature_metrics and len(inlet_columns)!=1:
        raise ValueError('One inlet-temperature column is required for ΔT formulas.')
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    for entry in plan['audit']:
        if entry['status'] != 'Matched':
            continue
        for metric,value in entry['results'].items():
            cell=ws.cell(entry['reference_row'],raw_cols[metric])
            cell.value=value
            cell.number_format='0.0'
    if delta_cols:
        inlet=inlet_columns[0]
        keycols=[titles[name] for name in plan['mappings']]
        for r in range(plan['header_row']+1,ws.max_row+1):
            if not any(ws.cell(r,c).value is not None for c in keycols):
                continue
            for metric,col in delta_cols.items():
                raw=f'{get_column_letter(raw_cols[metric])}{r}'; tin=f'{get_column_letter(inlet)}{r}'
                ws.cell(r,col).value=f'=IF(AND(ISNUMBER({raw}),ISNUMBER({tin})),{raw}-{tin},"")'
                ws.cell(r,col).number_format='0.0'
    correlation_format_and_chart(wb,ws,plan['header_row'])
    if wb.calculation:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    # Focus the newly inserted results on opening.
    ws.sheet_view.topLeftCell = 'A1'
    wb.active = wb.sheetnames.index(ws.title)
    output.parent.mkdir(parents=True, exist_ok=True)
    backup = None
    handle, temp = tempfile.mkstemp(suffix='.xlsx', dir=output.parent)
    os.close(handle)
    try:
        wb.save(temp)
        wb.close()
        # Verify the written archive before touching the destination.
        check = openpyxl.load_workbook(temp, read_only=True)
        check.close()
        if output.exists():
            backup = output.with_name(f'{output.stem}.backup-{stamp}.xlsx')
            shutil.copy2(output, backup)
        os.replace(temp, output)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
    return {'output': str(output), 'backup': str(backup) if backup else None, 'columns': names, 'first_column': get_column_letter(first)}


def correlation_scroll_area(root):
    from tkinter import Canvas, ttk
    frame = ttk.Frame(root)
    frame.pack(fill='both', expand=True, padx=12, pady=8)
    canvas = Canvas(frame, highlightthickness=0)
    scrollbar = ttk.Scrollbar(frame, orient='vertical', command=canvas.yview)
    canvas.configure(yscrollcommand=scrollbar.set)
    scrollbar.pack(side='right', fill='y')
    canvas.pack(side='left', fill='both', expand=True)
    body = ttk.Frame(canvas)
    item = canvas.create_window((0, 0), window=body, anchor='nw')
    body.bind('<Configure>', lambda event: canvas.configure(scrollregion=canvas.bbox('all')))
    canvas.bind('<Configure>', lambda event: canvas.itemconfigure(item, width=event.width))
    return body


def correlation_choose_sheet(root, path):
    import tkinter as tk
    from tkinter import ttk, simpledialog
    workbook = openpyxl.load_workbook(path, read_only=True)
    names = workbook.sheetnames
    workbook.close()
    name = names[0]
    if len(names) > 1:
        dialog = tk.Toplevel(root)
        dialog.title('Select worksheet')
        chosen = tk.StringVar(value=name)
        result = []
        ttk.Label(dialog, text=Path(path).name).pack(padx=15, pady=10)
        ttk.Combobox(dialog, textvariable=chosen, values=names, state='readonly', width=45).pack(padx=15, pady=10)
        def accept():
            result.append(chosen.get())
            dialog.destroy()
        ttk.Button(dialog, text='Continue', command=accept).pack(pady=10)
        dialog.grab_set()
        root.wait_window(dialog)
        if not result:
            return None, None
        name = result[0]
    row = simpledialog.askinteger('Column titles', f'{Path(path).name}\nWhich row contains the column titles?', initialvalue=1, minvalue=1, parent=root)
    return name, row


def run_correlation_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    root = tk.Tk()
    root.withdraw()
    try:
        reference = filedialog.askopenfilename(parent=root, title='Select reference experimental / simulation workbook', filetypes=[('Excel workbook', '*.xlsx')])
        if not reference:
            return
        sheet, header_row = correlation_choose_sheet(root, reference)
        if not sheet or not header_row:
            return
        source = filedialog.askopenfilename(parent=root, title='Select newly simulated data', filetypes=[('Excel or CSV', '*.xlsx *.csv')])
        if not source:
            return
        if Path(source).resolve() == Path(reference).resolve():
            raise ValueError('Select a separate simulation file.')
        source_sheet, source_header = None, 1
        if Path(source).suffix.lower() == '.xlsx':
            source_sheet, source_header = correlation_choose_sheet(root, source)
            if not source_sheet or not source_header:
                return
        source_titles, records = correlation_read_source(source, source_sheet, source_header)
        workbook = openpyxl.load_workbook(reference, read_only=True, data_only=True)
        headers = correlation_headers(workbook[sheet], header_row)
        defaults = correlation_auto_mapping(workbook[sheet], header_row, source_titles)
        reference_values = {title: sorted({str(row[col-1]) for row in workbook[sheet].iter_rows(min_row=header_row+1, values_only=True)
                                          if row[col-1] is not None}) for title, col in headers.items()}
        workbook.close()
        root.title('Correlation analysis — test conditions, physics and solver')
        root.geometry('1180x740')
        root.deiconify()
        match = re.search(r'v\d+(?:\.\d+)+', Path(source).name, re.I)
        existing_labels=list(dict.fromkeys(re.split(r'(?i)\bsim\s+',h,maxsplit=1)[-1] for h in headers if correlation_role(h)=='Simulation' and 'delta' not in correlation_token(h)))
        suggested=next((v for v in existing_labels if 'evo' in v.lower()),'')
        version = tk.StringVar(value=suggested)
        signatures=list(dict.fromkeys(correlation_physics_signature(record) for _,record in records))
        signature_labels=['; '.join(f'{k}={v or "unspecified"}' for k,v in signature if v) or 'No physics metadata in export' for signature in signatures]
        physics_selection=tk.StringVar(value=signature_labels[0] if len(signature_labels)==1 else '')
        top = ttk.Frame(root)
        top.pack(fill='x', padx=12, pady=10)
        ttk.Label(top, text=f'Reference: {Path(reference).name} / {sheet}\nSimulation: {Path(source).name} — {len(records)} rows').pack(anchor='w')
        ttk.Label(top, text='Physics / solver label (select existing to update it; enter a new label to add a comparison)').pack(anchor='w', pady=(8, 0))
        ttk.Combobox(top, textvariable=version, values=existing_labels, width=65).pack(anchor='w')
        ttk.Label(top,text='Source configuration').pack(anchor='w')
        ttk.Combobox(top,textvariable=physics_selection,values=signature_labels,state='readonly',width=135).pack(anchor='w')
        ttk.Label(top, text='Shared conditions are mapped automatically. Missing conditions that are constant in the reference are filled automatically.\nHTC takes priority over VFR/MFR. PHP/project, FR and orientation must match. If absent in the export, enter explicit constants or map angle/roll pairs to ES/H/U.', wraplength=1120).pack(anchor='w', pady=8)
        body = correlation_scroll_area(root)
        for c, title in enumerate(['Reference column', 'Role', 'Simulation column', 'Constant value', 'Optional value mapping']):
            ttk.Label(body, text=title).grid(row=0, column=c, sticky='w', padx=4)
        controls = {}
        def edit_values(title, column, lookup, button):
            if column.get() not in source_titles:
                messagebox.showerror('Choose a source column', 'Select a simulation column before mapping values.', parent=root)
                return
            dialog = tk.Toplevel(root)
            dialog.title(f'{column.get()} → {title}')
            dialog.geometry('610x520')
            ttk.Label(dialog, text='Enter the corresponding reference value for each source value.\nBlank entries will be flagged as missing conditions.', wraplength=570).pack(padx=12, pady=12)
            fields = correlation_scroll_area(dialog)
            unique = sorted({str(record.get(column.get())) for _, record in records if record.get(column.get()) is not None})
            old = json.loads(lookup.get()) if lookup.get() else {}
            variables = {}
            for index, value in enumerate(unique):
                ttk.Label(fields, text=value, wraplength=230).grid(row=index, column=0, padx=8, pady=4, sticky='w')
                variable = tk.StringVar(value=str(old.get(value, '')))
                ttk.Combobox(fields, textvariable=variable, values=reference_values[title], width=30).grid(row=index, column=1, padx=8, pady=4)
                variables[value] = variable
            buttons = ttk.Frame(dialog)
            buttons.pack(fill='x', padx=12, pady=12)
            def apply_values(clear=False):
                values = {value: var.get() for value, var in variables.items() if var.get().strip()}
                lookup.set('' if clear else json.dumps(values))
                button.configure(text='Map values…' if clear else f'Edit mapping ({len(values)})')
                dialog.destroy()
            ttk.Button(buttons, text='Use values unchanged', command=lambda: apply_values(True)).pack(side='left')
            ttk.Button(buttons, text='Apply mapping', command=apply_values).pack(side='right')
            dialog.grab_set()
        for r, title in enumerate(headers, 1):
            preset = defaults[title]
            role = tk.StringVar(value=preset['role'])
            column = tk.StringVar(value=preset['column'])
            constant, lookup = tk.StringVar(value=preset['constant']), tk.StringVar()
            ttk.Label(body, text=title, wraplength=210).grid(row=r, column=0, sticky='w', padx=4, pady=4)
            ttk.Combobox(body, values=['Condition', 'Experimental', 'Simulation', 'Ignore'], textvariable=role, state='readonly', width=14).grid(row=r, column=1, padx=4)
            ttk.Combobox(body, values=['', '<constant>']+source_titles, textvariable=column, state='readonly', width=26).grid(row=r, column=2, padx=4)
            ttk.Entry(body, textvariable=constant, width=16).grid(row=r, column=3, padx=4)
            button = ttk.Button(body, text='Map values…')
            button.configure(command=lambda t=title, c=column, l=lookup, b=button: edit_values(t, c, l, b))
            button.grid(row=r, column=4, padx=4)
            # A mapping belongs to one source column; never reuse it silently.
            def reset_mapping(*args, l=lookup, b=button):
                l.set('')
                b.configure(text='Map values…')
            column.trace_add('write', reset_mapping)
            controls[title] = (role, column, constant, lookup)
        metric_start = len(headers)+2
        ttk.Label(body, text='Simulation results to import (choose the averaged evaporator temperature)').grid(row=metric_start, column=0, columnspan=5, sticky='w', pady=12)
        metric_controls = {}
        for r, title in enumerate(source_titles, metric_start+1):
            suggested = correlation_is_temperature(title)
            selected = tk.BooleanVar(value=suggested)
            ttk.Checkbutton(body, text=title, variable=selected).grid(row=r, column=0, columnspan=5, sticky='w', pady=2)
            metric_controls[title] = selected
        footer = ttk.Frame(root)
        footer.pack(fill='x', padx=12, pady=12)

        def preview_save():
            try:
                mappings = {}
                for title, (role, column, constant, lookup) in controls.items():
                    if role.get() != 'Condition':
                        continue
                    if not column.get():
                        raise ValueError(f'Choose a simulation column or explicit constant for {title}.')
                    mapping = {'constant': constant.get()} if column.get() == '<constant>' else {'column': column.get()}
                    if 'constant' in mapping and not mapping['constant'].strip():
                        raise ValueError(f'Enter the constant for {title}.')
                    if lookup.get().strip():
                        mapping['values'] = json.loads(lookup.get())
                        if not isinstance(mapping['values'], dict):
                            raise ValueError(f'The value map for {title} must be a JSON object.')
                    mappings[title] = mapping
                metrics = [title for title, selected in metric_controls.items() if selected.get()]
                if physics_selection.get() not in signature_labels:
                    raise ValueError('Select one source physics / solver configuration.')
                selected_signature=signatures[signature_labels.index(physics_selection.get())]
                selected_records=[item for item in records if correlation_physics_signature(item[1])==selected_signature]
                plan = correlation_plan(reference, sheet, header_row, selected_records, mappings, metrics, version.get())
                plan['source'] = str(Path(source).resolve())
                summary = '\n'.join(f'{key}: {value}' for key, value in plan['counts'].items())
                preview = tk.Toplevel(root)
                preview.title('Review correlation import')
                preview.geometry('900x570')
                ttk.Label(preview, text=f'{summary}\n\nOnly matched, successful averages will be imported. Existing values for unmatched or failed rows are preserved.\nThe full physics / solver label identifies the destination column; different labels stay separate.', wraplength=850).pack(anchor='w', padx=12, pady=12)
                area = ttk.Frame(preview)
                area.pack(fill='both', expand=True, padx=12)
                tree = ttk.Treeview(area, columns=('source', 'ref', 'status', 'detail'), show='headings')
                for name, width in [('source', 80), ('ref', 80), ('status', 170), ('detail', 450)]:
                    tree.heading(name, text=name.title())
                    tree.column(name, width=width)
                scroll = ttk.Scrollbar(area, orient='vertical', command=tree.yview)
                tree.configure(yscrollcommand=scroll.set)
                scroll.pack(side='right', fill='y')
                tree.pack(fill='both', expand=True)
                for entry in plan['audit']:
                    tree.insert('', 'end', values=(entry['source_row'], entry['reference_row'] or '', entry['status'], entry['detail']))
                def save():
                    try:
                        output = filedialog.asksaveasfilename(parent=preview, title='Save updated reference (overwrites are backed up)', initialdir=str(Path(reference).parent), initialfile=Path(reference).name, defaultextension='.xlsx', filetypes=[('Excel workbook', '*.xlsx')])
                        if not output:
                            return
                        if Path(output).resolve() == Path(source).resolve():
                            raise ValueError('The output cannot overwrite the simulation source.')
                        result = correlation_save(plan, output)
                        text = f'Saved {plan["counts"].get("Matched", 0)} matched rows.\n{output}\nNew results start in column {result["first_column"]}.'
                        if result['backup']:
                            text += f'\nBackup: {result["backup"]}'
                        messagebox.showinfo('Correlation complete', text, parent=preview)
                        preview.destroy()
                        root.destroy()
                    except Exception as error:
                        messagebox.showerror('Cannot save correlation', str(error), parent=preview)
                ttk.Button(preview, text='Save updated workbook', command=save, state='normal' if plan['counts'].get('Matched') else 'disabled').pack(pady=12)
                preview.grab_set()
            except Exception as error:
                messagebox.showerror('Check correlation settings', str(error), parent=root)

        ttk.Button(footer, text='Preview matches', command=preview_save).pack(side='right')
        ttk.Button(footer, text='Cancel', command=root.destroy).pack(side='left')
        root.mainloop()
    except Exception as error:
        messagebox.showerror('Correlation analysis', str(error), parent=root)
    finally:
        try:
            root.destroy()
        except tk.TclError:
            pass


def launch_analysis_gui():
    import tkinter as tk
    from tkinter import ttk
    root = tk.Tk()
    root.title('JJCooling Data Processor')
    root.geometry('480x230')
    selected = []
    ttk.Label(root, text='What would you like to analyse?', font=('Segoe UI', 14)).pack(pady=24)
    def choose(mode):
        selected.append(mode)
        root.destroy()
    ttk.Button(root, text='Test data — generate Excel and PDF report', command=lambda: choose('test')).pack(fill='x', padx=30, pady=7)
    ttk.Button(root, text='Correlation — compare physics / solvers with tests', command=lambda: choose('correlation')).pack(fill='x', padx=30, pady=7)
    root.mainloop()
    if selected == ['test']:
        main()
    elif selected == ['correlation']:
        run_correlation_gui()


# Shift2DC: enter new schedules once, then reuse persisted per-CPU heat loads.
def is_shift2dc_file(path):
    return re.sub(r'[^a-z0-9]', '', Path(path).stem.split('_')[0].lower()) == 'shift2dc'


def shift2dc_time(data):
    column = find_optional_column(data, ['RelTime', 'REL_TIME', 'RelativeTime', 'ElapsedTime', 'Seconds', 'Time_s', 'Time [s]'])
    if column is None:
        raise ValueError('Shift2DC requires an elapsed-time column in seconds (for example RelTime).')
    values = pd.to_numeric(data[column], errors='coerce').to_numpy(dtype=float)
    if len(values) < 2 or not np.isfinite(values).all() or np.any(np.diff(values) <= 0):
        raise ValueError('Elapsed time must be finite and strictly increasing, without duplicate timestamps.')
    return values


def shift2dc_board_columns(data):
    """Accept CPU sensors and legacy BOARD names, counting each physical ID once.

    If both names exist for one ID, the explicit T_CPU channel takes precedence.
    """
    sensors = {}
    for column in data.columns:
        match = re.fullmatch(r'T_(CPU|BOARD)_(\d+)(?:\s*\[.*\])?', str(column).strip(), re.I)
        if match:
            sensor_id = int(match[2])
            if sensor_id not in sensors or match[1].upper() == 'CPU':
                sensors[sensor_id] = column
    return [sensors[sensor_id] for sensor_id in sorted(sensors)]


def shift2dc_columns_for_ids(data, sensor_ids):
    available = shift2dc_board_columns(data)
    by_id = {int(re.search(r'_(\d+)', column).group(1)): column for column in available}
    try:
        return [by_id[int(sensor_id)] for sensor_id in sensor_ids]
    except KeyError as error:
        raise ValueError(f'No T_CPU or T_BOARD sensor found for ID {error.args[0]}.') from error


def shift2dc_selected_boards(data, config):
    columns = shift2dc_board_columns(data)
    count = int(config['board_count'])
    if count != float(config['board_count']) or count < 1 or count > len(columns):
        raise ValueError(f'Connected board count must be between 1 and {len(columns)}.')
    selected = config.get('board_columns') or columns[:count]
    if len(selected) != count or len(set(selected)) != count or any(c not in columns for c in selected):
        raise ValueError('Choose exactly one available temperature sensor per connected board.')
    return selected


def detect_shift2dc_start(data, board_columns):
    """Estimate the first sustained board rise above a quiet 30-second baseline.

    The median board-to-water temperature suppresses common coolant drift and
    isolated sensor spikes. A local two-line fit refines the threshold crossing.
    This is a thermal estimate, not a recorded electrical switch-on timestamp.
    """
    t = shift2dc_time(data)
    board = data[board_columns].apply(pd.to_numeric, errors='coerce')
    if not np.isfinite(board.to_numpy()).all():
        raise ValueError('Selected board channels contain missing or invalid temperatures.')
    y = board.median(axis=1).to_numpy(dtype=float)
    inlet = find_optional_column(data, ['T_WATER_IN'])
    if inlet:
        coolant = pd.to_numeric(data[inlet], errors='coerce').to_numpy(dtype=float)
        if np.isfinite(coolant).all():
            y -= coolant
    if t[-1]-t[0] < 70:
        raise ValueError('At least 70 seconds of data are needed for automatic start detection. Enter the start manually.')
    grid = np.arange(t[0], t[-1]+.1, 1.0)
    smoothed = pd.Series(np.interp(grid,t,y)).rolling(5,center=True,min_periods=1).median().to_numpy()
    baseline = smoothed[:30]
    level = float(np.median(baseline))
    noise = 1.4826*float(np.median(np.abs(baseline-level)))
    threshold = max(.25, 8*noise)
    if abs(float(np.polyfit(np.arange(30),baseline,1)[0])) > .015:
        raise ValueError('The recording starts with a temperature trend. Enter the test start manually.')
    crossing = None
    for i in range(30,len(grid)-30):
        future = smoothed[i:i+25]
        if (smoothed[i] > level+threshold and np.mean(future > level+threshold*.7) >= .9
                and smoothed[i+20]-smoothed[max(0,i-10)] > max(.4,10*noise)):
            crossing=i
            break
    if crossing is None:
        raise ValueError('No sustained temperature rise was found. Enter the test start manually.')
    lo=max(0,crossing-40);hi=min(len(grid),crossing+15)
    x=grid[lo:hi];signal=smoothed[lo:hi]
    candidates=[]
    for index in range(max(lo+10,crossing-30),crossing+1):
        onset=grid[index]
        design=np.column_stack([np.ones(len(x)),np.maximum(x-onset,0)])
        coeff=np.linalg.lstsq(design,signal,rcond=None)[0]
        if coeff[1] > .005:
            candidates.append((float(np.sum((signal-design@coeff)**2)),onset))
    onset=min(candidates)[1] if candidates else grid[crossing]
    return float(onset), f'First sustained board-temperature rise; thermal start estimate {onset:.1f} s.'


def cpu_power_result_name(board):
    return 'W_CPU_' + str(int(re.search(r'_(\d+)', board).group(1))) + ' [W]'


def validate_cpu_power_matrix(values, step_count, cpu_count):
    matrix = np.asarray(values, dtype=float)
    if matrix.shape != (step_count, cpu_count):
        raise ValueError(f'Expected {step_count} rows (steps) and {cpu_count} columns (CPUs).')
    if not np.isfinite(matrix).all() or (matrix < 0).any() or (matrix.sum(axis=1) <= 0).any():
        raise ValueError('CPU heat loads must be finite and nonnegative, with a positive total for each step.')
    return matrix


def parse_cpu_power_table(text, boards, step_count):
    """Accept pasted/CSV step rows, optionally headed W_CPU_<physical ID>."""
    lines = [line for line in str(text).lstrip('\ufeff').splitlines() if line.strip()]
    if not lines:
        raise ValueError('Enter or import the per-CPU heat-load table.')
    delimiter = '\t' if '\t' in lines[0] else ';' if ';' in lines[0] else ','
    rows = list(csv.reader(lines, delimiter=delimiter))
    header_ids = []
    for token in rows[0]:
        match = re.fullmatch(r'(?:W_)?CPU_?(\d+)(?:\s*\[W\])?', token.strip(), re.I)
        header_ids.append(int(match[1]) if match else None)
    expected_ids = [int(re.search(r'_(\d+)', board).group(1)) for board in boards]
    if any(i is not None for i in header_ids):
        if None in header_ids or len(set(header_ids)) != len(header_ids) or set(header_ids) != set(expected_ids):
            raise ValueError('CSV headers must list each selected CPU exactly once, e.g. W_CPU_1,W_CPU_3.')
        rows = rows[1:]
        if any(len(row) != len(header_ids) for row in rows):
            raise ValueError('Every heat-load row must have one value per CPU.')
        rows = [[row[header_ids.index(i)] for i in expected_ids] for row in rows]
    try:
        values = [[float(value.strip().replace(',', '.') if delimiter != ',' else value.strip()) for value in row] for row in rows]
    except ValueError as error:
        raise ValueError('Use numeric heat loads in watts; one row per step, one column per selected CPU.') from error
    return validate_cpu_power_matrix(values, step_count, len(boards)).tolist()


def shift2dc_step_power_label(row):
    uniform = row.get('Power per board [W]', np.nan)
    return f'{uniform:g} W/CPU' if np.isfinite(uniform) else f"{row['Scheduled total power [W]']:g} W total (mixed CPU loads)"


def build_shift2dc_steps(data, config):
    t=shift2dc_time(data)
    boards=shift2dc_selected_boards(data,config)
    durations=np.asarray(config['durations_s'],dtype=float)
    if durations.ndim!=1 or not len(durations) or not np.isfinite(durations).all() or np.any(durations<100):
        raise ValueError('Each step needs a duration of at least 100 seconds.')
    if config.get('cpu_powers') is not None:
        cpu_powers=validate_cpu_power_matrix(config['cpu_powers'],len(durations),len(boards))
    else:
        powers=np.asarray(config['powers_per_board'],dtype=float)
        if powers.ndim!=1 or len(powers)!=len(durations) or not np.isfinite(powers).all() or np.any(powers<=0):
            raise ValueError('Enter a positive heat load per CPU for every step.')
        cpu_powers=np.repeat(powers[:,None],len(boards),axis=1)
    if config.get('start_s') is None:
        start,note=detect_shift2dc_start(data,boards)
    else:
        start=float(config['start_s']);note=config.get('start_note','User-specified start time.')
    dt=float(np.median(np.diff(t)))
    if not np.isfinite(start) or start<t[0] or start+sum(durations)>t[-1]+dt+1e-8:
        raise ValueError('The complete schedule must fit inside the recorded data; shorten it or adjust the start.')
    steps=[];schedule=[]
    for index,(cpu_loads,duration) in enumerate(zip(cpu_powers,durations),1):
        total_power=float(cpu_loads.sum())
        uniform_power=float(cpu_loads[0]) if np.allclose(cpu_loads,cpu_loads[0],rtol=0,atol=1e-8) else np.nan
        cpu_values={cpu_power_result_name(board):float(load) for board,load in zip(boards,cpu_loads)}
        end=start+float(duration);window_start=end-100
        # [start, end): never include a point from the next power step/cool-down.
        mask=(t>=window_start)&(t<end)
        positions=np.flatnonzero(mask)
        if len(positions)<2:
            raise ValueError(f'Step {index}: too few samples in the final 100 seconds.')
        gaps=np.diff(np.r_[window_start,t[positions],end])
        if max(gaps)>max(5,3*dt):
            raise ValueError(f'Step {index}: a logging gap prevents a complete final-100-second average.')
        window=data.iloc[positions].copy()
        # Piecewise-constant, time-weighted averages. At 1 Hz this is exactly
        # the normal 100-sample mean; irregular sampling is weighted by time.
        previous=int(np.searchsorted(t,window_start,side='right')-1)
        if t[positions[0]]>window_start:
            positions=np.r_[previous,positions]
            window=data.iloc[positions].copy()
        if len(positions) > 1 and np.max(np.diff(t[positions])) > max(5, 3*dt):
            raise ValueError(f'Step {index}: a logging gap crosses the averaging-window boundary.')
        left=np.maximum(t[positions],window_start)
        right=np.r_[t[positions[1:]],end]
        weights=right-left
        required=boards+list(get_medium_columns(data,'Water','VFR'))+lts_temperature_columns
        for column in required:
            actual=find_column(window,[column])
            values=pd.to_numeric(window[actual],errors='coerce').to_numpy(dtype=float)
            if not np.isfinite(values).all():
                raise ValueError(f'Step {index}: invalid values in {actual}.')
            window[actual]=float(np.average(values,weights=weights))
        window['W_PSU']=total_power  # Internal calculation adapter only.
        window.attrs.update(plateau_duration_s=float(duration),sample_count=int(mask.sum()),averaging_duration_s=100.0)
        steps.append(window)
        schedule.append({'Step':index,'Start [s]':start,'End [s]':end,'Average from [s]':window_start,
                         'Average to [s]':end,'Duration [s]':float(duration),'Power per board [W]':uniform_power, **cpu_values,
                         'Connected boards':len(boards),'Scheduled total power [W]':total_power,
                         'Sample Count':int(mask.sum()),'Start method':note})
        start=end
    return boards,steps,schedule


SHIFT2DC_LEGACY_SCHEDULE_COLUMNS = (
    'Heat_Load_Step', 'Heat_Load_Step_Start_s', 'Heat_Load_Step_End_s',
    'Total_Heat_Load_W', 'Scheduled_Power_Per_Board_W',
)


def shift2dc_cpu_power_columns(data):
    """Return saved CPU channels by physical sensor ID, including unit headers."""
    columns = {}
    for column in data.columns:
        match = re.fullmatch(r'W_CPU_(\d+)(?:\s*\[W\])?', str(column).strip(), re.I)
        if match:
            sensor_id = int(match[1])
            if sensor_id in columns:
                raise ValueError(f'Duplicate W_CPU power channels for CPU {sensor_id}.')
            columns[sensor_id] = column
    return dict(sorted(columns.items()))


def clean_shift2dc_power_channels(data):
    """Keep CPU loads and measured channels; remove legacy schedule/PSU exports."""
    legacy = {name.casefold() for name in SHIFT2DC_LEGACY_SCHEDULE_COLUMNS}
    obsolete = [column for column in data.columns
                if str(column).strip().casefold() in legacy or re.fullmatch(
                    r'[WIV]_PSU(?:_\d+)?(?:_SP)?', re.split(r'[\[(]', str(column))[0].strip(), re.I)]
    cleaned = data.drop(columns=obsolete).copy()
    cpu_columns = list(shift2dc_cpu_power_columns(cleaned).values())
    if cpu_columns:
        # Earlier exports left every CPU blank before/after the applied schedule.
        # Those are off intervals. Partial blanks/invalid readings remain visible
        # to validation rather than silently becoming a partial or zero total.
        off = cleaned[cpu_columns].isna().all(axis=1)
        cleaned.loc[off, cpu_columns] = 0.0
    return cleaned


def save_shift2dc_schedule_columns(data, boards, schedule):
    """Persist only per-CPU loads, with zero before/after the applied schedule."""
    saved = clean_shift2dc_power_channels(data)
    saved = saved.drop(columns=list(shift2dc_cpu_power_columns(saved).values()))
    power_columns = ["W_CPU_" + str(int(re.search(r"_(\d+)", column).group(1))) for column in boards]
    for column in power_columns:
        saved[column] = 0.0
    t = shift2dc_time(saved)
    for row in schedule:
        mask = (t >= row['Start [s]']) & (t < row['End [s]'])
        saved.loc[mask, power_columns] = [row.get(c+' [W]', row['Power per board [W]']) for c in power_columns]
    return saved


def build_shift2dc_saved_steps(data):
    """Recover plateaus from changes in the complete CPU power vector.

    Zero/blank intervals separate applied heat-load steps. Legacy schedule
    metadata is ignored; boundaries are inferred from sample timestamps.
    """
    columns = shift2dc_cpu_power_columns(data)
    if not columns:
        return None
    t = shift2dc_time(data)
    dt = float(np.median(np.diff(t)))
    power_frame = data[list(columns.values())].apply(pd.to_numeric, errors='coerce')
    powers = power_frame.to_numpy(dtype=float)
    nonempty = data[list(columns.values())].notna().to_numpy()
    if np.any(nonempty & ~np.isfinite(powers)) or np.any(powers[np.isfinite(powers)] < 0):
        raise ValueError('Saved W_CPU columns contain invalid or negative loads. Correct the CPU power data.')
    complete = np.isfinite(powers).all(axis=1)
    if np.any(np.isfinite(powers).any(axis=1) & ~complete):
        raise ValueError('Saved W_CPU columns are incomplete: some CPUs have power while others are blank.')
    active = complete & (np.nansum(powers, axis=1) > 0)
    if not active.any():
        raise ValueError('W_CPU columns exist but contain no positive heat-load steps. Remove empty W_CPU columns to enter a new schedule.')
    boards = shift2dc_columns_for_ids(data, columns.keys())
    steps, schedule = [], []
    index = 0
    while index < len(t):
        if not active[index]:
            index += 1
            continue
        first = index
        while index+1 < len(t) and active[index+1] and np.allclose(powers[index+1], powers[first], rtol=0, atol=1e-8):
            index += 1
        stop = index+1
        start = t[first]
        end = t[stop] if stop < len(t) else t[-1]+dt
        # Ignore short synchronization artifacts, consistently with 100 s averaging.
        if end-start >= 100-1e-8:
            config = dict(board_count=len(boards), board_columns=boards, start_s=float(start),
                          cpu_powers=[powers[first].tolist()], durations_s=[float(end-start)],
                          start_note='Saved W_CPU plateaus (entered heat loads, not measured electrical power).')
            _, new_steps, new_schedule = build_shift2dc_steps(data, config)
            new_schedule[0]['Step'] = len(schedule)+1
            steps.extend(new_steps)
            schedule.extend(new_schedule)
        index = stop
    if not steps:
        raise ValueError('No W_CPU plateau lasts at least 100 seconds.')
    return boards, steps, schedule


def trim_shift2dc_cooldown(data):
    """Keep recorded samples through 100 s after the last CPU power interval.

    Called after saved-step validation. Use all powered samples, including a
    short final interval excluded from steady-state averages, so no later heat
    load is mistaken for cooldown. The first off sample marks the step end.
    """
    columns = list(shift2dc_cpu_power_columns(data).values())
    if not columns:
        return data
    t = shift2dc_time(data)
    powers = data[columns].apply(pd.to_numeric, errors='raise').to_numpy(dtype=float)
    powered = np.flatnonzero(np.any(powers > 0, axis=1))
    if not len(powered) or powered[-1] == len(t)-1:
        return data
    cutoff = t[powered[-1]+1] + 100.0
    return data.loc[t <= cutoff].copy()


def overwrite_shift2dc_csv(path, data):
    """Replace one CSV atomically so failed writes cannot truncate raw data."""
    path = Path(path)
    handle = tempfile.NamedTemporaryFile(mode='w', dir=path.parent, prefix=path.stem+'.',
                                         suffix='.tmp', delete=False, encoding='utf-8-sig', newline='')
    temporary = Path(handle.name)
    try:
        with handle:
            data.to_csv(handle, index=False, lineterminator='\n')
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def show_shift2dc_dialog(data, source_name):
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
    root=tk.Tk();root.title('Shift2DC server test')
    root.geometry(f"{min(1100,root.winfo_screenwidth()-60)}x{min(900,root.winfo_screenheight()-80)}")
    actions=ttk.Frame(root);actions.pack(side='bottom',fill='x',padx=14,pady=10)
    body_canvas=tk.Canvas(root,highlightthickness=0)
    body_scroll=ttk.Scrollbar(root,orient='vertical',command=body_canvas.yview)
    body_scroll.pack(side='right',fill='y');body_canvas.pack(fill='both',expand=True)
    body_canvas.configure(yscrollcommand=body_scroll.set)
    body=ttk.Frame(body_canvas)
    body_window=body_canvas.create_window((0,0),window=body,anchor='nw')
    body.bind('<Configure>',lambda _event:body_canvas.configure(scrollregion=body_canvas.bbox('all')))
    body_canvas.bind('<Configure>',lambda event:body_canvas.itemconfigure(body_window,width=event.width))
    available=shift2dc_board_columns(data)
    if not available:
        root.destroy();raise ValueError('No T_CPU_<number> or legacy T_BOARD_<number> channels found.')
    fields=ttk.Frame(body);fields.pack(fill='x',padx=14,pady=12)
    ttk.Label(fields,text=source_name,wraplength=980).grid(row=0,column=0,columnspan=2,sticky='w',pady=5)
    fields.columnconfigure(1,weight=1)
    variables={}
    entries_by_key={}
    entries=[('board_count','Connected boards',str(len(available))),
             ('board_ids','CPU IDs (optional; default first N sensors)',''),
             ('step_count','Number of heat-load steps','3'),
             ('powers','Heat load per CPU [W], comma-separated',''),
             ('durations','Step durations [s]: one value for all, or comma-separated',''),
             ('start','Test start [s], detected below and editable','')]
    for row,(key,label,default) in enumerate(entries,1):
        ttk.Label(fields,text=label).grid(row=row,column=0,sticky='w',padx=4,pady=3)
        variables[key]=tk.StringVar(value=default)
        entry=ttk.Entry(fields,textvariable=variables[key],width=55)
        entry.grid(row=row,column=1,sticky='ew',padx=4,pady=3)
        entries_by_key[key]=entry
    ttk.Label(fields, text='On confirmation, CPU heat-load columns are saved in this raw CSV for future reports.',
              wraplength=980).grid(row=7,column=0,columnspan=2,sticky='w',pady=(6,0))
    heterogeneous=tk.BooleanVar(value=False)
    ttk.Checkbutton(fields,text='Heterogeneous CPU heat loads (different watts for each CPU)',
                    variable=heterogeneous).grid(row=8,column=0,columnspan=2,sticky='w',pady=(8,0))
    matrix_frame=ttk.LabelFrame(fields,text='CPU heat loads [W]',padding=8)
    matrix_frame.grid(row=9,column=0,columnspan=2,sticky='ew',pady=5)
    ttk.Label(matrix_frame,text='One row per step; one column per selected CPU, in the selected CPU order. '
              'Optional headers: W_CPU_1,W_CPU_2,... . Separate values with commas, semicolons or tabs.',
              wraplength=950).pack(anchor='w')
    matrix_text=tk.Text(matrix_frame,height=5,width=80,wrap='none')
    matrix_scroll=ttk.Scrollbar(matrix_frame,orient='horizontal',command=matrix_text.xview)
    matrix_text.configure(xscrollcommand=matrix_scroll.set)
    matrix_text.pack(fill='x',pady=4);matrix_scroll.pack(fill='x')
    matrix_buttons=ttk.Frame(matrix_frame);matrix_buttons.pack(fill='x')
    def import_cpu_csv():
        path=filedialog.askopenfilename(parent=root,title='Import per-CPU heat loads',
                                        filetypes=[('CSV / text tables','*.csv *.txt *.tsv'),('All files','*.*')])
        if not path:return
        try:
            text=Path(path).read_text(encoding='utf-8-sig')
            matrix_text.delete('1.0','end');matrix_text.insert('1.0',text)
        except (OSError,UnicodeError) as error:
            messagebox.showerror('Cannot import heat loads',str(error),parent=root)
    def insert_cpu_template():
        try:
            columns=shift2dc_selected_boards(data,selected_config())
            n=int(variables['step_count'].get())
            if n<1:raise ValueError('Enter at least one step.')
            header=','.join(cpu_power_result_name(c).replace(' [W]','') for c in columns)
            text=header+'\n'+('\n'.join([','.join(['']*len(columns)) for _ in range(n)]))
            matrix_text.delete('1.0','end');matrix_text.insert('1.0',text)
        except (ValueError,TypeError) as error:
            messagebox.showerror('Check CPU selection',str(error),parent=root)
    ttk.Button(matrix_buttons,text='Import CSV...',command=import_cpu_csv).pack(side='left')
    ttk.Button(matrix_buttons,text='Insert empty template',command=insert_cpu_template).pack(side='left',padx=6)
    def update_load_mode(*_args):
        entries_by_key['powers'].state(['disabled'] if heterogeneous.get() else ['!disabled'])
        if heterogeneous.get():matrix_frame.grid()
        else:matrix_frame.grid_remove()
    heterogeneous.trace_add('write',update_load_mode);update_load_mode()
    status=tk.StringVar()
    ttk.Label(body,textvariable=status,wraplength=1000).pack(fill='x',padx=18)
    fig,ax=plt.subplots(figsize=(10,3.6));fig.tight_layout()
    widget=FigureCanvasTkAgg(fig,master=body);widget.get_tk_widget().pack(fill='both',expand=True,padx=12,pady=8)
    output=[];detected=None;detected_columns=None
    def selected_config():
        count=int(variables['board_count'].get());ids=variables['board_ids'].get().strip()
        selected=None
        if ids:
            selected=shift2dc_columns_for_ids(data, [int(v.strip()) for v in ids.split(',')])
        return {'board_count':count,'board_columns':selected}
    def draw(schedule=None):
        ax.clear();t=shift2dc_time(data)
        for c in available:
            ax.plot(t,pd.to_numeric(data[c],errors='coerce'),label=c,linewidth=.9)
        if schedule:
            for row in schedule:
                ax.axvline(row['Start [s]'],color='#344D63',linestyle='--',linewidth=.8)
                ax.axvspan(row['Average from [s]'],row['End [s]'],color='#2C9877',alpha=.18)
            ax.axvline(schedule[-1]['End [s]'],color='#344D63',linestyle='--',linewidth=.8)
        ax.set(xlabel='Elapsed time [s]',ylabel='Board temperature [°C]')
        ax.legend(ncol=4,fontsize=8);ax.grid(alpha=.2);fig.tight_layout();widget.draw()
    def detect():
        nonlocal detected,detected_columns
        try:
            columns=shift2dc_selected_boards(data,selected_config())
            detected,note=detect_shift2dc_start(data,columns)
            detected_columns=columns
            variables['start'].set(f'{detected:.1f}');status.set(note+' Check alignment against the trace.')
        except Exception as error:status.set(str(error))
    def build():
        nonlocal detected,detected_columns
        config=selected_config();n=int(variables['step_count'].get())
        durations=[float(v.strip()) for v in variables['durations'].get().split(',')]
        if len(durations)==1:durations*=n
        if n<1 or len(durations)!=n:raise ValueError('Step count and durations must agree.')
        start=float(variables['start'].get())
        columns=shift2dc_selected_boards(data,config)
        if start==detected and columns!=detected_columns:
            start,_=detect_shift2dc_start(data,columns)
            detected=start;detected_columns=columns
            variables['start'].set(f'{start:.1f}')
        if heterogeneous.get():
            config['cpu_powers']=parse_cpu_power_table(matrix_text.get('1.0','end'),columns,n)
        else:
            powers=[float(v.strip()) for v in variables['powers'].get().split(',')]
            if len(powers)!=n:raise ValueError('Enter one heat load per step.')
            config['powers_per_board']=powers
        config.update(durations_s=durations,start_s=start,
                      start_note='Automatically detected thermal start (reviewed).' if start==detected else 'User-adjusted start time.')
        boards,steps,schedule=build_shift2dc_steps(data,config);config['board_columns']=boards
        draw(schedule);status.set('Green areas: final 100 s used for averaging. Dashed lines: scheduled step boundaries.')
        return config
    def preview():
        try:build()
        except Exception as error:messagebox.showerror('Check schedule',str(error),parent=root)
    def accept():
        try:
            output.append(build());root.destroy()
        except Exception as error:messagebox.showerror('Check schedule',str(error),parent=root)
    for label,command in [('Detect start',detect),('Preview schedule',preview),('Continue to report setup',accept),('Cancel',root.destroy)]:
        ttk.Button(actions,text=label,command=command).pack(side='left',padx=5)
    draw();detect();root.mainloop();plt.close(fig)
    return output[0] if output else None


def shift2dc_result(step,metadata,boards,schedule):
    result=calculate_step_result(step,metadata,boards)
    # Retain the proven water-side and LTS calculations, but expose board names
    # and scheduled power rather than copper/PSU measurement labels.
    result={key.replace('T_CU','T_CPU').replace('DeltaT_CPU','ΔT_CPU'):value for key,value in result.items()}
    # Canonical sensor keys let old and renamed CSVs stack in the same columns.
    for key in list(result):
        match = re.fullmatch(r'T_(?:CPU|BOARD)_(\d+)(?:\s*\[.*\])?', key, re.I)
        if match:
            canonical = f'T_CPU_{int(match[1])} [°C]'
            if canonical != key:
                result[canonical] = result.pop(key)

    result['Scheduled total power [W]']=result.pop('W_IN [W]')
    result.update(schedule)
    result['Power basis']='Sum of entered per-CPU heat loads (not measured electrical power)'
    result['Water heat removal [W]']=result.pop('W_OUT [W]')
    result['Water / scheduled power [%]']=100*result['Water heat removal [W]']/result['Scheduled total power [W]']
    result['Unrecovered scheduled power [W]']=result['Scheduled total power [W]']-result['Water heat removal [W]']
    return result


def write_shift2dc_excel(results, output):
    board_columns = sorted({
        key for row in results for key in row
        if re.fullmatch(r'T_(?:CPU|BOARD)_\d+ \[°C\]', key, re.I)
    }, key=natural_text_sort_key)
    create_master_excel(shift2dc_report_rows(results), board_columns, 'LTS', output)


def write_shift2dc_pdf(details, output, report_configuration=None, fluid_properties=None):
    results = [row for detail in details for row in detail['results']]
    return create_pdf_report(
        shift2dc_report_rows(results), 'LTS', output, test_details=details,
        report_configuration=report_configuration, fluid_properties=fluid_properties,
    )


def process_shift2dc_files(raw_files,output_folder,configurations=None,open_pdf_when_done=True,report_configuration=None):
    details=[];results=[]
    for path in raw_files:
        metadata=parse_file_name(path);metadata['source_file']=path.name
        if metadata['part_type']!='LTS' or metadata['medium']!='Water':
            raise ValueError('Shift2DC currently supports water-cooled LTS tests.')
        original=read_csv_automatically(path)
        data=clean_shift2dc_power_channels(original)
        complete_nominal_conditions(metadata,data)
        if metadata.get('test_mode') == 'TR':
            if not data.equals(original):
                overwrite_shift2dc_csv(path, data)
            boards = shift2dc_board_columns(data)
            details.append(dict(source_file=path.name, metadata=metadata, data=data, cleaned=data.copy(),
                                boards=boards, t_cu_columns=boards, steps=[], schedule=[], results=[], raw_data_only=True))
            continue
        recovered=build_shift2dc_saved_steps(data)
        if recovered is None:
            config=configurations.get(path.name) if configurations is not None else show_shift2dc_dialog(data,path.name)
            if config is None:
                if configurations is not None:raise ValueError(f'Missing Shift2DC schedule for new test {path.name}.')
                print('Shift2DC report cancelled.');return
            boards,steps,schedule=build_shift2dc_steps(data,config)
            data=save_shift2dc_schedule_columns(data,boards,schedule)
            # Use the persisted representation immediately, ensuring repeatable averages.
            boards,steps,schedule=build_shift2dc_saved_steps(data)
            print(f'Saved CPU heat-load steps in {path.name}.')
        else:
            boards,steps,schedule=recovered
            print(f'Reused {len(schedule)} saved CPU heat-load step(s) from {path.name}.')
        trimmed=trim_shift2dc_cooldown(data)
        if len(trimmed) < len(data):
            print(f'Removed {len(data)-len(trimmed)} rows more than 100 s after the final heat-load step from {path.name}.')
        data=trimmed
        if not data.equals(original):
            overwrite_shift2dc_csv(path,data)
        test_results=[shift2dc_result(step,metadata,boards,row) for step,row in zip(steps,schedule)]
        results.extend(test_results)
        cleaned=data.copy()
        details.append({'source_file':path.name,'metadata':metadata,'data':data,'cleaned':cleaned,'boards':boards,'schedule':schedule,'results':test_results,'steps':steps,'t_cu_columns':boards,'raw_data_only':False})
    report_configuration, fluid_properties = configure_report(
        shift2dc_report_rows(results), 'LTS', details, report_configuration,
    )
    if report_configuration is None:
        print('Report generation cancelled by the user.')
        return
    output_folder=Path(output_folder);output_folder.mkdir(parents=True,exist_ok=True)
    excel=output_folder/'Shift2DC_LTS_Test_Averages.xlsx';pdf=output_folder/'Shift2DC_LTS_Test_Report.pdf'
    for detail in details:
        overwrite_shift2dc_csv(output_folder/detail['source_file'], detail['cleaned'])
    results = add_lts_psat_results(results, details, report_configuration)
    write_shift2dc_excel(results,excel)
    write_shift2dc_pdf(details,pdf,report_configuration,fluid_properties)
    print(f'Shift2DC Excel: {excel}\nShift2DC PDF: {pdf}')
    if open_pdf_when_done:open_pdf_automatically(pdf)
    return {'excel':excel,'pdf':pdf,'results':results,'details':details}




if __name__ == "__main__":
    launch_analysis_gui()
